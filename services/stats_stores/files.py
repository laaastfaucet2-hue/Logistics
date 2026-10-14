# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملفات إحصائيات المخازن: المسارات والأسماء وكتابة الإكسل ومعرفة الحاجة للتحديث.

الشهري: database/<السنة>/<الشهر>/09-الاحصائيات/المخازن والثلاجات/<N. التبويب>/<التبويب> — <الشهر>.xlsx
السنوي: database/<السنة>/00-السنوي/09-الاحصائيات/المخازن والثلاجات/<N. التبويب>/<التبويب> — سنة <السنة>.xlsx

* كل ورقة = جدول من جداول التبويب بنفس العنوان والأعمدة والصفوف (الشاشة = الإكسل).
* الدباجة واللوجو + سطر التبويب والفترة + العناوين BDD7EE + التوقيعان (الفحص الذكي).
* الكتابة عبر dataguard.atomic_save فقط (الملف المفتوح في Excel يُكتب أول ما يُقفل).
* الملف الذي لم يتغير محتواه لا يُعاد كتابته (بصمة المحتوى في ملف حالة مخفي).
* ملف الحالة خارج مجلدات الأقسام: <الشهر>/.stats_stores_state.json و 00-السنوي/.stats_stores_state.json
"""
import hashlib
import json
import re
from pathlib import Path

from core import arabic_numbers as arnum
from core import egtime
from core.config import SECTIONS, month_folder
from data_access import dataguard, storage

from . import facts as F
from .subtabs import SUBTAB_MAP, file_name, folder_name, period_text

ROOT_NAME = "المخازن والثلاجات"
STATE_NAME = ".stats_stores_state.json"
FORMAT_VERSION = 1
HEAD_FILL = "BDD7EE"
FOOT_FILL = "FFF2CC"
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


def _stats_index():
    return next(i for i, s in enumerate(SECTIONS, start=1) if s["key"] == "statistics")


def section_name():
    return f"{_stats_index():02d}-{SECTIONS[_stats_index() - 1]['name']}"


# ---------------------------------------------------------------- المسارات
def root(year, month=None):
    """مجلد «المخازن والثلاجات» داخل الإحصائيات — للشهر أو لمجلد السنة 00-السنوي."""
    if month:
        return storage.section_files_path(year, month, _stats_index()) / ROOT_NAME
    return storage.yearly_section_path(year, _stats_index()) / ROOT_NAME


def tab_dir(key, year, month=None):
    return root(year, month) / folder_name(key)


def tab_file(key, year, month=None):
    return tab_dir(key, year, month) / file_name(key, year, month)


def path_text(key, month=None):
    """المسار المختصر المكتوب في الشريط: 09-الاحصائيات › المخازن والثلاجات › 2. أرصدة المخازن."""
    parts = [section_name(), ROOT_NAME, folder_name(key)]
    if not month:
        parts.insert(0, storage.YEARLY_FOLDER)
    return " › ".join(parts)


def ensure_folders(year, month=None, keys=None):
    from . import visible_keys
    for key in keys or visible_keys():
        tab_dir(key, year, month).mkdir(parents=True, exist_ok=True)
    return root(year, month)


def state_path(year, month=None):
    if month:
        return storage.year_path(year) / month_folder(month) / STATE_NAME
    return storage.yearly_path(year) / STATE_NAME


def load_state(year, month=None):
    try:
        data = json.loads(state_path(year, month).read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def save_state(year, month, state):
    raw = json.dumps(state, ensure_ascii=False, indent=1, sort_keys=True).encode("utf-8")

    def writer(tmp):
        Path(tmp).write_bytes(raw)
    dataguard.atomic_save(writer, state_path(year, month), zip_check=False)


# ---------------------------------------------------------------- البصمات
def _letterhead_dirs_identity(year, months):
    out = []
    for month in months:
        folder = storage.year_path(year) / month_folder(month) / "الدباجة"
        try:
            out.append(sorted((p.name, p.stat().st_size, p.stat().st_mtime_ns)
                              for p in folder.iterdir() if p.is_file()))
        except OSError:
            out.append([])
    return out


def _signatures_identity():
    from data_access import db_letterhead as lhdb
    try:
        return lhdb.get_signatures()
    except Exception:  # noqa: BLE001 — البصمة لا تُسقط التحديث
        return {}


def fingerprint(year, month=None, today=None):
    """بصمة مدخلات ملفات النطاق: قواعد الشهور + سجل المخازن + التوقيعان + تاريخ المرجع."""
    today = today or egtime.today()
    parts = {"v": FORMAT_VERSION, "reg": F.registry_signature(), "sig": _signatures_identity()}
    if month:
        py, pm = (year, month - 1) if month > 1 else (year - 1, 12)
        parts.update(m=F.month_signature(year, month), p=F.month_signature(py, pm),
                     ly=F.month_signature(year - 1, month),
                     ref=str(F.ref_date(year, month, today)),
                     el=F.elapsed_days(year, month, today),
                     lh=_letterhead_dirs_identity(year, [month]))
    else:
        parts.update(m=[F.month_signature(year, m) for m in range(1, 13)],
                     p=[F.month_signature(year - 1, m) for m in range(1, 13)],
                     ref=str(today) if year == today.year else "",
                     lh=_letterhead_dirs_identity(year, range(1, 13)))
    raw = json.dumps(parts, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def content_hash(view, stamp, letterhead):
    tables = [[t["title"], t["head"], [[c["t"] for c in row] for row in t["rows"]],
               [[c["t"] for c in row] for row in t["foot"]]] for t in view["tables"]]
    raw = json.dumps({"v": FORMAT_VERSION, "t": tables, "s": stamp, "lh": letterhead},
                     ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def letterhead_month(year, month=None, today=None):
    """شهر الدباجة: الشهر نفسه، وفي السنوي آخر شهر فيه بيانات (فالملف ثابت أيًا كان المفتوح)."""
    if month:
        return month
    from .year import year_facts
    yf = year_facts(year, today)
    if yf["last"]:
        return yf["last"]
    today = today or egtime.today()
    return today.month if year == today.year else 12


def _letterhead_values(year, month):
    from data_access import db_letterhead as lhdb
    values = dict(lhdb.get_all(year, month))
    logo = lhdb.logo_path(year, month)
    if logo:
        stat = logo.stat()
        values["_logo"] = [logo.name, stat.st_size, stat.st_mtime_ns]
    return values


# ---------------------------------------------------------------- الكتابة
def sheet_title(title, used):
    """اسم ورقة صالح: بلا رموز Excel الممنوعة · ٣١ حرفًا كحد أقصى · بلا تكرار."""
    base = _BAD_SHEET.sub(" ", str(title or "ورقة")).strip().strip("'") or "ورقة"
    base = " ".join(base.split())[:31].strip()
    name, i = base, 2
    while name.lower() in used:
        suffix = f" {arnum.to_arabic_indic(i)}"
        name = base[:31 - len(suffix)].rstrip() + suffix
        i += 1
    used.add(name.lower())
    return name


def write_book(path, view, tab_title, period, year, lh_month):
    """ملف التبويب: ورقة لكل جدول — دباجة + سطر الفترة + العناوين + الصفوف + التوقيعان."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from documents.official_xlsx import DATA_ROW, TABLE_ROW, add_letterhead
    from services.raghibin import signatures_rows

    head_fill = PatternFill("solid", fgColor=HEAD_FILL)
    foot_fill = PatternFill("solid", fgColor=FOOT_FILL)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    book = Workbook()
    used = set()
    for index, tbl in enumerate(view["tables"]):
        ws = book.active if index == 0 else book.create_sheet()
        ws.title = sheet_title(tbl["title"], used)
        ncols = max(len(tbl["head"]), 3)
        add_letterhead(ws, year, lh_month, ncols)
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ncols)
        stamp = ws.cell(5, 1, f"{tab_title} — {tbl['title']} — {period}")
        stamp.font = Font(name="Cairo", size=11, bold=True, color="132638")
        stamp.alignment = center
        ws.row_dimensions[5].height = 22
        for col, name in enumerate(tbl["head"], start=1):
            cell = ws.cell(TABLE_ROW, col, name)
            cell.fill = head_fill
            cell.font = Font(bold=True)
            cell.alignment = center
        row_no = DATA_ROW
        widths = [len(str(h)) for h in tbl["head"]]
        for rows, is_foot in ((tbl["rows"], False), (tbl["foot"], True)):
            for row in rows:
                for col, item in enumerate(row, start=1):
                    text = item["t"] if isinstance(item, dict) else item
                    cell = ws.cell(row_no, col, text if text not in (None, "") else None)
                    if isinstance(text, str) and text.startswith("="):
                        cell.data_type = "s"          # نص يبدأ بـ«=» يبقى نصًا لا معادلة
                    cell.alignment = Alignment(horizontal="center", vertical="center",
                                               readingOrder=2)
                    if is_foot:
                        cell.font = Font(bold=True)
                        cell.fill = foot_fill
                    if col <= len(widths):
                        widths[col - 1] = max(widths[col - 1], len(str(text or "")))
                row_no += 1
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = min(max(width + 4, 10), 48)
        signatures_rows(ws, row_no + 1, year, lh_month, ncols)
    dataguard.atomic_save(book.save, path)
    return path


# ---------------------------------------------------------------- التحديث
def scope_has_data(year, month=None, today=None):
    if month:
        return F.month_facts(year, month, today)["has_data"]
    from .year import year_facts
    return bool(year_facts(year, today)["months"])


def is_fresh(year, month=None, today=None, keys=None):
    """الملفات محدّثة؟ — نفس البصمة وكل الملفات موجودة (أو نطاق فارغ مسجل بلا ملفات)."""
    from . import visible_keys
    keys = keys or visible_keys()
    state = load_state(year, month)
    if state.get("fp") != fingerprint(year, month, today):
        return False
    exists = [tab_file(k, year, month).exists() for k in keys]
    return all(exists) or (state.get("empty") and not any(exists))


def write_scope(year, month=None, keys=None, force=False, create=False, today=None):
    """يكتب ملفات نطاق (شهر أو سنة) — المتغير فقط. يرجّع عدد الملفات المكتوبة.

    create: يكتب حتى لو النطاق بلا بيانات (الشهر المفتوح الآن) — وإلا الشهر الفارغ
    يأخذ مجلداته فقط بلا ملفات فارغة.
    """
    from . import build_view, visible_keys
    all_keys = visible_keys()
    keys = [k for k in (keys or all_keys) if k in SUBTAB_MAP]
    full = set(keys) == set(all_keys)
    fp = fingerprint(year, month, today)
    state = load_state(year, month)
    paths = {k: tab_file(k, year, month) for k in keys}
    exists = {k: p.exists() for k, p in paths.items()}
    if not force and full and state.get("fp") == fp and (
            all(exists.values()) or (state.get("empty") and not create
                                     and not any(exists.values()))):
        return 0
    ensure_folders(year, month)
    if not (create or force or any(exists.values()) or scope_has_data(year, month, today)):
        if full:
            save_state(year, month, {"fp": fp, "files": {}, "empty": True,
                                     "at": egtime.now().isoformat(timespec="seconds")})
        return 0
    lh_month = letterhead_month(year, month, today)
    letterhead = _letterhead_values(year, lh_month)
    period = period_text(year, month)
    hashes = dict(state.get("files") or {})
    written = 0
    for key in keys:
        view = build_view(key, year, month or lh_month, "month" if month else "year",
                          None, today)
        title = SUBTAB_MAP[key]["name"]
        digest = content_hash(view, f"{title}|{period}", letterhead)
        if force or not exists[key] or hashes.get(key) != digest:
            write_book(paths[key], view, title, period, year, lh_month)
            written += 1
        hashes[key] = digest
    new_state = {"files": hashes, "empty": False,
                 "at": egtime.now().isoformat(timespec="seconds")}
    new_state["fp"] = fp if full else state.get("fp", "")
    save_state(year, month, new_state)
    return written


def status(key, year, month=None, busy=False):
    """حالة ملف التبويب للشريط: محدّث بوقته · جارٍ التحديث · مفتوح في Excel."""
    path = tab_file(key, year, month)
    if str(path) in dataguard.pending_paths():
        return {"state": "locked", "text": "🔒 مفتوح في Excel"}
    if busy or not path.exists():
        return {"state": "busy", "text": "⏳ جارٍ التحديث"}
    stamp = egtime.from_timestamp(path.stat().st_mtime).strftime("%H:%M")
    return {"state": "fresh", "text": f"✓ محدّث {arnum.to_arabic_indic(stamp)}"}
