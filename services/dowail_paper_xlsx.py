# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ورقتا إذن تفاريد الدولى في XLSX: مجمع الخط وورقة كل نقطة.

الناتج يطابق ورقة البرنامج: دباجة ولوجو رسميان، بيانات آلية، فترة شاملة،
جدول RTL، ومناطق التسليم والاعتماد والاستلام بالتوقيعات الرسمية المحفوظة.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import arabic_numbers as arnum
from core import dates, egtime
from core.config import MONTH_NAMES
from data_access import db_letterhead as lhdb
from documents.official_xlsx import add_letterhead

NAVY = "132638"
HEAD = "E9F1F9"
SOFT = "F8FAFC"
GRID = Side(style="thin", color="8793A0")
DARK_GRID = Side(style="thin", color="111827")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
HEAD_BORDER = Border(left=DARK_GRID, right=DARK_GRID, top=DARK_GRID, bottom=DARK_GRID)
NC = 7


def _headers(kind):
    total = "إجمالي الحصة المجمعة للخط" if kind == "line" else "إجمالي حصة النقطة"
    return ["م", "اسم الصنف والمادة الغذائية", "الوحدة", "المعدل اليومي للفرد",
            "أيام الصرف", total, "ملاحظات"]


def _safe_sheet(name):
    return "".join(char for char in (name or "") if char not in "[]:*?/\\")[:31] or "عنصر"


def _uniq_sheet(book, name):
    name = _safe_sheet(name)
    if name not in book.sheetnames:
        return name
    index = 2
    while True:
        candidate = f"{name[:29]}-{index}"
        if candidate not in book.sheetnames:
            return candidate
        index += 1


def _quantity(value):
    """كمية ظاهرة بثلاث خانات وأرقام عربية، مثل ورقة البرنامج."""
    return arnum.fmt_qty(value or 0)


def _cell(ws, row, column, value="", bold=False, border=BORDER, align="center",
          fill=None, size=10):
    cell = ws.cell(row, column, value)
    cell.font = Font(name="Cairo", size=size, bold=bold, color="000000")
    cell.alignment = Alignment(horizontal=align, vertical="center", readingOrder=2,
                               wrap_text=True)
    cell.border = border
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _merged(ws, row, first, last, text, bold=False, align="right", fill=None,
            border=None, size=10):
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    cell = ws.cell(row, first, text)
    cell.font = Font(name="Cairo", size=size, bold=bold, color=NAVY)
    cell.alignment = Alignment(horizontal=align, vertical="center", readingOrder=2,
                               wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        for column in range(first, last + 1):
            ws.cell(row, column).border = border
            if fill:
                ws.cell(row, column).fill = PatternFill("solid", fgColor=fill)
    return cell


def _doc_rows(data, line, kind, point=None):
    for paper in data.get("papers") or []:
        if int(paper["line"]["id"]) != int(line["id"]):
            continue
        if kind == "line":
            return paper["line_paper"]["rows"]
        for point_paper in paper["points_papers"]:
            if point is not None and int(point_paper["point"]["id"]) == int(point["id"]):
                return point_paper["rows"]
    return []


def _signature(ws, row, first, last, label, rank, name):
    _merged(ws, row, first, last, label, bold=True, align="center", size=10)
    _merged(ws, row + 1, first, last, rank or "......................................",
            bold=True, align="center", size=10)
    _merged(ws, row + 2, first, last, name or "......................................",
            bold=True, align="center", size=10)


def _paper_period(data, line):
    """فترة الورقة نفسها؛ لا الفترة التجميعية إذا اختلفت مدد الخطوط."""
    for paper in data.get("papers") or []:
        if int(paper["line"]["id"]) == int(line["id"]):
            day_from = int(paper.get("date_from") or data.get("date_from") or 1)
            day_to = int(paper.get("date_to") or data.get("date_to") or day_from)
            return day_from, day_to
    day_from = int(data.get("date_from") or 1)
    return day_from, int(data.get("date_to") or day_from)


def _write_doc(ws, data, line, kind, point=None):
    """ديباجة → بيانات آلية → جدول → ملاحظة → تسليم/اعتماد/استلام."""
    year, month = int(data["year"]), int(data["month"])
    day_from, day_to = _paper_period(data, line)
    day_count = max(1, day_to - day_from + 1)
    add_letterhead(ws, year, month, NC)

    today = dates.format_date(egtime.now().date())
    _merged(ws, 5, 1, 5,
            f"{MONTH_NAMES[month - 1]} {arnum.to_arabic_indic(year)}",
            bold=True, align="right", size=9)
    _merged(ws, 5, 6, 7, f"التاريخ: {today}", bold=True, align="left", size=9)

    title = "إذن صرف بون تعيينات مجمع" if kind == "line" else "إذن صرف بون تعيينات نقطة"
    _merged(ws, 6, 1, NC, title, bold=True, align="center", size=16)
    ws.row_dimensions[6].height = 27

    if kind == "line":
        intro = [
            f"الجهة الطالبة: نقاط {line['name']}",
            "نوع المستند: مستند إجمالي عام",
        ]
        metadata = [
            f"اسم القطاع الجغرافي: {line['name']} مجمع",
            f"النقاط الممثلة بالصرف: عدد {arnum.to_arabic_indic(line.get('count') or len(line.get('points') or []))} نقطة",
            f"إجمالي القوة العسكرية للخط: {arnum.to_arabic_indic(line.get('total_force') or 0)} فرد / يوم",
        ]
    else:
        intro = [
            f"النقطة الطالبة: نقطة «{point['name']}»",
            "نوع المستند: طلب صرف تعيينات",
        ]
        metadata = [
            "نوع الخدمة: نقطة حدودية",
            f"الخط الجغرافي: {line['name']}",
            f"قوة النقطة: {arnum.to_arabic_indic(point.get('force') or 0)} فرد / يوم",
        ]

    row = 7
    _merged(ws, row, 1, 4, intro[0], bold=True, fill=SOFT, border=BORDER)
    _merged(ws, row, 5, 7, intro[1], bold=True, fill=SOFT, border=BORDER)
    row += 1
    for text in metadata:
        _merged(ws, row, 1, NC, text, bold=True, border=BORDER)
        row += 1
    period = (
        "فترة بون الصرف والتحميل: عن الفترة من "
        f"{dates.format_date(f'{year:04d}-{month:02d}-{day_from:02d}')} إلى "
        f"{dates.format_date(f'{year:04d}-{month:02d}-{day_to:02d}')} "
        f"({arnum.to_arabic_indic(day_count)} أيام)"
    )
    _merged(ws, row, 1, NC, period, bold=True, border=BORDER)
    row += 1
    if kind == "point":
        _merged(ws, row, 1, NC,
                "اسم المندوب المستلم: ............................................................",
                bold=True, border=BORDER)
        row += 1

    header_row = row + 1
    for column, name in enumerate(_headers(kind), start=1):
        _cell(ws, header_row, column, name, bold=True, border=HEAD_BORDER,
              fill=HEAD, size=9)
    ws.row_dimensions[header_row].height = 32

    row = header_row + 1
    for serial, item in enumerate(_doc_rows(data, line, kind, point), start=1):
        values = [
            arnum.to_arabic_indic(serial), item["name"], item["unit"] or "—",
            _quantity(item["rate"]), arnum.to_arabic_indic(item["days"]),
            _quantity(item["total"]), "",
        ]
        for column, value in enumerate(values, start=1):
            _cell(ws, row, column, value, bold=column in (2, 6),
                  align="right" if column == 2 else "center", size=9)
        ws.row_dimensions[row].height = 22
        row += 1
    if row == header_row + 1:
        _merged(ws, row, 1, NC,
                "لا توجد أصناف مفعّلة — فعّل معدلات من «تعديل معدلات الأصناف».",
                align="center", border=BORDER)
        row += 1

    if kind == "line":
        row += 1
        note = ("هذا الكشف يمثل إجمالي أذونات ومطالب التعيينات المصروفة لكافة نقاط "
                f"{line['name']} دفعة واحدة مجمعة للتحميل.")
        _merged(ws, row, 1, NC, note, bold=True, align="center", border=BORDER, size=9)
        ws.row_dimensions[row].height = 28

    signatures = lhdb.get_signatures()
    row += 2
    if kind == "line":
        _signature(ws, row, 1, 3, "أمين العهدة (المُسلِّم)",
                   signatures["sig_right_rank"], signatures["sig_right_name"])
        _signature(ws, row, 5, 7, "الاعتماد",
                   signatures["sig_left_rank"], signatures["sig_left_name"])
    else:
        _signature(ws, row, 1, 2, "أمين العهدة (المُسلِّم)",
                   signatures["sig_right_rank"], signatures["sig_right_name"])
        _signature(ws, row, 3, 5, "الاعتماد",
                   signatures["sig_left_rank"], signatures["sig_left_name"])
        _signature(ws, row, 6, 7, "المندوب المستلم (المُستَلِم)", "", "")

    widths = [5, 31, 10, 14, 11, 18, 12]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    ws.page_setup.fitToHeight = 1
    ws.print_area = f"A1:G{row + 2}"


def _line_from_data(data, line_id):
    for paper in data.get("papers") or []:
        if int(paper["line"]["id"]) == int(line_id):
            return paper["line"]
    return None


def build_paper_wb(data, year, month):
    """كل الخطوط: شيت مجمع ثم شيت لكل نقطة، بنفس الورقة الرسمية."""
    book = Workbook()
    book.remove(book.active)
    for paper in data.get("papers") or []:
        line = paper["line"]
        ws = book.create_sheet(_uniq_sheet(book, f"مجمع-{line['name']}"))
        _write_doc(ws, data, line, "line")
        for point in line.get("points") or []:
            point_ws = book.create_sheet(_uniq_sheet(book, f"{line['name']}·{point['name']}"))
            _write_doc(point_ws, data, line, "point", point=point)
    if not book.sheetnames:
        ws = book.create_sheet("الخطوط")
        add_letterhead(ws, year, month, NC)
        _merged(ws, 6, 1, NC, "لا توجد خطوط — ابدأ من «إدارة الخطوط والنقاط»",
                bold=True, align="center")
    return book


def build_line_summary_wb(data, line):
    """ملف مجمع الخط فقط."""
    book = Workbook()
    ws = book.active
    ws.title = _uniq_sheet(book, f"مجمع-{line['name']}")
    _write_doc(ws, data, line, "line")
    return book


def build_points_wb(data, line):
    """ملف نقاط الخط: شيت رسمي مستقل لكل نقطة."""
    book = Workbook()
    book.remove(book.active)
    for point in line.get("points") or []:
        ws = book.create_sheet(_uniq_sheet(book, f"{line['name']}·{point['name']}"))
        _write_doc(ws, data, line, "point", point=point)
    if not book.sheetnames:
        ws = book.create_sheet("النقاط")
        add_letterhead(ws, int(data["year"]), int(data["month"]), NC)
        _merged(ws, 6, 1, NC, f"خط «{line['name']}» — لا توجد نقاط",
                bold=True, align="center")
    return book


def build_line_paper_wb(data, line):
    """ورقة خط واحد: المجمع أولًا ثم أوراق نقاطه."""
    saved_line = _line_from_data(data, line["id"])
    if saved_line:
        line = saved_line
    book = Workbook()
    book.remove(book.active)
    ws = book.create_sheet(_uniq_sheet(book, f"مجمع-{line['name']}"))
    _write_doc(ws, data, line, "line")
    for point in line.get("points") or []:
        point_ws = book.create_sheet(_uniq_sheet(book, f"{line['name']}·{point['name']}"))
        _write_doc(point_ws, data, line, "point", point=point)
    return book


def write_paper_file(path, data, line):
    """كتابة ورقة خط كاملة داخل أرشيف التوزيعة."""
    build_line_paper_wb(data, line).save(str(path))
