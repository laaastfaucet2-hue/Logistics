# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""الملفان المحليان لـ«دفتر 4 مخازن» — التصميم بالمللى (توجيه ١٠/١٠).

جوا {سجلات الامداد،سجلات المتعهد}/4-مخازن/ :
- دفتر الشطب اليومى.4 مخازن.xlsx :
  ديباجة → شريط أزرق «دفتر الشطب اليومى للمواد التمونية ٢ع (موجود 4 مخازن)…»
  → «عن شهر: … | إجمالي أيام الشهر: … | …» → م | البند | الوحدة | يوم 1..N | المجموع
  → إجمالي المستهلكات التموينية → 3 توقيعات.
- دفتر الشطب الشهرى.4 مخازن.xlsx :
  شيت لكل صنف: ديباجة → «دفتر الشطب» → (الوحدة/حالة الصنف/اسم الصنف/رقم الصنف)
  → هيدر أزرق (التاريخ/رقم الإذن/وارد من-منصرف إلى/الكمية المضافة/الكمية المنصرفة
  فعلي-صرفيات أخرى/الرصيد/سعر الوحدة/القيمة/ملاحظات) → السطور الملونة → صفوف فراغ.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import arabic_numbers as arnum
from core import dates
from core.config import MONTH_NAMES
from data_access import db_letterhead as lh
from documents.official_xlsx import add_letterhead
from services import wh4_build as wb
from services.warehouses_fs import cycle_dir

DAILY_FILE = "دفتر الشطب اليومى.4 مخازن.xlsx"
MONTHLY_FILE = "دفتر الشطب الشهرى.4 مخازن.xlsx"

NAVY = "132638"   # نص كحلي على أبيض
HEAD = "E9F1F9"   # ازرق فاتح خالص لرؤوس الجدوال (نص أسود عادي)
YELLOW = "FFF2CC"
PINK = "FCE4EC"
GRID = Side(style="thin", color="B8B8B8")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)


def _navy(ws, row, col, value, merge_to=None):
    cell = ws.cell(row, col, value)
    cell.fill = PatternFill("solid", fgColor=HEAD)
    cell.font = Font(name="Arial", size=11, bold=True, color="000000")
    cell.alignment = Alignment(horizontal="center", vertical="center",
                               readingOrder=2, wrap_text=True)
    cell.border = BORDER
    if merge_to:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_to)
        for c in range(col, merge_to + 1):
            ws.cell(row, c).border = BORDER
            ws.cell(row, c).fill = PatternFill("solid", fgColor=HEAD)
    return cell


def _cell(ws, row, col, value="", fill=None, bold=False, red=False):
    cell = ws.cell(row, col, value)
    cell.font = Font(name="Arial", size=11, bold=bold,
                     color="C62828" if red else "000000")
    cell.alignment = Alignment(horizontal="center", vertical="center",
                               readingOrder=2)
    cell.border = BORDER
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _q(value):
    """عرض الكمية بالعربية المشرقية (زي شاشات البرنامج بالظبط)."""
    v = float(value or 0)
    return arnum.fmt_qty(v) if v else "٠"


def snapshot(year, month, cycle):
    """يعيد كتابة الملفين (يتستدعى من مرايا الدورة بعد أي حفظ)."""
    d = cycle_dir(year, month, cycle, "wh4")
    _write_daily(year, month, cycle, d / DAILY_FILE)
    _write_monthly(year, month, cycle, d / MONTHLY_FILE)


def _write_daily(year, month, cycle, path):
    data = wb.daily_matrix(year, month, cycle)
    last = data["last"]
    ncols = 3 + last + 1
    book = Workbook()
    ws = book.active
    ws.title = "دفتر الشطب اليومى"
    add_letterhead(ws, year, month, ncols)

    # شريط العنوان الأزرق + سطر الشهر
    _navy(ws, 7, 1, "دفتر الشطب اليومى للمواد التمونية ٢ع (موجود 4 مخازن) "
                    "— سجلات الإمداد والمتعهد الحكومي", merge_to=ncols)
    ws.row_dimensions[7].height = 24
    vals = lh.get_all(year, month)
    info = (f"عن شهر: {MONTH_NAMES[int(month) - 1]} "
            f"{arnum.to_arabic_indic(str(year))} | إجمالي أيام الشهر: "
            f"{arnum.to_arabic_indic(str(last))} يومًا | {vals.get('lh_3') or ''}")
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=ncols)
    info_c = ws.cell(8, 1, info)
    info_c.font = Font(name="Arial", size=10, bold=True, color="132638")
    info_c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)

    # الهيدر: م | البند التمويني | الوحدة | يوم 1..N | المجموع
    hr = 9
    for c, name in enumerate(["م", "البند التمويني", "الوحدة"] +
                             [f"يوم {arnum.to_arabic_indic(str(d))}"
                              for d in range(1, last + 1)] + ["المجموع"], start=1):
        _navy(ws, hr, c, name)
    ws.row_dimensions[hr].height = 22

    r = hr + 1
    for it in data["rows"]:
        _cell(ws, r, 1, arnum.to_arabic_indic(str(it["serial"])))
        _cell(ws, r, 2, it["name"], bold=True)
        _cell(ws, r, 3, it["unit"] or "—")
        for d in range(1, last + 1):
            # ١٣/١٠: الخلية = تلقائي (أذون ٢ مخازن) + شطب يدوي
            v = (it["auto"].get(d) or 0) + (it["days"].get(d) or 0)
            _cell(ws, r, 3 + d, arnum.fmt_qty(v) if v else "—")
        _cell(ws, r, 4 + last, arnum.fmt_qty(it["total"]) if it["total"] else "—",
              bold=True)
        r += 1
    # إجمالي المستهلكات التموينية
    _navy(ws, r, 1, "إجمالي المستهلكات التموينية", merge_to=3)
    for d in range(1, last + 1):
        v = data["days_total"].get(d)
        _navy(ws, r, 3 + d, arnum.fmt_qty(v) if v else "—")
    _navy(ws, r, 4 + last, arnum.fmt_qty(data["grand_total"])
          if data["grand_total"] else "—")
    r += 2
    # التوقيعات الثلاثة (١٠/١٠): أمين المخزن = التوقيع الرسمي اليمين،
    # رئيس الفرع = التوقيع الرسمي الشمال — من «الدباجة والتوقيعات» —
    # وضابط التوضيح خط توقيع فاضي (زي المندوب المستلم في الدولى)
    for label, value, col in [
            ("أمين المخزن والتعيينات",
             f"{vals.get('sig_right_rank', '')} {vals.get('sig_right_name', '')}".strip(), 1),
            ("ضابط توضيح التعيينات", "(..........................)", (ncols + 1) // 2),
            ("رئيس فرع التعيينات بالقنا",
             f"{vals.get('sig_left_rank', '')} {vals.get('sig_left_name', '')}".strip(), ncols)]:
        _cell(ws, r, col, label, bold=True)
        _cell(ws, r + 1, col, value)
    for col in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(col)].width = (
            16 if col in (2,) else 6)
    book.save(str(path))


def _safe_sheet(name):
    return "".join(ch for ch in name if ch not in "[]:*?/\\")[:31] or "صنف"


def _write_monthly(year, month, cycle, path):
    items = wb.ration_items(year, month, cycle)
    book = Workbook()
    first = True
    for it in items:
        led = wb.monthly_ledger(year, month, cycle, it["name"])
        if not led:
            continue
        ws = book.active if first else book.create_sheet()
        first = False
        ws.title = _safe_sheet(it["name"])
        ncols = 10
        add_letterhead(ws, year, month, ncols)
        m = ws.cell(6, 1, f"عن شهر: {MONTH_NAMES[int(month) - 1]} "
                          f"{arnum.to_arabic_indic(str(year))}")
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=ncols)
        m.font = Font(name="Arial", size=11, bold=True, color="132638")
        m.alignment = Alignment(horizontal="center", vertical="center",
                                readingOrder=2)

        # العنوان + بيانات الصنف
        t = ws.cell(7, 1, "دفتر الشطب")
        ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=ncols)
        t.font = Font(name="Arial", size=16, bold=True, color="132638")
        t.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)
        meta = [("الوحدة: " + (led["unit"] or "—"), 1),
                ("حالة الصنف: " + (led["status"] or "—"), 3),
                ("اسم الصنف: " + led["item_name"], 4),
                ("رقم الصنف: " + arnum.to_arabic_indic(str(led["item_no"] or "—")), 6)]
        for text, col in meta:
            c = ws.cell(8, col, text)
            c.font = Font(name="Arial", size=11, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center", readingOrder=2)

        # الهيدر الرئيسي + الفرعي
        main = ["التاريخ", "رقم الإذن", "وارد من / منصرف إلى", "الكمية المضافة",
                "الكمية المنصرفة", "", "الرصيد", "سعر الوحدة", "القيمة", "ملاحظات"]
        for c, name in enumerate(main, start=1):
            if name:
                _navy(ws, 9, c, name)
        ws.merge_cells(start_row=9, start_column=5, end_row=9, end_column=6)
        ws.cell(9, 6).fill = PatternFill("solid", fgColor=HEAD)
        ws.cell(9, 6).border = BORDER
        _navy(ws, 10, 5, "فعلي")
        _navy(ws, 10, 6, "صرفيات أخرى")
        for c in (1, 2, 3, 4, 7, 8, 9, 10):
            _navy(ws, 10, c, "")
        ws.row_dimensions[9].height = 20
        ws.row_dimensions[10].height = 16

        # السطور
        r = 11
        for row in led["rows"]:
            kind = row["kind"]
            fill = YELLOW if kind in ("opener", "closing") else None
            strike_fill = PINK if kind == "strike" else None
            _cell(ws, r, 1, dates.format_date(row["date_iso"]) if row["date_iso"] else "—",
                  fill=fill)
            _cell(ws, r, 2, arnum.to_arabic_indic(str(row["permit_no"]))
                  if row["permit_no"] else "—", fill=fill)
            _cell(ws, r, 3, row["label"], fill=fill, red=(kind == "strike"),
                  bold=(kind in ("opener", "closing")))
            _cell(ws, r, 4, _q(row["added"]) if row["added"] else "—", fill=fill)
            _cell(ws, r, 5, _q(row["real"]) if (kind == "strike" and row["real"]) else "—",
                  fill=strike_fill or fill)
            _cell(ws, r, 6, _q(row["other"]) if row["other"] else "—", fill=fill)
            _cell(ws, r, 7, _q(row["balance"]), fill=fill, bold=True)
            _cell(ws, r, 8, _q(row["price"]) if row["price"] else "—", fill=fill)
            _cell(ws, r, 9, _q(row["value"]) if row["value"] else "—", fill=fill)
            _cell(ws, r, 10, row.get("notes") or "", fill=fill)
            r += 1
        for _ in range(12):      # صفوف فراغ منسقة للتدوين
            for c in range(1, ncols + 1):
                _cell(ws, r, c, "")
            r += 1
        # التوقيعان الرسميان من «الدباجة والتوقيعات» (١٠/١٠): يمين وشمال
        vals = lh.get_all(year, month)
        r += 1
        _cell(ws, r, 1, f"{vals.get('sig_right_rank', '')} "
                        f"{vals.get('sig_right_name', '')}".strip(), bold=True)
        _cell(ws, r, ncols, f"{vals.get('sig_left_rank', '')} "
                            f"{vals.get('sig_left_name', '')}".strip(), bold=True)
        widths = [13, 9, 22, 11, 9, 11, 10, 10, 10, 14]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if first:
        # مفيش بيانات شهرية — شيت واحد بالدباجة عشان الملف يبقى رسمي (١٠/١٠)
        ws = book.active
        ws.title = "لا بيانات"
        ncols = 10
        add_letterhead(ws, year, month, ncols)
        m = ws.cell(6, 1, f"عن شهر: {MONTH_NAMES[int(month) - 1]} "
                          f"{arnum.to_arabic_indic(str(year))}")
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=ncols)
        m.font = Font(name="Arial", size=11, bold=True, color="132638")
        m.alignment = Alignment(horizontal="center", vertical="center",
                                readingOrder=2)
        t = ws.cell(7, 1, "دفتر الشطب — لا توجد بيانات شهرية لهذا الشهر")
        ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=ncols)
        t.font = Font(name="Arial", size=14, bold=True, color="132638")
        t.alignment = Alignment(horizontal="center", vertical="center",
                                readingOrder=2)
        vals = lh.get_all(year, month)
        r = 10
        _cell(ws, r, 1, f"{vals.get('sig_right_rank', '')} "
                        f"{vals.get('sig_right_name', '')}".strip(), bold=True)
        _cell(ws, r, ncols, f"{vals.get('sig_left_rank', '')} "
                            f"{vals.get('sig_left_name', '')}".strip(), bold=True)
    book.save(str(path))
