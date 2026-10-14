# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملفات «3. تجميع الكشف الشهري العام/» — جدول أيام الشهر لكل جهة.

ملف واحد لكل جهة «<الجهة>.xlsx» وشيتان للضباط والأفراد. يعرض كل شيت:
م | الرتبة / الاسم الكامل | أيام الشهر | الإجمالي. خلية يوم الرغبة تحمل
ترتيب الوجبة الفعلية للشخص، والأيام بلا وجبة تحمل نقطة، ثم صف إجماليات وتوقيعان.
يُبنى الملف لحظيًا بعد كل علامة يومية (بناء مستهدف لملف الجهة).
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import arabic_numbers as arnum
from core import colors as palette
from core import egtime
from core.config import MONTH_NAMES
from data_access import dataguard
from data_access import db_raghibin as dr
from documents.official_xlsx import add_letterhead

from . import signatures_rows, tab_dir

CENTER = Alignment(horizontal="center", vertical="center", readingOrder=2)


def _headers(year, month):
    days = range(1, egtime.days_in_month(year, month) + 1)
    day_headers = [
        f"{arnum.to_arabic_indic(str(day))}\n{egtime.weekday_ar(date(year, month, day))}"
        for day in days
    ]
    return ["م", "الرتبة / الاسم الكامل", *day_headers, "الإجمالي"]


def _sheet(ws, year, month, entity_name, category_label, group_label,
           persons, meal_days):
    headers = _headers(year, month)
    cols = len(headers)
    month_length = len(headers) - 3
    add_letterhead(ws, year, month, cols)
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=cols)
    title_text = (f"تجميع الكشف الشهري العام — {category_label} — {entity_name} — "
                  f"{MONTH_NAMES[month - 1]} {year}")
    title = ws.cell(6, 1, arnum.to_arabic_indic(title_text))
    title.font = Font(name="Cairo", size=13, bold=True, color="132638")
    title.alignment = CENTER
    ws.row_dimensions[6].height = 22

    for col, name in enumerate(headers, start=1):
        cell = ws.cell(7, col, name)
        cell.font = Font(name="Cairo", size=9 if 2 < col < cols else 10,
                         bold=True, color="000000")
        cell.fill = PatternFill("solid", fgColor="E9F1F9")
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True, readingOrder=2)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 40
    for col in range(3, cols):
        ws.column_dimensions[get_column_letter(col)].width = 8
    ws.column_dimensions[get_column_letter(cols)].width = 14
    ws.row_dimensions[7].height = 34
    ws.freeze_panes = "C8"

    row = 8
    total_meals = 0
    for serial, person in enumerate(persons, start=1):
        days = meal_days.get(person["id"], [])
        sequence = {day: index for index, day in enumerate(days, 1)}
        total_meals += len(days)
        values = [
            arnum.to_arabic_indic(str(serial)),
            arnum.to_arabic_indic(
                f"{person['rank'] + ' / ' if person['rank'] else ''}{person['full_name']}"
            ),
            *[arnum.to_arabic_indic(str(sequence[day])) if day in sequence else "·"
              for day in range(1, month_length + 1)],
            f"{arnum.to_arabic_indic(str(len(days)))} وجبة",
        ]
        name_color = (palette.color_for("person", person["full_name"]) or "").lstrip("#")
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.alignment = CENTER
            cell.font = Font(size=10, name="Cairo", bold=(col in (2, cols)),
                             color=name_color if col == 2 else "000000")
            if row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F4F6FB")
        ws.row_dimensions[row].height = 22
        row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    people_total = ws.cell(
        row, 1,
        f"إجمالي {group_label}: {arnum.to_arabic_indic(str(len(persons)))} فرد",
    )
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=cols)
    meals_total = ws.cell(
        row, 3,
        f"إجمالي الوجبات: {arnum.to_arabic_indic(str(total_meals))} وجبة",
    )
    for cell in (people_total, meals_total):
        cell.font = Font(size=11, name="Cairo", bold=True)
        cell.alignment = CENTER
        cell.fill = PatternFill("solid", fgColor="FFF3D6")
    ws.row_dimensions[row].height = 24
    signatures_rows(ws, row + 2, year, month, cols)
    ws.print_area = f"A1:{get_column_letter(cols)}{row + 3}"
    return total_meals


def write_monthly_file(year, month, entity_id, entity_name):
    """بناء/تحديث كشف التجميع الشهري لجهة واحدة (شيتا ضباط وأفراد)."""
    folder = tab_dir(year, month, "monthly")
    meal_days = dr.month_meal_days(year, month, entity_id)
    book = Workbook()
    book.remove(book.active)
    categories = (
        ("officers", "ضابط — الضباط", "الضباط"),
        ("individuals", "فرد — الأفراد والصف", "الأفراد"),
    )
    for category_key, category_label, group_label in categories:
        persons = dr.list_persons(year, month, entity_id=entity_id,
                                  category=category_key, excluded=False)
        ws = book.create_sheet(category_label)
        _sheet(ws, year, month, entity_name, category_label, group_label,
               persons, meal_days)
    path = folder / f"{entity_name}.xlsx"
    dataguard.atomic_save(book.save, str(path))
    return path
