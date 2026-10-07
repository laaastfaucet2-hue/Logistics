# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملفات «1. (يومي) الراغبين/يوم N/<الجهة>/» — ملفان لكل جهة في اليوم.

قرار المستخدم ٠٦/١٠/٢٠٢٦: فولدر لكل يوم، وداخل فولدر اليوم **فولدر لكل جهة**،
وداخل فولدر الجهة **ملفا إكسل**: «ض — الضباط.xlsx» و«أ — الأفراد والصف.xlsx».
كل ملف بدباجة + توقيعين رسميين + لوجو (القواعد الذهبية)، وصف إجمالي في الآخر،
وبناء مستهدف عبر atomic_save (live-sync لو الملف مفتوح في Excel).
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core import arabic_numbers as arnum
from core import egtime
from core.config import MONTH_NAMES
from data_access import dataguard
from data_access import db_raghibin as dr
from documents.official_xlsx import add_letterhead

from . import FILL_EXCLUDED, signatures_rows, tab_dir
from core import colors as palette


def person_color(full_name):
    """لون الشخص الثابت (بلا #) لخط الاسم في الإكسل — نفس لون قوسه في الشاشة."""
    return (palette.color_for("person", full_name) or "").lstrip("#") or "000000"

FILL_WILLING = "E8F5EE"
CENTER = Alignment(horizontal="center", vertical="center", readingOrder=2)
# الملفان بالحروف (توجيه ض/أ/م) — العنوان يحمل الحرف ثم الاسم الكامل
CATEGORIES = (
    ("officers", "ض — الضباط.xlsx", "ض — الضباط"),
    ("individuals", "أ — الأفراد والصف.xlsx", "أ — الأفراد والصف"),
)
HEADERS = ["م", "الرتبة", "الاسم الرتبعي الكامل", "حالة اليوم", "ملاحظات"]
COLS = len(HEADERS)


def day_dir(year, month, day):
    """فولدر «يوم N» بأرقام عربية (نفس تسمية فولدرات التأميدات)."""
    path = tab_dir(year, month, "daily") / "يوم {}".format(arnum.to_arabic_indic(str(day)))
    path.mkdir(parents=True, exist_ok=True)
    return path


def entity_dir(year, month, day, entity_name):
    """فولدر الجهة داخل اليوم — داخله ملفا ض/أ (يُنشأ عند أول كتابة)."""
    path = day_dir(year, month, day) / (entity_name or "").strip()
    path.mkdir(parents=True, exist_ok=True)
    return path


def ensure_day_folders(year, month):
    """فولدرات «يوم ١..آخر الشهر» تتولد من أول الشهر (بلا ملفات — تُبنى عند التسجيل)."""
    root = tab_dir(year, month, "daily")
    for day in range(1, egtime.days_in_month(year, month) + 1):
        (root / "يوم {}".format(arnum.to_arabic_indic(str(day)))).mkdir(exist_ok=True)
    return root


def _signatures(ws, row, year, month):
    """التوقيعان الرسميان — ديليجيت للمشترك في الحزمة."""
    signatures_rows(ws, row, year, month, COLS)


def _letterhead(ws, year, month, day, entity_name, cat_title):
    from core import labels
    add_letterhead(ws, year, month, COLS)
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=COLS)
    title = ws.cell(6, 1, f"الراغبين في وجبة الطعام — يوم "
                          f"{arnum.to_arabic_indic(str(day))} — {entity_name} — "
                          f"{cat_title} — {MONTH_NAMES[month - 1]} "
                          f"{arnum.to_arabic_indic(year)}")
    title.font = Font(name="Cairo", size=13, bold=True, color="132638")
    title.alignment = CENTER
    ws.row_dimensions[6].height = 22
    legend = ws.cell(7, 1, labels.LEGEND_TWO)
    legend.font = Font(name="Cairo", size=10, color="6B7280")
    legend.alignment = CENTER
    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(8, col, name)
        cell.font = Font(name="Cairo", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="132638")
        cell.alignment = CENTER
        ws.column_dimensions[chr(64 + col)].width = max(12, len(name) * 2 + 6)
    ws.row_dimensions[8].height = 20


def _write_category(folder, year, month, day, entity_name, persons, state, cat,
                    cat_file, cat_title):
    """ملف فئة واحدة (ضباط أو أفراد) — نفس التصميم للاثنين."""
    book = Workbook()
    ws = book.active
    ws.title = cat_title[:31]
    _letterhead(ws, year, month, day, entity_name, cat_title)
    row = 9
    willing = 0
    for serial, person in enumerate(persons, start=1):
        if person["excluded"]:
            mark, fill = "⊘ غير راغب", FILL_EXCLUDED
        elif state.get(person["id"]):
            mark, fill = "✓ راغب بالوجبة", FILL_WILLING
            willing += 1
        else:
            mark, fill = "✗ لا", None
        values = [arnum.to_arabic_indic(str(serial)), person["rank"] or "—",
                  person["full_name"], mark, person["exclude_note"] or ""]
        name_color = person_color(person["full_name"])
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.alignment = CENTER
            cell.font = Font(size=11, name="Cairo", bold=(col == 3),
                             color=name_color if col == 3 else "000000")
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
        row += 1
    from core import labels
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLS)
    total = ws.cell(row, 1,
                    f"الإجمالي: {arnum.to_arabic_indic(str(willing))} راغبون من أصل "
                    f"{arnum.to_arabic_indic(str(len(persons)))} — "
                    f"{labels.SHORT[cat]} ({cat_title.split(' — ')[-1]}) — {labels.LEGEND_TWO}")
    total.font = Font(size=12, name="Cairo", bold=True)
    total.alignment = CENTER
    _signatures(ws, row + 2, year, month)
    path = folder / cat_file
    dataguard.atomic_save(book.save, str(path))   # ذرّي + لحظي لو الملف مفتوح
    return path


def write_day_file(year, month, day, entity_id, entity_name):
    """بناء/تحديث ملفي الجهة (ض/أ) داخل فولدرها في يومها — يُرجع فولدر الجهة."""
    folder = entity_dir(year, month, day, entity_name)
    persons = dr.list_persons(year, month, entity_id=entity_id)   # ضباط ثم أفراد
    state = dr.day_state(year, month, day, entity_id)
    for cat, cat_file, cat_title in CATEGORIES:
        group = [p for p in persons if (p["category"] or "individuals") == cat]
        _write_category(folder, year, month, day, entity_name, group, state,
                        cat, cat_file, cat_title)
    return folder


def day_files(year, month, day, entity_name):
    """مسارات ملفي الجهة — يُستخدمان في التنزيل والفحص الذكي."""
    folder = entity_dir(year, month, day, entity_name)
    return {cat: folder / name for cat, name, _title in CATEGORIES}


def day_zip(year, month, day, entity_name):
    """يجمع ملفي الجهة في ZIP للتنزيل من نسخة الويب — يُرجع مسار الأرشيف."""
    import tempfile
    import zipfile
    files = day_files(year, month, day, entity_name)
    safe = "_".join((entity_name or "entity").split())
    out = Path(tempfile.gettempdir()) / f"raghibin-day{day}-{abs(hash(safe)) % 10**8}.zip"
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in files.values():
            if path.is_file():
                zf.write(path, arcname=path.name)
    return out
