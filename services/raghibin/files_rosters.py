# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملفات «2. عدم الراغبين/» — سجل المستثنين (ضباط/أفراد) لكل جهة.

لكل جهة ملفان: «<الجهة> (ضباط غير راغبين).xlsx» و«<الجهة> (أفراد غير راغبين).xlsx»
الأعمدة: م | الرتبة | الاسم الرتبعي الكامل | ملاحظات / حالة الغياب + صف الإجمالي.
بالدباجة والتوقيعات الرسمية، وبالترتيب بالأقدمية (توجيه ٣٠/٠٩/٢٠٢٦) — لحظي.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core import arabic_numbers as arnum
from core.config import MONTH_NAMES
from data_access import dataguard
from data_access import db_raghibin as dr
from documents.official_xlsx import add_letterhead

from . import FILL_EXCLUDED, signatures_rows, tab_dir

CENTER = Alignment(horizontal="center", vertical="center", readingOrder=2)
HEADERS = ["م", "الرتبة", "الاسم الرتبعي الكامل", "ملاحظات / حالة الغياب"]
COLS = len(HEADERS)


def _sheet(ws, year, month, entity_name, category_label, persons):
    """category_label يحمل الحرف أولًا («ض — الضباط») توجيه ٠٦/١٠/٢٠٢٦."""
    add_letterhead(ws, year, month, COLS)
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=COLS)
    title = ws.cell(6, 1, f"سجل {category_label} — غير الراغبين — {entity_name} — "
                          f"{MONTH_NAMES[month - 1]} {arnum.to_arabic_indic(year)}")
    title.font = Font(name="Cairo", size=13, bold=True, color="132638")
    title.alignment = CENTER
    ws.row_dimensions[6].height = 22
    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(7, col, name)
        cell.font = Font(name="Cairo", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="132638")
        cell.alignment = CENTER
        ws.column_dimensions[chr(64 + col)].width = max(12, len(name) * 2 + 6)
    ws.row_dimensions[7].height = 20
    row = 8
    for serial, person in enumerate(persons, start=1):
        values = [arnum.to_arabic_indic(str(serial)), person["rank"] or "—",
                  person["full_name"], person["exclude_note"] or "—"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.alignment = CENTER
            from core import colors as palette
            name_color = (palette.color_for("person", person["full_name"]) or "").lstrip("#")
            cell.font = Font(size=11, name="Cairo", bold=(col == 3),
                             color=name_color if col == 3 else "000000")
            cell.fill = PatternFill("solid", fgColor=FILL_EXCLUDED)
        row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLS)
    from core import labels
    short = labels.SHORT["officers"] if category_label.startswith(labels.SHORT["officers"]) \
        else labels.SHORT["individuals"]
    total = ws.cell(row, 1, f"الإجمالي: {arnum.to_arabic_indic(str(len(persons)))} "
                            f"{short} غير راغبين  —  {labels.LEGEND_TWO}")
    total.font = Font(size=12, name="Cairo", bold=True)
    total.alignment = CENTER
    signatures_rows(ws, row + 2, year, month, COLS)


def write_excluded_files(year, month, entity_id, entity_name):
    """بناء ملفي المستثنين (ضباط/أفراد) لجهة واحدة — بعد كل تعليم استثناء."""
    folder = tab_dir(year, month, "excluded")
    for category_key, category_label, suffix in (
            ("officers", "ض — الضباط", "ضباط غير راغبين"),
            ("individuals", "أ — الأفراد والصف", "أفراد غير راغبين")):
        persons = dr.list_persons(year, month, entity_id=entity_id,
                                  category=category_key, excluded=True)
        book = Workbook()
        ws = book.active
        # ≤ ٣١ حرفًا (حد Excel لأسماء الشيتات)
        ws.title = f"{category_label.split(' — ')[0]} — {category_label.split(' — ')[-1]} (غير راغبين)"
        _sheet(ws, year, month, entity_name, category_label, persons)
        path = folder / f"{entity_name} ({suffix}).xlsx"
        dataguard.atomic_save(book.save, str(path))
    return folder
