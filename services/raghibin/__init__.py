# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""مولدات ملفات قسم «الراغبين» المحلية — فولدر قسم رقم ٣ «03-قسم الراغبين».

المرحلة الأولى (ج١): «4. الجهات والكوادر المعتمدة» — ملفان لكل جهة
(<الجهة> (ضباط).xlsx / <الجهة> (أفراد).xlsx) بالسجل الكامل بالحالة ✓/⊘.
كل الملفات بالدباجة الرسمية المعتمدة (الترويسة + اللوجو + التوقيعان)
وليس شكل الصور الخام — توجيه المستخدم ٣٠/٠٩/٢٠٢٦.
الكتابة لحظية عبر dataguard.atomic_save (ولو الملف مفتوح عند المستخدم
يتحدث أول ما يتقفل بنظام live-sync الموجود).
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core import arabic_numbers as arnum
from core.config import MONTH_NAMES
from data_access import dataguard, db_letterhead as lhdb, storage
from documents.official_xlsx import add_letterhead
from data_access import db_raghibin as dr

CENTER = Alignment(horizontal="center", vertical="center", readingOrder=2)
BORDER_COLOR = "8A97A8"
FILL_HEAD = "132638"
FILL_EXCLUDED = "FCE9E7"


def _section_index():
    for index, section in enumerate(SECTION_LIST(), start=1):
        if section["key"] == "raghebeen":
            return index
    raise KeyError("raghebeen")


def SECTION_LIST():
    from core.config import SECTIONS
    return SECTIONS


# أسماء الفولدرات الأربعة داخل قسم الراغبين (قرار المستخدم ٣٠/٠٩/٢٠٢٦)
FOLDERS = {
    "daily": "1. (يومي) الراغبين",
    "excluded": "2. عدم الراغبين",
    "monthly": "3. تجميع الكشف الشهري العام",
    "cadres": "4. الجهات والكوادر المعتمدة",
}


def base_dir(year, month):
    path = storage.section_files_path(year, month, _section_index())
    path.mkdir(parents=True, exist_ok=True)
    return path


def tab_dir(year, month, key):
    path = base_dir(year, month) / FOLDERS[key]
    path.mkdir(parents=True, exist_ok=True)
    return path


def ensure_folders(year, month):
    """ينشئ الفولدرات الأربعة + فولدرات يوم ١..آخر الشهر (عند فتح القسم وكل حفظ)."""
    for key in FOLDERS:
        tab_dir(year, month, key)
    from .files_daily import ensure_day_folders
    ensure_day_folders(year, month)
    return base_dir(year, month)


def cadres_dir(year, month):
    return tab_dir(year, month, "cadres")


def _title(ws, text, ncols):
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=ncols)
    cell = ws.cell(6, 1, text)
    cell.font = Font(name="Cairo", size=13, bold=True, color="132638")
    cell.alignment = CENTER
    ws.row_dimensions[6].height = 22


def _header(ws, headers):
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(7, col, name)
        cell.font = Font(name="Cairo", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=FILL_HEAD)
        cell.alignment = CENTER
        ws.column_dimensions[chr(64 + col)].width = max(12, len(name) * 2 + 6)
    ws.row_dimensions[7].height = 20


def _cell(ws, row, col, value, bold=False, fill=None):
    cell = ws.cell(row, col, value)
    cell.alignment = CENTER
    cell.font = Font(size=11, name="Cairo", bold=bold)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def signatures_rows(ws, row, year, month, ncols):
    """التوقيعان الرسميان أسفل الجدول — دالة مشتركة لكل ملفات الراغبين
    (نفس آلية ملفات التأميدات الرسمية: قيمة الرتبة سطر والاسم تحته)."""
    blocks = [(1, True, lhdb.get_setting(year, month, "sig_right_rank"),
               lhdb.get_setting(year, month, "sig_right_name"))]
    # الملفات الضيقة (٣ أعمدة زي كشف الشهر): التوقيع الثاني عمود مستقل بلا دمج
    left_col = ncols - 1 if ncols >= 4 else 3
    blocks.append((left_col, left_col + 1 <= ncols,
                   lhdb.get_setting(year, month, "sig_left_rank"),
                   lhdb.get_setting(year, month, "sig_left_name")))
    for col, merge, rank, name in blocks:
        if merge:
            ws.merge_cells(start_row=row, start_column=col,
                           end_row=row, end_column=col + 1)
        top = ws.cell(row, col, rank or "........................")
        top.font = Font(bold=True, size=12, name="Cairo")
        top.alignment = CENTER
        if merge:
            ws.merge_cells(start_row=row + 1, start_column=col,
                           end_row=row + 1, end_column=col + 1)
        bottom = ws.cell(row + 1, col, name or "........................")
        bottom.font = Font(bold=True, size=12, name="Cairo")
        bottom.alignment = CENTER


def _signatures(ws, row, ncols, year, month):
    """توافقية: الملفات القديمة تستدعي _signatures — ديليجيت للمشترك."""
    signatures_rows(ws, row, year, month, ncols)


def _cadres_sheet(ws, year, month, entity_name, category_key, category_label, persons):
    """ورقة قوة واحدة: م | الرتبة | الاسم الرتبعي الكامل | الحالة | ملاحظات."""
    headers = ["م", "الرتبة", "الاسم الرتبعي الكامل", "الحالة", "ملاحظات"]
    add_letterhead(ws, year, month, len(headers))
    _title(ws, f"الجهات والكوادر المعتمدة — {category_label} — {entity_name} — "
               f"{MONTH_NAMES[month - 1]} {arnum.to_arabic_indic(year)}", len(headers))
    _header(ws, headers)
    row = 8
    for serial, person in enumerate(persons, start=1):
        excluded = bool(person["excluded"])
        state = "⊘ غير راغب" if excluded else "✓ راغب"
        note = person["exclude_note"] if excluded else ""
        fill = FILL_EXCLUDED if excluded else ("F4F6FB" if row % 2 == 0 else None)
        values = [arnum.to_arabic_indic(serial), person["rank"] or "—",
                  person["full_name"], state, note or ""]
        for col, value in enumerate(values, start=1):
            _cell(ws, row, col, value, bold=(col == 3), fill=fill)
        row += 1
    active = sum(1 for p in persons if not p["excluded"])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    total = _cell(ws, row, 1,
                  f"الإجمالي: {arnum.to_arabic_indic(len(persons))} "
                  f"({arnum.to_arabic_indic(active)} منهم راغبون معتمدون و"
                  f"{arnum.to_arabic_indic(len(persons) - active)} غير راغبين)")
    total.font = Font(size=12, name="Cairo", bold=True)
    _signatures(ws, row + 2, len(headers), year, month)


def write_cadres_files(year, month, entity_id, entity_name):
    """بناء ملفي القوة (ضباط/أفراد) لجهة واحدة لحظيًا — يُستدعى بعد كل تعديل."""
    ensure_folders(year, month)
    persons = dr.list_persons(year, month, entity_id=entity_id)
    for category_key, category_label in (("officers", "الضباط"),
                                         ("individuals", "الأفراد والصف")):
        book = Workbook()
        ws = book.active
        ws.title = category_label
        _cadres_sheet(ws, year, month, entity_name, category_key, category_label,
                      [p for p in persons if p["category"] == category_key])
        suffix = "ضباط" if category_key == "officers" else "أفراد"
        path = cadres_dir(year, month) / f"{entity_name} ({suffix}).xlsx"
        dataguard.atomic_save(book.save, str(path))   # ذرّي + لحظي لو الملف مفتوح
    return cadres_dir(year, month)


def write_all_cadres(year, month):
    """بناء ملفات القوة لكل جهات القاموس (بعد نسخ القوة مثلًا)."""
    from data_access import db_tameedat as dt
    for entity in dt.list_entities(year, month):
        write_cadres_files(year, month, entity["id"], entity["name"])
    return cadres_dir(year, month)


def write_entity_files(year, month, entity_id, entity_name, day=None):
    """كل ملفات جهة واحدة بعد أي تعديل عليها: الكوادر + عدم الراغبين + الشهري
    (+ ملف اليوم المحدد إن تحدد) — كلها بناء مستهدف لحظي."""
    write_cadres_files(year, month, entity_id, entity_name)
    from .files_rosters import write_excluded_files
    write_excluded_files(year, month, entity_id, entity_name)
    from .files_monthly import write_monthly_file
    write_monthly_file(year, month, entity_id, entity_name)
    if day:
        from .files_daily import write_day_file
        write_day_file(year, month, day, entity_id, entity_name)


def write_all(year, month):
    """بناء كل ملفات الراغبين لكل الجهات (بعد نسخ القوة أو إعادة البذر)."""
    from data_access import db_tameedat as dt
    from .files_rosters import write_excluded_files
    from .files_monthly import write_monthly_file
    for entity in dt.list_entities(year, month):
        write_cadres_files(year, month, entity["id"], entity["name"])
        write_excluded_files(year, month, entity["id"], entity["name"])
        write_monthly_file(year, month, entity["id"], entity["name"])
    return base_dir(year, month)
