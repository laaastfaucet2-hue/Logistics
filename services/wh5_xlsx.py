# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملف «دفتر حركة صادر ووارد المخزن (5 مخازن).xlsx» — التصميم بالمللى زي المرجع
(توجيه المستخدم ١١/١٠): دباجة رسمية + لوجو → شريط كحلي داكن بالعنوان → سطر
الشهر → هيدر أزرق (10 أعمدة) → الأسطر أو «لا توجد حركات…» → الإجمالي العام
لحركات الشهر → 3 سطور توقيعات (أمين عهد المخزن / ضابط تأميني التعبيئات /
رئيس فرع التعبيئات). نفس الملف بالتصميم نفسه في الدورات الثلاث: الإمداد/
المتعهد/الترفيه — يفرق عنوان الشريط فقط.

جوا {سجلات الإمداد،سجلات المتعهد}/٥ مخازن/ و 11-الترفية/٥ مخازن/.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import arabic_numbers as arnum
from core import dates
from core.config import MONTH_NAMES
from data_access import db_letterhead as lh
from documents.official_xlsx import add_letterhead
from services import wh5_build as w5

FILE_NAME = "دفتر حركة صادر ووارد المخزن (5 مخازن).xlsx"
SHEET = "دفتر 5 مخازن"

DARK = "1F3864"       # الشريط الكحلي الداكن (تصميم المرجع)
BLUE = "2E5B9F"       # الهيدر الأزرق
HEAD = "E9F1F9"       # أزرق فاتح لسطر الشهر
NAVY = "132638"       # نص كحلي على أبيض
TOTAL_FILL = "FFF2CC" # صف الإجمالي
GRID = Side(style="thin", color="B8B8B8")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
NCOLS = 10
HEADERS = ["م", "التاريخ", "رقم الإذن", "الجهة", "المصدر", "مصروف إلى",
           "قيمة الأصناف المضافة", "قيمة الأصناف المنصرفة",
           "البيان والملحوظات والتفاصيل", "المسؤول والمراقب"]
WIDTHS = [4, 12, 9, 16, 15, 16, 11, 11, 46, 16]
EMPTY_MSG = ("لا توجد حركات مسجلة لهذا الشهر بعد — يظهر أول إذن إضافة أو صرف "
             "هنا تلقائيًا (يبدأ السجل من الإذونات مباشرةً)")


def _center(cell, wrap=True):
    cell.alignment = Alignment(horizontal="center", vertical="center",
                               readingOrder=2, wrap_text=wrap)
    cell.border = BORDER


def _cell(ws, row, col, value="", fill=None, bold=False, white=False,
          size=11, merge_to=None):
    cell = ws.cell(row, col, value)
    cell.font = Font(name="Arial", size=size, bold=bold,
                     color="FFFFFF" if white else NAVY)
    _center(cell)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if merge_to:
        ws.merge_cells(start_row=row, start_column=col, end_row=row,
                       end_column=merge_to)
        for c in range(col, merge_to + 1):
            ws.cell(row, c).border = BORDER
            if fill:
                ws.cell(row, c).fill = PatternFill("solid", fgColor=fill)
    return cell


def _qty(ws, row, col, value):
    """الكمية بلا أصفار زائدة (٢١٠ مش ٢١٠٫٠٠٠) — زي الأسكرين شوت المرجع."""
    if not value:
        _cell(ws, row, col, "—")
    else:
        _cell(ws, row, col, arnum.fmt_qty_trim(value), bold=True)


def _permit_label(r):
    if r["kind"] == "add1":
        return arnum.to_arabic_indic(str(r["permit_no"]))
    fiscal = int(r.get("fiscal") or 0)
    if fiscal:
        return "{}/{}".format(arnum.to_arabic_indic(str(r["permit_no"])),
                              arnum.to_arabic_indic(str(fiscal)))
    return arnum.to_arabic_indic(str(r["permit_no"]))


def write_book(path, year, month, rows, title):
    """يكتب الدفتر بالمللى — rows من wh5_build (بلا رصيد أول المدة أبدًا)."""
    sig = lh.get_signatures()
    book = Workbook()
    ws = book.active
    ws.title = SHEET
    add_letterhead(ws, year, month, NCOLS)

    # الشريط الكحلي الداكن — العنوان الكامل (صف 7)
    _cell(ws, 7, 1, title, fill=DARK, bold=True, white=True, size=14,
          merge_to=NCOLS)
    ws.row_dimensions[7].height = 30

    # سطر الشهر (صف 8)
    n = len(rows)
    info = ("عن شهر: {} {} | عدد الحركات: {}".format(
        MONTH_NAMES[int(month) - 1], arnum.to_arabic_indic(str(year)),
        arnum.to_arabic_indic(str(n))) if n else
        "عن شهر: {} {}".format(MONTH_NAMES[int(month) - 1],
                               arnum.to_arabic_indic(str(year))))
    _cell(ws, 8, 1, info, fill=HEAD, bold=True, size=10, merge_to=NCOLS)
    ws.row_dimensions[8].height = 18

    # الهيدر الأزرق (صف 9)
    for c, name in enumerate(HEADERS, start=1):
        _cell(ws, 9, c, name, fill=BLUE, bold=True, white=True, size=10)
    ws.row_dimensions[9].height = 30

    r = 10
    if not rows:
        _cell(ws, r, 1, EMPTY_MSG, merge_to=NCOLS, size=10)
        ws.row_dimensions[r].height = 22
        r += 1
    for idx, row in enumerate(rows, start=1):
        _cell(ws, r, 1, arnum.to_arabic_indic(str(idx)))
        _cell(ws, r, 2, dates.format_date(row["date_iso"]))
        _cell(ws, r, 3, _permit_label(row), bold=True)
        _cell(ws, r, 4, row["party"] or "—")
        _cell(ws, r, 5, row["source"] or "—")
        _cell(ws, r, 6, row["issued_to"] or "—")
        _qty(ws, r, 7, row["added"])
        _qty(ws, r, 8, row["issued"])
        _cell(ws, r, 9, row["details"], size=10)
        _cell(ws, r, 10, row["responsible"] or "—")
        ws.row_dimensions[r].height = max(20, 13 * (
            1 + len(row["details"]) // 60))
        r += 1

    # الإجمالي العام لحركات الشهر (كم / ورود)
    total_added = sum(float(x["added"] or 0) for x in rows)
    total_issued = sum(float(x["issued"] or 0) for x in rows)
    r += 1
    _cell(ws, r, 1, "الإجمالي العام لحركات الشهر (كم / ورود):",
          fill=TOTAL_FILL, bold=True, merge_to=6)
    _cell(ws, r, 7, arnum.fmt_qty_trim(total_added) if total_added else "—",
          fill=TOTAL_FILL, bold=True)
    _cell(ws, r, 8, arnum.fmt_qty_trim(total_issued) if total_issued else "—",
          fill=TOTAL_FILL, bold=True)
    _cell(ws, r, 9, "الحركات: {}".format(
        arnum.to_arabic_indic(str(len(rows)))), fill=TOTAL_FILL, bold=True)
    _cell(ws, r, 10, "", fill=TOTAL_FILL)
    ws.row_dimensions[r].height = 22

    # التوقيعات الثلاثة (تصميم المرجع) — التوقيعات الرسمية من «الدباجة والتوقيعات»
    r += 2
    sig_right = " ".join((sig.get("sig_right_rank") or "",
                          sig.get("sig_right_name") or "")).strip()
    sig_left = " ".join((sig.get("sig_left_rank") or "",
                         sig.get("sig_left_name") or "")).strip()
    for label, value, col in [
            ("أمين عهد المخزن", sig_right or "—", 1),
            ("ضابط تأميني التعبيئات", "(..........................)", 5),
            ("رئيس فرع التعبيئات", sig_left or "—", NCOLS)]:
        _cell(ws, r, col, label, bold=True)
        _cell(ws, r + 1, col, value)

    for c, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    book.save(str(path))
    return path


def snapshot(year, month, cycle):
    """مرايا الدورتين الإمداد/المتعهد: يكتب الدفتر في فولدر الدورة."""
    from services import warehouses_fs as wf
    path = wf.cycle_dir(year, month, cycle, "wh5") / FILE_NAME
    return write_book(path, year, month, w5.rows(year, month, cycle),
                      w5.BOOK_TITLE.format(w5.TITLES[cycle]))


def snapshot_tarfea(year, month):
    """مرآة الترفية — نفس الملف بالتصميم نفسه (توجيه ١١/١٠)."""
    from services import tarfea_fs as tf
    path = tf.sub_dir(year, month, "wh5") / FILE_NAME
    return write_book(path, year, month, w5.tarfea_rows(year, month),
                      w5.BOOK_TITLE.format(w5.TITLES["tarfea"]))
