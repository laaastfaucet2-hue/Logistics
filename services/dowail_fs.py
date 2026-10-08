# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""حفظ تفاريد الدولى على القرص: فولدر جوا بيانات الشهر (JSON + Excel)."""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from core import arabic_numbers as arnum
from data_access import months

SECTION_FOLDER = "14-تفاريد الدولى"


def _safe(name):
    return re.sub(r"[^\w\u0600-\u06FF\-]+", "_", name).strip("_") or "item"


def _sheet_name(kind, name):
    """اسم الشيت: «نقطة ٤٧» بدل «نقطة نقطة_٤٧» لو الاسم بياخد الكلمة داخلة."""
    n = (name or "").strip()
    if n.startswith(kind):
        n = n[len(kind):].strip(" _-")
    return (kind + " " + _safe(n)).strip()[:31] or kind


def section_dir(year, month):
    """فولدر القسم داخل الشهر: 14-تفاريد الدولى (جنب بقية أقسام الشهر)."""
    base = months.month_db_path(year, month).parent
    d = base / SECTION_FOLDER
    d.mkdir(parents=True, exist_ok=True)
    return d


def _sheet(ws, title, header, rows):
    ws.append([title])
    ws.append(header)
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for cell in ws[2]:
        cell.font = Font(bold=True)


def save_to_disk(year, month, data):
    """يحفظ الفولدر ويرجع (قائمة الملفات النسبية، مسار الفولدر)."""
    sec = section_dir(year, month)
    n = len([p for p in sec.iterdir() if p.is_dir()]) + 1
    folder = sec / ("حفظ " + arnum.to_arabic_indic(n) + " — يوم "
                    + arnum.to_arabic_indic(data["date_from"]) + " إلى "
                    + arnum.to_arabic_indic(data["date_to"]))
    folder.mkdir(parents=True, exist_ok=True)
    files = []
    (folder / "data.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    files.append("data.json")
    wb = Workbook()
    wb.remove(wb.active)
    for lp in data.get("papers") or []:
        line = lp["line"]
        for pp in lp.get("points_papers") or []:
            p = pp["point"]
            ws = wb.create_sheet(_sheet_name("نقطة", p["name"]))
            header = ["م", "المادة التموينية المصروفة", "الوحدة", "المعدل اليومى",
                      "أيام الصرف", "الحصيلة الإجمالية للنقطة", "ملاحظات"]
            rows = [[i + 1, r["name"], r["unit"], r["rate"], r["days"],
                     r["total"], ""] for i, r in enumerate(pp.get("rows") or [])]
            _sheet(ws, f"«{p['name']}» — «{line['name']}» (القوة {p['force']})",
                   header, rows)
        lp_rows = lp.get("line_paper") or {"rows": []}
        ws = wb.create_sheet(_sheet_name("خط", line["name"]))
        header = ["م", "اسم الصنف والعلامة التجارية", "الوحدة", "المعدل اليومى",
                  "أيام الصرف", "الحصيلة الإجمالية للخط", "ملاحظات"]
        rows = [[i + 1, r["name"], r["unit"], r["rate"], r["days"],
                 r["total"], ""] for i, r in enumerate(lp_rows.get("rows") or [])]
        _sheet(ws,
               f"إذن صرف «{line['name']}» — إجمالي القوة {line['total_force']} فرد/يوم",
               header, rows)
    matrix = data.get("matrix") or {}
    ws = wb.create_sheet("الإجمالي العام")
    item_names = [r["name"] for r in matrix.get("rates") or []]
    header = ["الخط / النقطة", "القوة"] + item_names
    rows = []
    for l in matrix.get("lines") or []:
        for pt in l.get("points") or []:
            p = pt["point"]
            rows.append(["   " + p["name"], p["force"]]
                        + [pt["cells"].get(n, 0) for n in item_names])
        rows.append([l["line"]["name"] + " — إجمالي الخط", l["force"]]
                    + [l["line_total"].get(n, 0) for n in item_names])
    rows.append(["الإجمالي العام", matrix.get("grand_force", 0)]
                + [matrix.get("items_total", {}).get(n, 0) for n in item_names])
    _sheet(ws, "بيان توزيع وإجمالي خطوط حراسة الدولى", header, rows)
    xlsx_name = "تفاريد-الدولى.xlsx"
    wb.save(folder / xlsx_name)
    files.append(xlsx_name)
    rel = [f"{SECTION_FOLDER}/{folder.name}/{f}" for f in files]
    return rel, folder
