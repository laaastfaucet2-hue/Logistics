# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""إكسل «دفتر 6 مخازن (جرد الأصناف)» — مطابق حرفيًا لملف التصميم المعتمد (1-9-2026.xlsx).

التخطيط (ورقة Sheet1، الأعمدة D..P):
- رأس: E3 نموذج 6 مخازن حكومه · O4 اسم الجهة / · J5 محضر جرد الاصناف ·
  F6 اسم صاحب العهدة / · O6 مخزن / · N8 تاريخ الجرد: <سنة>/<شهر>/<يوم>
- جدول (12-13 عناوين):
  D القيمه | E سعر الوحدة | F العجز | G:H الزيادة | I حالة الصنف |
  J:K الرصيد الدفتري | L حالة الصنف | M الموجود من واقع الجرد | N الوحدة |
  O اسم الصنف | P رقم الصنف
- توقيعات (٥ بلوكات): رئيس المصلحه (D) · مدير المخازن (F:G) ·
  توقيع لجنة الجرد (H:J، ٣ أسطر) · توقيع صاحب العهدة (M) · توقيع كاتب الشطب (N:O)
- الأرقام أعداد حقيقية (0.###) والخط Arial 11 bold في كل الجدول والتوقيعات.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

ARIAL_B = Font(name="Arial", size=11, bold=True)
THIN = Side(style="thin", color="FF000000")
MEDIUM = Side(style="medium", color="FF000000")
CENTER = Alignment(horizontal="center", vertical="center")
CENTER_W = Alignment(horizontal="center", vertical="center", wrap_text=True)

COL_WIDTHS = {"A": 8.0, "C": 21.25, "D": 12.375, "E": 8.625, "F": 10.625,
              "G": 8.5, "H": 5.0, "I": 7.875, "J": 12.5, "K": 3.625,
              "L": 8.25, "M": 16.5, "N": 7.625, "O": 14.5, "P": 7.5}
HEADER1, HEADER2 = 12, 13
FIRST_DATA = 14
NUM_FMT = "0.###"
INT_FMT = "0"
NONE_TXT = "لا يوجد"


def _numfmt(x):
    """زي ملف التصميم: رقم صحيح (فيهم ٠) → «0»، وليه كسور → «0.###»."""
    n = float(x)
    return INT_FMT if n == int(n) else NUM_FMT

# (عمود، عنوان) — العمود المدموج يُمرر كجورنال
HEADERS = [
    ("D", "القيمه"), ("E", "سعر الوحدة"), ("F", "العجز"),
    ("G", "الزيادة", "H"), ("I", "حالة الصنف"),
    ("J", "الرصيد الدفتري", "K"), ("L", "حالة الصنف"),
    ("M", "الموجود من"), ("N", "الوحدة"), ("O", "اسم الصنف"),
    ("P", "رقم\nالصنف"),
]


def _put(ws, ref, value, align=None, fmt=None):
    c = ws[ref]
    c.value = value
    c.font = ARIAL_B
    if align is not None:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    return c


def _block(ws, col_a, col_b, row, value, border=False):
    """سطر بلوك: قيمة على العمود الأول + دمج على عرض البلوك."""
    if col_b and col_b != col_a:
        ws.merge_cells(f"{col_a}{row}:{col_b}{row}")
    _put(ws, f"{col_a}{row}", value, align=CENTER)
    if border:
        c = ws[f"{col_a}{row}"]
        c.border = Border(top=MEDIUM)


def write_day_xlsx(path, year, month, day, header, rows):
    """يكتب ملف يوم الجرد في المسار المعطى (path = Path)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for letter, width in COL_WIDTHS.items():
        ws.column_dimensions[letter].width = width
    for r in range(1, 12):
        ws.row_dimensions[r].height = 15
    ws.row_dimensions[HEADER1].height = 20
    ws.row_dimensions[HEADER2].height = 20

    h = header or {}
    _put(ws, "E3", "نموذج 6 مخازن حكومه")
    _put(ws, "O4", f"اسم الجهة /{h.get('sector') or ''}")
    _put(ws, "J5", "محضر جرد الاصناف")
    _put(ws, "F6", f"اسم صاحب العهدة /  {h.get('owner_rank') or ''} / "
                   f"{h.get('owner_name') or ''}")
    _put(ws, "O6", f"مخزن /{h.get('warehouse') or ''}")
    ws.merge_cells("N8:P8")
    _put(ws, "N8", f"تاريخ الجرد: {int(year)}/{int(month)}/{int(day)}",
         align=Alignment(horizontal="center"))

    # ===== عناوين الجدول =====
    for item in HEADERS:
        col, text = item[0], item[1]
        if col == "M":      # م12 + م13 خلية على حدة (زي الملف)
            _put(ws, f"M{HEADER1}", text, align=CENTER_W)
            continue
        if len(item) > 2:   # دمج مستطيل (G:H / J:K)
            ws.merge_cells(f"{col}{HEADER1}:{item[2]}{HEADER2}")
        else:
            ws.merge_cells(f"{col}{HEADER1}:{col}{HEADER2}")
        _put(ws, f"{col}{HEADER1}", text, align=CENTER_W)
    _put(ws, f"M{HEADER2}", "واقع الجرد", align=CENTER_W)

    # ===== صفوف الأصناف =====
    n = len(rows)
    for i, r in enumerate(rows):
        row = FIRST_DATA + i
        ws.row_dimensions[row].height = 16
        _put(ws, f"P{row}", int(r.get("serial") or i + 1), align=CENTER)
        _put(ws, f"O{row}", r.get("name") or "", align=CENTER)
        _put(ws, f"N{row}", r.get("unit") or "", align=CENTER)
        counted = round(float(r.get("counted") or 0), 3)
        book = round(float(r.get("book") or 0), 3)
        _put(ws, f"M{row}", counted, align=CENTER, fmt=_numfmt(counted))
        _put(ws, f"L{row}", r.get("status_item") or "جيدة", align=CENTER)
        ws.merge_cells(f"J{row}:K{row}")
        _put(ws, f"J{row}", book, align=CENTER, fmt=_numfmt(book))
        _put(ws, f"I{row}", r.get("status_issue") or "جيدة", align=CENTER)
        ws.merge_cells(f"G{row}:H{row}")
        inc = round(float(r.get("inc") or 0), 3)
        _put(ws, f"G{row}", inc if inc else NONE_TXT, align=CENTER,
             fmt=_numfmt(inc) if inc else None)
        dec = round(float(r.get("dec") or 0), 3)
        _put(ws, f"F{row}", dec if dec else NONE_TXT, align=CENTER,
             fmt=_numfmt(dec) if dec else None)
        price = round(float(r.get("price") or 0), 3)
        _put(ws, f"E{row}", price if price else None, align=CENTER,
             fmt=_numfmt(price) if price else None)
        value = r.get("value")
        if value is None:      # بيانات قديمة بلا قيمه — نحسبها
            value = round(price * counted, 3)
        value = round(float(value or 0), 3)
        _put(ws, f"D{row}", value if value else None, align=CENTER,
             fmt=_numfmt(value) if value else None)

    last = FIRST_DATA + n - 1
    # ===== حدود الجدول — جدول حروف مطابق لملف التصميم (يسار/يمين/فوق/تحت) =====
    _h1 = {"D": ("medium", "thin", "thin", "medium"),
           "E": ("thin", "thin", "thin", "medium"),
           "F": ("thin", "thin", "thin", "medium"),
           "G": ("thin", "thin", "thin", "medium"),
           "H": (None, "thin", "thin", None),
           "I": ("thin", "thin", "thin", "medium"),
           "J": ("thin", "thin", "thin", "medium"),
           "K": (None, "thin", "thin", None),
           "L": ("thin", "thin", "thin", "medium"),
           "M": ("thin", "thin", "medium", "thin"),
           "N": ("thin", "thin", "thin", "medium"),
           "O": ("thin", "thin", "thin", "medium"),
           "P": ("thin", "medium", "thin", "medium")}
    _h2 = {"D": ("medium", "thin", None, "medium"),
           "E": ("thin", "thin", None, "medium"),
           "F": ("thin", "thin", None, "medium"),
           "G": ("thin", None, None, "medium"),
           "H": (None, "thin", None, "medium"),
           "I": ("thin", "thin", None, "medium"),
           "J": ("thin", None, None, "medium"),
           "K": (None, "thin", None, "medium"),
           "L": ("thin", "thin", None, "medium"),
           "M": ("thin", "thin", "thin", "medium"),
           "N": ("thin", "thin", None, "medium"),
           "O": ("thin", "thin", None, "medium"),
           "P": ("thin", "medium", None, "medium")}

    def _s(style):
        return THIN if style == "thin" else MEDIUM if style == "medium" else None

    def _border(spec):
        return Border(left=_s(spec[0]), right=_s(spec[1]),
                      top=_s(spec[2]), bottom=_s(spec[3]))

    for col, spec in _h1.items():
        ws[f"{col}{HEADER1}"].border = _border(spec)
    for col, spec in _h2.items():
        ws[f"{col}{HEADER2}"].border = _border(spec)
    for row in range(FIRST_DATA, last + 1):
        for col in "DEFGHIJKLMNOP":
            ws[f"{col}{row}"].border = Border(
                left=MEDIUM if col == "D" else THIN,
                right=MEDIUM if col == "P" else THIN, top=THIN, bottom=THIN)

    # ===== التوقيعات الخمسة — دمج وحدود زي ملف التصميم حرفيًا =====
    sigs = h.get("sigs") or {}
    t = last + 1
    for row in range(t, t + 4):
        ws.row_dimensions[row].height = 16
    _block(ws, "D", None, t, "رئيس المصلحه", border=True)
    _put(ws, f"D{t + 1}", (sigs.get("chief") or {}).get("rank") or "", align=CENTER)
    _put(ws, f"D{t + 2}", (sigs.get("chief") or {}).get("name") or "", align=CENTER)
    _block(ws, "F", "G", t, "مدير المخازن", border=True)
    ws.merge_cells(f"F{t + 1}:G{t + 1}")
    ws.merge_cells(f"F{t + 2}:G{t + 2}")
    _put(ws, f"F{t + 1}", (sigs.get("manager") or {}).get("rank") or "", align=CENTER)
    _put(ws, f"F{t + 2}", (sigs.get("manager") or {}).get("name") or "", align=CENTER)
    _block(ws, "H", "J", t, "توقيع لجنة الجرد", border=True)
    comm = sigs.get("committee") or ["", "", ""]
    for i in range(3):
        ws.merge_cells(f"H{t + 1 + i}:J{t + 1 + i}")
    _put(ws, f"H{t + 1}", comm[0], align=CENTER)
    _put(ws, f"H{t + 2}", comm[1] if len(comm) > 1 else "", align=CENTER)
    _put(ws, f"H{t + 3}", comm[2] if len(comm) > 2 else "", align=CENTER)
    _block(ws, "M", None, t, "توقيع صاحب العهدة", border=True)
    ws.merge_cells(f"K{t + 1}:M{t + 1}")
    _put(ws, f"K{t + 1}", (sigs.get("owner") or {}).get("rank") or "", align=CENTER)
    _put(ws, f"M{t + 2}", (sigs.get("owner") or {}).get("name") or "", align=CENTER)
    _block(ws, "N", "O", t, "توقيع كاتب الشطب", border=True)
    ws.merge_cells(f"N{t + 1}:O{t + 1}")
    _put(ws, f"N{t + 1}", (sigs.get("writer") or {}).get("rank") or "", align=CENTER)
    _put(ws, f"O{t + 2}", (sigs.get("writer") or {}).get("name") or "", align=CENTER)
    # خط medium فوقي على كل صف التوقيات (D..P) — غير خلايا الدمج غير الأساسية
    for col in "DEFGHIJKLMNOP":
        if col in ("G", "I", "J", "O"):
            continue
        ws[f"{col}{t}"].border = Border(top=MEDIUM)

    wb.save(str(path))
    return True
