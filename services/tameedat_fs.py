# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""الملف المحلي لقسم التاميدات — فولدرات بأسماء التويبات داخل فولدر الشهر.

المكان المعتمد (يطلب المستخدم ربط كل شيء محليًا):
    database/<السنة>/<الشهر>/07-التاميدات/
        تأميدات اليوم المحدد/        ← لقطة JSON محدثة لكل سجلات الشهر
        الجهات المومدة بالشهر الحالي/ ← ملخص JSON لكل جهة مومدة
        قاموس ودليل الجهات/          ← لقطة JSON لقاموس الشهر
        طباعة التقرير الشامل/        ← ملفات التقرير (Excel بالدباجة واللوجو + Word)

قاعدة البيانات الحية تظل month.db (حماية WAL والنسخ الاحتياطي تشملها تلقائيًا)،
وهذه الملفات مرآة مقروءة تتحدث بعد كل حفظ/تعديل/حذف ناجح فقط.
"""
import json
import logging

from core import dates, egtime
from core.config import SECTIONS, MONTH_NAMES
from data_access import dataguard, storage
from data_access import db_tameedat as dt

TAB_FOLDERS = {
    "day": "تأميدات اليوم المحدد",
    "momoda": "الجهات المومدة بالشهر الحالي",
    "dict": "قاموس ودليل الجهات",
    "report": "طباعة التقرير الشامل",
}

# ملف كل تبويب (يظهر اسمه على زر «فتح ملف» — قاعدة أزرار الملفات الملزمة)
TAB_FILES = {
    "day": "سجلات التأميدات.json",
    "momoda": "ملخص الجهات المومدة.json",
    "dict": "قاموس الجهات.json",
}
# الملفات التي يفتحها المستخدم بزر «فتح الملف» — كلها Excel (توجيه المستخدم ٢٣/٠٩)
TAB_XLSX = {
    "day": "سجلات التأميدات.xlsx",
    "momoda": "ملخص الجهات المومدة.xlsx",
    "dict": "قاموس الجهات.xlsx",
}

HEADER_1 = "منطقة وسط وجنوب للأمن المركزي"
HEADER_2 = "قطاع وسط سيناء - قسم التميينات"


def _section_index():
    for index, section in enumerate(SECTIONS, start=1):
        if section["key"] == "tameedat":
            return index
    raise KeyError("tameedat")


def base_dir(year, month):
    """فولدر قسم التاميدات داخل الشهر: database/<سنة>/<شهر>/07-التاميدات"""
    path = storage.section_files_path(year, month, _section_index())
    path.mkdir(parents=True, exist_ok=True)
    return path


def tab_dir(year, month, tab):
    """فولدر تبويب واحد داخل التاميدات — يُنشأ إن لم يوجد."""
    path = base_dir(year, month) / TAB_FOLDERS[tab]
    path.mkdir(parents=True, exist_ok=True)
    return path


def ensure_folders(year, month):
    """ينشئ الفولدرات الأربعة بأسماء التويبات (تُستدعى عند فتح القسم وكل حفظ)."""
    for key in TAB_FOLDERS:
        tab_dir(year, month, key)
    return base_dir(year, month)


def report_dir(year, month):
    return tab_dir(year, month, "report")


def _meta(year, month):
    return {
        "الجهة": HEADER_1,
        "القسم": HEADER_2,
        "الشهر": MONTH_NAMES[month - 1],
        "السنة": year,
        "آخر تحديث": egtime.now().isoformat(timespec="seconds"),
    }


def _save_json(path, payload):
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    data = text.encode("utf-8")
    dataguard.atomic_save(
        lambda tmp: __import__("pathlib").Path(tmp).write_bytes(data),
        path, zip_check=False)


def _record_json(rec):
    return {
        "رقم": rec["id"],
        "اليوم": rec["day"],
        "إلى يوم": rec["day_to"],
        "عدد أيام المدة": rec["range_days"],
        "الجهة": rec["entity_name"],
        "نوع الجهة": rec["entity_type"],
        "ضباط": rec["officers"],
        "أفراد": rec["individuals"],
        "مجندين": rec["recruits"],
        "الإجمالي": rec["total"],
        "ملحقة": [{"الاسم": a["name"], "النوع": a["entity_type"],
                   "ضباط": a["officers"], "أفراد": a["individuals"],
                   "مجندين": a["recruits"]} for a in rec["attachments"]],
        "إجمالي التأميدة مع الملحقات": rec["grand_total"],
        "ملاحظات": rec["notes"],
        "آخر تعديل": rec["updated_at"],
    }


def snapshot_all(year, month):
    """يعيد كتابة اللقطات الثلاث بعد أي حفظ/تعديل/حذف ناجح — محليًا وذرّيًا."""
    ensure_folders(year, month)

    records = dt.month_records(year, month)
    _save_json(tab_dir(year, month, "day") / "سجلات التأميدات.json", {
        **_meta(year, month),
        "عدد السجلات": len(records),
        "السجلات": [_record_json(r) for r in records],
    })

    entities = dt.list_entities(year, month)
    _save_json(tab_dir(year, month, "dict") / "قاموس الجهات.json", {
        **_meta(year, month),
        "عدد الجهات": len(entities),
        "الجهات": [{
            "مسلسل": e["serial"], "الاسم": e["name"], "النوع": e["entity_type"],
            "راغبين ضباط": e["rag_officers"], "راغبين أفراد": e["rag_individuals"],
            "ملاحظات": e["notes"],
        } for e in entities],
    })

    summary, totals = dt.month_summary(year, month)
    _save_json(tab_dir(year, month, "momoda") / "ملخص الجهات المومدة.json", {
        **_meta(year, month),
        "الجهات": [{
            "الجهة": r["name"], "النوع": r["entity_type"],
            "ملحقة": r["kind"] == "attachment",
            "أيام التميد": r["active_days"], "عدد التأميدات": r["records"],
            "ضباط": r["total_officers"], "أفراد": r["total_individuals"],
            "مجندين": r["total_recruits"], "الإجمالي": r["grand_total"],
            "أيام التميد بالتواريخ": [{
                "التأميدة رقم": part["record_id"],
                "من": dates.format_date(f"{year:04d}-{month:02d}-{part['day']:02d}"),
                "إلى": dates.format_date(f"{year:04d}-{month:02d}-{part['day_to']:02d}"),
                "عدد الأيام": part["range_days"],
                "تابعة لتأميدة": (part.get("attachment") and part["parent_name"]) or "",
            } for part in r["participations"]],
            "متوسط يومي": round(r["grand_total"] and
                (r["total_officers"] + r["total_individuals"] + r["total_recruits"])
                / r["active_days"], 1) if r["active_days"] else 0,
        } for r in summary],
        "إجمالي القوة": totals["grand"],
        "الجهات المومدة": totals["entities"],
        "أيام التميد الفعلية": totals["active_days"],
    })
    _snapshot_xlsx(year, month, records, entities, summary)   # مرآة Excel التي يفتحها المستخدم (قاعدة)
    dataguard.auto_backup("write", min_minutes=20)
    return True


def _save_xlsx(path, sheets):
    """يكتب ملف Excel عربي RTL من قائمة (اسم الورقة, العناوين, الصفوف) — مرآة مقروءة للمستخدم."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    first = True
    for title, headers, rows in sheets:
        sheet = book.active if first else book.create_sheet()
        sheet.title = title
        sheet.sheet_view.rightToLeft = True
        first = False
        sheet.append(headers)
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill("solid", fgColor="F59E0B")
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append(list(row))
        for column_cells in sheet.columns:
            width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 4, 12), 46)
    book.save(path)


def _snapshot_xlsx(year, month, records, entities, summary):
    """نسخ Excel من جداول التويبات الثلاثة — توجيه المستخدم (٢٣/٠٩): الملفات كلها Excel.
    أي عطل هنا لا يوقف الحفظ: مرآة JSON تظل سليمة ويُسجَّل التحذير."""
    try:
        day_rows = []
        for idx, rec in enumerate(records, 1):
            atts = "، ".join(f"{a['name']} ({a['entity_type'] or '—'}): "
                            f"{a['officers'] + a['individuals'] + a['recruits']}"
                            for a in rec.get("attachments", []))
            day_rows.append((idx, rec["entity_name"], rec["entity_type"],
                             dates.format_date(f"{year:04d}-{month:02d}-{rec['day']:02d}"),
                             dates.format_date(f"{year:04d}-{month:02d}-{rec['day_to']:02d}"),
                             rec["officers"], rec["individuals"], rec["recruits"],
                             rec["grand_total"], atts or "—", rec.get("notes") or "—"))
        _save_xlsx(tab_dir(year, month, "day") / TAB_XLSX["day"], [(
            "سجلات التأميدات",
            ["م", "الجهة", "النوع", "من يوم", "إلى يوم", "ضباط", "أفراد", "مجندين",
             "إجمالي التأميدة", "الملحقات", "ملاحظات"], day_rows)])

        stats = dt.dict_month_stats(year, month)   # المصدر الموحد — لا منطق متوازٍ هنا
        dict_rows = []
        for e in entities:
            st = stats.get(e["id"], {})
            own = st.get("own")
            dict_rows.append((e["serial"], e["name"], e["entity_type"],
                              e["rag_officers"] if e["rag_officers"] is not None else "—",
                              e["rag_individuals"] if e["rag_individuals"] is not None else "—",
                              round(own["avg_officers"], 3) if own else "—",
                              round(own["avg_individuals"], 3) if own else "—",
                              round(own["avg_recruits"], 3) if own else "—",
                              st.get("records_total", 0),
                              (f"×{st['att_count']} — {'، '.join(st['att_parents'])}"
                               if st.get("att_count") else "—"),
                              e.get("notes") or "—"))
        _save_xlsx(tab_dir(year, month, "dict") / TAB_XLSX["dict"], [(
            "قاموس الجهات",
            ["م", "الجهة", "النوع", "راغبين ضباط", "راغبين أفراد",
             "متوسط ضباط", "متوسط أفراد", "متوسط مجندين", "عدد التأميدات", "ملحقة على", "ملاحظات"],
            dict_rows)])

        momoda_rows = [(idx, r["name"], "ملحقة" if r["kind"] == "attachment" else "رئيسية",
                        r["entity_type"], r["active_days"], r["records"], r["total_officers"],
                        r["total_individuals"], r["total_recruits"], r["grand_total"])
                       for idx, r in enumerate(summary, 1)]
        date_rows = []
        for r in summary:
            for part in r["participations"]:
                date_rows.append((r["name"], part["record_id"],
                                  dates.format_date(f"{year:04d}-{month:02d}-{part['day']:02d}"),
                                  dates.format_date(f"{year:04d}-{month:02d}-{part['day_to']:02d}"),
                                  part["range_days"],
                                  part["parent_name"] if part.get("attachment") else "—"))
        _save_xlsx(tab_dir(year, month, "momoda") / TAB_XLSX["momoda"], [
            ("ملخص الجهات المومدة",
             ["م", "الجهة", "الحالة", "النوع", "أيام التميد", "عدد التأميدات",
              "ضباط", "أفراد", "مجندين", "الإجمالي"], momoda_rows),
            ("التواريخ من - إلى",
             ["الجهة", "تأميدة رقم", "من", "إلى", "عدد الأيام", "ملحقة على"], date_rows)])
    except Exception:
        logging.exception("tameedat xlsx mirror failed — JSON snapshots are intact")
