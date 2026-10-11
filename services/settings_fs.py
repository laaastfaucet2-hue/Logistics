# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملفات «الإعدادات» المحلية — كل الملفات مشفرة بالباستورد الملكي 7799.

المكان (نظامي — مش شهري):
    database/الإعدادات/
        المستخدمون/<الاسم>.xlsx          بيانات + صلاحيات + إحصائيات كل مستخدم
        سجلات-المستخدمين/<username>.xlsx  سجل حركاته بالتاريخ والوقت بالثانية
        الملاحظات/الملاحظات.xlsx          شيت «ملاحظات» + شيت «ردود»
"""
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core import arabic_numbers as arnum
from core.paths import DATA_DIR

MASTER_PASSWORD = "7799"
ROOT = DATA_DIR / "الإعدادات"
USERS_DIR = ROOT / "المستخدمون"
LOGS_DIR = ROOT / "سجلات-المستخدمين"
NOTES_DIR = ROOT / "الملاحظات"
NOTES_FILE = NOTES_DIR / "الملاحظات.xlsx"


def _safe(name):
    """اسم ملف آمن: حروف/أرقام/مسافات + طول مقصوص (حد نظام الملفات ~255)."""
    clean = re.sub(r"[^\w\u0600-\u06FF.\- ]+", "", name or "?").strip() or "?"
    return " ".join(clean.split())[:80]


def user_file_path(user):
    return USERS_DIR / f"{_safe(user['full_name'])}.xlsx"


def log_file_path(user):
    return LOGS_DIR / f"{_safe(user['username'])}.xlsx"


# ======================================================================
# تشفير/فك — openpyxl بيكتب صريح والمكتبة تشفر الملف كله بالباستورد
# ======================================================================
def _save_encrypted(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    wb.save(buf)
    import msoffcrypto
    plain = io.BytesIO(buf.getvalue())
    with open(str(path), "wb") as f_out:
        msoffcrypto.OfficeFile(plain).encrypt(MASTER_PASSWORD, f_out)


def _load_decrypted(path):
    if not path.is_file():
        return None
    import msoffcrypto
    with open(str(path), "rb") as f:
        of = msoffcrypto.OfficeFile(f)
        of.load_key(password=MASTER_PASSWORD)
        buf = io.BytesIO()
        of.decrypt(buf)
    return load_workbook(buf)


# ======================================================================
# شيت بيانات المستخدم (بيانات + صلاحيات + إحصائيات)
# ======================================================================
def sync_user_file(user, tab_keys=None):
    """يحفظ/يحدّث شيت بيانات المستخدم (مشفر بـ 7799)."""
    from data_access import db_users as du
    total = du.total_stats(user["id"])
    wb = Workbook()
    ws = wb.active
    ws.title = "البيانات"
    rows = [
        ("الاسم", user["full_name"]),
        ("اسم المستخدم", user["username"]),
        ("الدور", "مدير النظام" if user["role"] == "admin" else "مستخدم"),
        ("الحالة", "نشط" if user["is_active"] else "معطل"),
        ("تاريخ الإنشاء", (user.get("created_at") or "")[:10]),
        ("إجمالي مدة الشغل", du.fmt_dhm(total["work_seconds"])),
        ("أيام الشغل", total["work_days"]),
        ("إجمالي الحركات", total["events"]),
    ]
    for r in rows:
        ws.append(r)
    ws2 = wb.create_sheet("الصلاحيات")
    ws2.append(["التاب المسموح"])
    for k in sorted(tab_keys or []):
        ws2.append([k])
    wb3 = wb.create_sheet("سجل الدخول")
    wb3.append(["آخر ظهور", total.get("last_ts") or ""])
    _save_encrypted(wb, user_file_path(user))
    return user_file_path(user)


def delete_user_file(user):
    for p in (user_file_path(user), log_file_path(user)):
        try:
            p.unlink(missing_ok=True)
        except OSError:
            pass


# ======================================================================
# سجل المستخدم (كل حركته بالثانية)
# ======================================================================
def sync_log_file(user, limit=20000):
    """يمراي سجل المستخدم (من أحدث لقديم) في شيت مشفر — نفس أعمدة دفتر السجلات بالظبط."""
    from data_access import db_users as du
    from core import action_names
    events = du.list_activity(user["id"], limit)
    wb = Workbook()
    ws = wb.active
    ws.title = "السجل"
    ws.append(["التاريخ والوقت", "المشغل / الرتبة", "اسم الإجراء", "التفاصيل والملخص الفني لنشاط المستخدم"])
    role = "مدير النظام 👮" if user.get("role") == "admin" else "مستخدم 👤"
    for e in events:
        act, det = action_names.log_row(e["event"], e.get("page"), e.get("details"))
        ws.append([action_names.fmt_log_time(e["ts"] or ""),
                   f"{user['username']} {role}", act, det or "—"])
    _save_encrypted(wb, log_file_path(user))
    return log_file_path(user)


def sync_all_logs():
    """يمراي سجلات كل المستخدمين النشطين (يُستدعى مع كل عملية إدخال جديدة)."""
    from data_access import db_users as du
    for u in du.list_users():
        try:
            sync_log_file(u)
        except Exception:      # noqa: BLE001 — المرايا لا تسقط الصفحة
            import logging
            logging.exception("settings log mirror failed for %s", u["username"])


def sync_all_user_files():
    """كل مستخدم يضمن ملف بياناته + ملف سجله موجودين (يُستدعى عند التشغيل).
    يغطي الحسابات القديمة اللي اتعملت قبل ما المرايا تكون موجودة."""
    from data_access import db_users as du
    for u in du.list_users():
        try:
            if not user_file_path(u).exists():
                sync_user_file(u, du.get_user_tabs(u["id"]))
            if not log_file_path(u).exists():
                sync_log_file(u)
        except Exception:      # noqa: BLE001 — المرايا لا تسقط التشغيل
            import logging
            logging.exception("settings user-file mirror failed for %s", u["username"])


# ======================================================================
# الملاحظات (شيت ملاحظات + شيت ردود — بحالة كل ملحوظة)
# ======================================================================
def sync_notes_file():
    from data_access import db_users as du
    conn = None
    from data_access.database import get_conn
    conn = get_conn(readonly=True)
    notes = conn.execute("""
        SELECT n.*, fu.full_name fu, tu.full_name tu,
               (SELECT MAX(r.created_at) FROM note_replies r WHERE r.note_id=n.id) AS last_reply
        FROM notes n
        JOIN users fu ON fu.id=n.from_user JOIN users tu ON tu.id=n.to_user
        ORDER BY n.id""").fetchall()
    replies = conn.execute("""
        SELECT r.*, n.id note_id, fu.full_name fu FROM note_replies r
        JOIN notes n ON n.id=r.note_id
        JOIN users fu ON fu.id=r.from_user
        ORDER BY r.id""").fetchall()
    conn.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "الملاحظات"
    ws.append(["#", "التاريخ", "الوقت", "من", "إلى", "النص", "الحالة", "آخر رد"])
    status_ar = {"open": "مفتوحة", "replied": "اترد عليها", "closed": "مقفولة"}
    for n in notes:
        ts = (n["created_at"] or "").split(" ")
        ws.append([n["id"], ts[0] if ts else "", ts[1] if len(ts) > 1 else "",
                   n["fu"], n["tu"], n["text"],
                   status_ar.get(n["status"], n["status"]),
                   (n["last_reply"] or "")])
    ws2 = wb.create_sheet("الردود")
    ws2.append(["#", "التاريخ", "الوقت", "من", "على الملحوظة", "النص"])
    for r in replies:
        ts = (r["created_at"] or "").split(" ")
        ws2.append([r["id"], ts[0] if ts else "", ts[1] if len(ts) > 1 else "",
                    r["fu"], r["note_id"], r["text"]])
    _save_encrypted(wb, NOTES_FILE)
    return NOTES_FILE
