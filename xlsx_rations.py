# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""مولّد ملفات Excel الحية للمقررات — تُعاد كتابتها فور أي تعديل في البرنامج.

- المقررات التمونيية: ملف واحد بـ3 شيتات (صيفي/شتوي/رمضان).
- مقررات المتعهد: نفس الـ3 شيتات + شيت مستقل لكل جهة توزيع.
- الأرقام داخل الإكسل بالأرقام العربية المشرقية، والاتجاه من اليمين لليسار.
"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import storage
import database as db
import db_rations as dr
import db_entities as de
from config import (SECTIONS, RATION_KINDS, RATION_KIND_MAP,
                    MONTH_NAMES, DAYS, MEAL_MAP)
import arabic_numbers as arnum

NAVY, GOLD, GOLD_L, GREEN, WHITE = "0A1230", "B8860B", "F6D47C", "1E7F4F", "FFFFFF"
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BORDER = Border(*[Side(style="thin", color="9AA6C7")] * 4)

SHORT = {"tamween": "tamween_rations", "contractor": "contractor_rations"}
INVALID_SHEET = re.compile(r"[\\/*?\[\]:]")


def _section_info(short):
    key = SHORT[short]
    for i, s in enumerate(SECTIONS, start=1):
        if s["key"] == key:
            return i, s["name"]
    raise KeyError(short)


def _num(v):
    """كمية بثلاثة أرقام عشرية عربية دائمًا (٠٫١٢٠ / ٧٥٫٠٠٠)؛ الفارغ أو الصفر = خلية فاضية."""
    if v in (None, "", 0):
        return ""
    return arnum.fmt_qty(v)


def _title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, text)
    c.font = Font(bold=True, size=15, color=GOLD_L, name="Cairo")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = CENTER
    ws.row_dimensions[1].height = 34


def _header(ws, row, headers):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row, j, h)
        c.font = Font(bold=True, size=12, color=WHITE, name="Cairo")
        c.fill = PatternFill("solid", fgColor=GOLD)
        c.alignment = CENTER
        c.border = _BORDER
    ws.row_dimensions[row].height = 26


def _rows(ws, start, data):
    for i, r in enumerate(data):
        for j, v in enumerate(r, start=1):
            c = ws.cell(start + i, j, v)
            c.alignment = CENTER
            c.border = _BORDER
            c.font = Font(size=11, name="Cairo",
                          bold=(j == 2))
            if (start + i) % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F4F6FB")
    return start + len(data)


def _signatures(ws, row, ncols):
    """توقيعان أسفل الجدول: اليمين (أعمدة RTL الأولى) والشمال (الأعمدة الأخيرة)."""
    right_rank = db.get_setting("sig_right_rank")
    right_name = db.get_setting("sig_right_name")
    left_rank = db.get_setting("sig_left_rank")
    left_name = db.get_setting("sig_left_name")
    blocks = [(1, right_rank, right_name), (ncols - 1, left_rank, left_name)]
    for col, rank, name in blocks:
        # قيمة الرتبة سطر وقيمة الاسم سطر تحته — بدون تسميات «الرتبة/الاسم»
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        c = ws.cell(row, col, rank or "........................")
        c.font = Font(bold=True, size=12, name="Cairo")
        c.alignment = CENTER
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        c2 = ws.cell(row + 1, col, name or "........................")
        c2.font = Font(bold=True, size=12, name="Cairo")
        c2.alignment = CENTER


def _custom_text(entries):
    parts = []
    for e in entries:
        parts.append(f"{DAYS[e['weekday']]} {MEAL_MAP.get(e['meal'], '')} {_num(e['qty'])}")
    return "، ".join(parts)


def _kind_sheet(wb, name, year, month, short, section_name, kind_key):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    kind = RATION_KIND_MAP[kind_key]
    headers = ["مسلسل", "اسم الصنف", "الوحدة", "فطار", "غداء", "عشاء", "المقرر المخصص"]
    _title(ws, f"{section_name} — مقرر {kind['name']} — {MONTH_NAMES[month-1]} {arnum.to_arabic_indic(year)}",
           len(headers))
    _header(ws, 2, headers)
    items, custom = dr.get_items(year, month, short, kind_key)
    rows = [[arnum.to_arabic_indic(it["serial"]), it["name"], it["unit"],
             _num(it["breakfast"]), _num(it["lunch"]), _num(it["dinner"]),
             _custom_text(custom.get(it["id"], []))] for it in items]
    next_row = _rows(ws, 3, rows) + 2
    _signatures(ws, next_row, len(headers))
    widths = [8, 26, 10, 10, 10, 10, 40]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A3"


def _entity_sheet(wb, entity, year, month):
    name = INVALID_SHEET.sub("", entity["name"])[:28] or "جهة"
    base = name
    i = 2
    while name in wb.sheetnames:
        name = f"{base} {i}"
        i += 1
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    headers = (["مسلسل", "اسم الصنف", "الوحدة", "فطار", "غداء", "عشاء"] + DAYS)
    _title(ws, f"توزيع مقررات المتعهد — {entity['name']} — {MONTH_NAMES[month-1]} {arnum.to_arabic_indic(year)}",
           len(headers))
    _header(ws, 2, headers)
    rows = []
    for it in entity["items"]:
        dq = it.get("day_qty", {})
        days = []
        for d in range(7):
            if d not in it["days"]:
                days.append("")
            elif dq.get(d) is None:
                days.append("✓")
            else:
                days.append(_num(dq[d]))
        rows.append([arnum.to_arabic_indic(it["serial"]), it["name"], it["unit"],
                     _num(it["breakfast"]), _num(it["lunch"]), _num(it["dinner"])] + days)
    next_row = _rows(ws, 3, rows) + 2
    _signatures(ws, next_row, len(headers))
    for j, w in enumerate([8, 24, 9, 9, 9, 9] + [9] * 7, start=1):
        ws.column_dimensions[chr(64 + j)].width = w
    ws.freeze_panes = "A3"


def xlsx_path(year, month, short):
    index, name = _section_info(short)
    folder = storage.section_files_path(year, month, index)
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"{name}.xlsx", name


def rebuild(year, month, short):
    """يعيد بناء ملف الإكسل كاملًا — يُستدعى فور أي تعديل."""
    index, section_name = _section_info(short)
    wb = Workbook()
    wb.remove(wb.active)
    for kind_key, _, _ in RATION_KINDS:
        label = f"مقرر {RATION_KIND_MAP[kind_key]['name']}"
        _kind_sheet(wb, label, year, month, short, section_name, kind_key)
    if short == "contractor":
        for entity in de.list_entities(year, month):
            _entity_sheet(wb, entity, year, month)
    path, _ = xlsx_path(year, month, short)
    wb.save(str(path))
    return path


def ensure(year, month, short):
    """يتأكد أن الملف موجود (لو اتحذف يُعاد بناؤه)، ويرجع مساره."""
    path, _ = xlsx_path(year, month, short)
    if not path.exists():
        rebuild(year, month, short)
    return path


def rebuild_all():
    """يعيد بناء كل ملفات إكسل المقررات الموجودة (بعد تغيير التوقيعات مثلًا).

    يرجع عدد الملفات التي أُعيد بناؤها. يتجاهل الشهور التي لا بيانات لها.
    """
    import months as _months
    count = 0
    for year in storage.list_years():
        for month in range(1, 13):
            if not _months.month_db_path(year, month).exists():
                continue
            for short in SHORT:
                rebuild(year, month, short)
                count += 1
    return count
