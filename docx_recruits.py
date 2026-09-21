# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""ملفات الوورد الرسمية للمجندين — محلية لحظية في database/المجندون/كشوف/:
- «كشف حالات الشهر»: مصفوفة المجندين × أيام الشهر (يُعاد بناؤه بعد كل تعديل يومية).
- «كشف إجازات الشهر»: مدى الإجازات مدموجًا (يُعاد بناؤه بعد كل تعديل).
- «تصريح إجازة»: يُولَّد عند الطلب لكل مجند/مدى ويُحفظ ويُنزَّل معًا.
كلها بالدباجة والتوقيعين الرسميين وبالأرقام العربية-الهندية، وتُكتب ذرّيًا (قاعدة ١٣).
"""
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import arabic_numbers as arnum
import database as db
import dataguard
import db_attendance as da
import egtime
import storage

NAVY = RGBColor(0x0A, 0x12, 0x30)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
MONTH_NAMES = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
               "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]
AR_D = {"السبت": "سبت", "الأحد": "أحد", "الاثنين": "اثنن", "الثلاثاء": "ثلا",
        "الأربعاء": "أرب", "الخميس": "خمي", "الجمعة": "جمع"}
STATUS_LETTER = {"حضور": "ح", "إجازة": "إ", "غياب": "غ", "مأمورية": "م", "مستشفى": "ط", "أخرى": "أ"}


def docs_dir():
    p = storage.DATA_DIR / "المجندون" / "كشوف"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _run(p, text, size=12, bold=True, color=NAVY):
    r = p.add_run(text)
    r.font.name = "Cairo"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return r


def _header(doc, title):
    logo = db.get_setting("logo_file")
    logo_path = storage.letterhead_dir() / logo if logo else None
    if logo_path and logo_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(logo_path), width=Cm(3))
    for i, size in ((1, 15), (2, 13), (3, 12), (4, 12)):
        text = db.get_setting(f"lh_{i}") or ""
        if text:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(p, text, size)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, title, 15, True, GOLD)
    doc.add_paragraph("")


def _signatures(doc):
    doc.add_paragraph("\n")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, side in ((0, "left"), (1, "right")):
        cell = table.cell(0, idx)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for key in ("rank", "name"):
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, db.get_setting(f"sig_{side}_{key}") or "........................", 12)


def _landscape(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width


# ==================== كشف حالات الشهر (مصفوفة) ====================
def rebuild_month(year, month):
    """يعيد بناء كشفي الشهر (الحالات + الإجازات) — يُستدعى بعد أي تعديل."""
    import db_recruits as dr  # تأخر الاستيراد لتفادي الدورات
    recruits = dr.list_recruits()
    matrix = da.month_matrix(year, month)
    eom = egtime.days_in_month(year, month)
    title_m = "كشف الحالات اليومية — شهر {} {}".format(
        MONTH_NAMES[month - 1], arnum.to_arabic_indic(year))

    doc = Document()
    _landscape(doc)
    _header(doc, title_m)
    table = doc.add_table(rows=1, cols=eom + 2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    _fill(hdr[0], "م", 9)
    _fill(hdr[1], "اسم المجند", 9)
    for d in range(1, eom + 1):
        _fill(hdr[1 + d], arnum.to_arabic_indic(d), 7)
    for i, r in enumerate(recruits, 1):
        cells = table.add_row().cells
        _fill(cells[0], arnum.to_arabic_indic(i), 8)
        _fill(cells[1], r["name"], 9)
        days = matrix.get(r["id"], {})
        for d in range(1, eom + 1):
            _fill(cells[1 + d], STATUS_LETTER.get(days.get(d, ""), "ـ"), 7, center=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(p, "الرموز: ح=حضور · إ=إجازة · غ=غياب · م=مأمورية · ط=مستشفى · أ=أخرى · ـ=بدون تسجيل", 9, True, GOLD)
    _signatures(doc)
    path1 = docs_dir() / "كشف-الحالات-{}-{:02d}.docx".format(year, month)
    dataguard.atomic_save(doc.save, path1)

    # ---- كشف إجازات الشهر ----
    doc2 = Document()
    _header(doc2, "كشف الإجازات — شهر {} {}".format(
        MONTH_NAMES[month - 1], arnum.to_arabic_indic(year)))
    ranges = da.leave_ranges(year, month, matrix, 1, eom)
    by_id = {r["id"]: r for r in recruits}
    t2 = doc2.add_table(rows=1, cols=5)
    t2.style = "Table Grid"
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(("م", "اسم المجند", "الرقم العسكري", "المدة", "من يوم — إلى يوم")):
        _fill(t2.rows[0].cells[j], h, 10)
    i = 0
    for rid, rs in ranges.items():
        r = by_id.get(rid)
        if not r:
            continue
        for f, t in rs:
            i += 1
            cells = t2.add_row().cells
            _fill(cells[0], arnum.to_arabic_indic(i), 9)
            _fill(cells[1], r["name"], 10)
            _fill(cells[2], arnum.to_arabic_indic(r["mil_no"]), 9)
            _fill(cells[3], arnum.to_arabic_indic(t - f + 1) + " " + ("يوم" if t - f + 1 <= 10 else "يومًا"), 9)
            _fill(cells[4], "{} — {}".format(arnum.to_arabic_indic(f), arnum.to_arabic_indic(t)), 9)
    if i == 0:
        p = doc2.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, "لا توجد إجازات مسجلة هذا الشهر", 12, True, GOLD)
    _signatures(doc2)
    path2 = docs_dir() / "كشف-الإجازات-{}-{:02d}.docx".format(year, month)
    dataguard.atomic_save(doc2.save, path2)
    return path1, path2


def _fill(cell, text, size, center=False):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    _run(p, str(text), size)


# ==================== تصريح إجازة ====================
def build_leave_permit(recruit, from_day, to_day, year, month):
    today = egtime.today()
    doc = Document()
    _header(doc, "تصريح إجازة")
    lines = [
        "صرّحت وحدة التعيينات للمجند / {}".format(recruit["name"]),
        "رقم عسكري / {}    محافظة / {}    المدينة / {}".format(
            arnum.to_arabic_indic(recruit.get("mil_no") or "—"),
            recruit.get("governorate") or "—", recruit.get("city") or "—"),
        "بإجازة اعتيادية من يوم {} شهر {} {} ".format(
            arnum.to_arabic_indic(from_day), MONTH_NAMES[month - 1], arnum.to_arabic_indic(year))
        + "وحتى يوم {} من ذات الشهر (مدة {} يومًا).".format(
            arnum.to_arabic_indic(to_day), arnum.to_arabic_indic(to_day - from_day + 1)),
        "وعلى المجند المذكور العودة إلى وحدة التعيينات فور انتهاء الإجازة،",
        "وقد حُرِّر له هذا التصريح لتقديمه لمن يهمه الأمر.",
        "حُرِّر في {} .".format(egtime.fmt_ar(today)),
    ]
    for text in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(p, text, 13)
    _signatures(doc)
    safe = "".join(ch for ch in recruit["name"] if ch.isalnum() or ch in " _-").strip().replace(" ", "_")
    path = docs_dir() / "تصاريح" / "تصريح-{}-{}-{:02d}-{:02d}.docx".format(
        safe, year, month, from_day)
    path.parent.mkdir(parents=True, exist_ok=True)
    dataguard.atomic_save(doc.save, path)
    return path


def month_sheets(year, month):
    """مسارات كشفي الشهر إن وُجدا (لأزرار التنزيل)."""
    base = docs_dir()
    p1 = base / "كشف-الحالات-{}-{:02d}.docx".format(year, month)
    p2 = base / "كشف-الإجازات-{}-{:02d}.docx".format(year, month)
    return (p1 if p1.exists() else None), (p2 if p2.exists() else None)
