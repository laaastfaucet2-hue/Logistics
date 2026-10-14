# -*- coding: utf-8 -*-
"""ملفات إحصائيات المخازن: مجلد وإكسل لكل تبويب فرعي — شهري وسنوي — متجدد مع البرنامج.

* أول فتح للتبويب ⇒ مجلدات التبويبات + ملف الشهر + ملف السنة بالأسماء الصحيحة.
* الورق = الإكسل: أعمدة وصفوف كل جدول على الشاشة هي نفسها في ورقته.
* الحفظ الناجح ⇒ طلب تحديث يتجمع (حفظات متتالية = تحديث واحد) ويكتب الملف فعلًا.
* الملف المفتوح في Excel يُكتب أول ما يُقفل · الملف الذي لم يتغير لا يُعاد كتابته.
* الفحص الذكي: الدباجة واللوجو والتوقيعان واسم الشهر في كل ملف شهري.
"""
import os
import time

from openpyxl import load_workbook

from data_access import dataguard, db_stores, storage
from services import local_audit
from services import stats_stores as ss
from services.stats_stores import files, refresher

from test_stats_stores import MONTH, YEAR, seed_basic


def _sheet_rows(ws):
    """صفوف الجدول من الورقة: العناوين في الصف ٧ والبيانات من الصف ٨ حتى التوقيعين."""
    head = [c for c in next(ws.iter_rows(min_row=7, max_row=7, values_only=True))]
    width = len([h for h in head if h is not None])
    rows = []
    for row in ws.iter_rows(min_row=8, max_row=ws.max_row, values_only=True):
        cells = list(row[:width])
        if not any(c not in (None, "") for c in cells):
            break
        rows.append(["" if c is None else c for c in cells])
    return head[:width], rows


def test_first_open_creates_tab_folders_and_month_and_year_files(client):
    seed_basic()
    assert client.get("/stats/?cycle=supply&tab=stores").status_code == 200
    month_root = storage.year_path(YEAR) / "09-سبتمبر" / "09-الاحصائيات" / "المخازن والثلاجات"
    year_root = storage.year_path(YEAR) / "00-السنوي" / "09-الاحصائيات" / "المخازن والثلاجات"
    for tab in ss.visible_subtabs():
        folder = f"{tab['no']}. {tab['name']}"
        assert (month_root / folder / f"{tab['name']} — سبتمبر ٢٠٣١.xlsx").is_file()
        assert (year_root / folder / f"{tab['name']} — سنة ٢٠٣١.xlsx").is_file()
        # مجلدات التبويبات في كل الشهور — والشهر الفارغ بلا ملفات فارغة
        empty_month = storage.year_path(YEAR) / "03-مارس" / "09-الاحصائيات" / "المخازن والثلاجات"
        assert (empty_month / folder).is_dir() and not any((empty_month / folder).iterdir())
    assert not (storage.year_path(YEAR) / "09-سبتمبر" / "10-المخازن والثلاجات"
                / "المخازن والثلاجات").exists()


def test_paper_equals_excel_for_every_subtab(app):
    seed_basic()
    ss.refresh_now(YEAR, MONTH)
    for tab in ss.visible_subtabs():
        for month in (MONTH, None):
            view = ss.build_view(tab["key"], YEAR, month or MONTH, "month" if month else "year")
            book = load_workbook(files.tab_file(tab["key"], YEAR, month))
            assert len(book.worksheets) == len(view["tables"]), tab["key"]
            for table, ws in zip(view["tables"], book.worksheets):
                head, rows = _sheet_rows(ws)
                assert head == table["head"], (tab["key"], table["title"])
                expected = [[c["t"] for c in row] for row in table["rows"] + table["foot"]]
                assert rows == expected, (tab["key"], table["title"])


def test_successful_save_requests_refresh_and_rewrites_the_file(app, client):
    refresher.reset()
    seed_basic()
    client.get("/stats/?cycle=supply&tab=stores")
    path = files.tab_file("summary", YEAR, MONTH)

    def stores_count():
        ws = load_workbook(path).worksheets[0]
        _head, rows = _sheet_rows(ws)
        return next(r[1] for r in rows if r[0] == "المخازن")

    assert stores_count() == "١"
    response = client.post("/stores/save", data={"name": "ثلاجة اللحوم"})
    assert response.status_code == 302 and "err=" not in response.headers["Location"]
    assert refresher.busy(YEAR)                     # الطلب سُجّل (بلا انتظار الشاشة)
    refresher.flush()
    assert stores_count() == "٢"
    assert not refresher.busy(YEAR)


def test_failed_save_does_not_request_refresh(app, client):
    refresher.reset()
    response = client.post("/stores/save", data={"name": ""})
    assert "err=" in response.headers["Location"]
    assert not refresher.busy(YEAR)


def test_many_saves_are_grouped_into_one_refresh(app, monkeypatch):
    calls = []
    monkeypatch.setattr(refresher, "run_year", lambda year, **kw: calls.append(year))
    refresher.reset()
    for _ in range(10):
        refresher.schedule(YEAR, MONTH, background=False)
    refresher.flush()
    assert calls == [YEAR]


def test_unchanged_content_is_not_rewritten(app):
    seed_basic()
    ss.refresh_now(YEAR, MONTH)
    path = files.tab_file("balances", YEAR, MONTH)
    before = path.stat().st_mtime_ns
    state = files.load_state(YEAR, MONTH)
    state["fp"] = "قديم"
    files.save_state(YEAR, MONTH, state)            # بصمة مختلفة والمحتوى نفسه
    assert files.write_scope(YEAR, MONTH) == 0
    assert path.stat().st_mtime_ns == before


def test_file_open_in_excel_is_written_after_it_is_closed(app, monkeypatch):
    seed_basic()
    ss.refresh_now(YEAR, MONTH)
    path = files.tab_file("summary", YEAR, MONTH)
    monkeypatch.setattr(dataguard, "RETRY_SECONDS", 0.05)
    real_replace = os.replace
    lock = {"on": True}

    def replace(src, dst):
        if lock["on"] and str(dst) == str(path):
            raise PermissionError("الملف مفتوح")
        return real_replace(src, dst)

    monkeypatch.setattr(os, "replace", replace)
    db_stores.add_store("مخزن جديد")
    files.write_scope(YEAR, MONTH, keys=["summary"], force=True)
    assert str(path) in dataguard.pending_paths()
    assert files.status("summary", YEAR, MONTH)["state"] == "locked"
    lock["on"] = False
    deadline = time.time() + 5
    while str(path) in dataguard.pending_paths() and time.time() < deadline:
        time.sleep(0.05)
    assert str(path) not in dataguard.pending_paths()
    rows = _sheet_rows(load_workbook(path).worksheets[0])[1]
    assert next(r[1] for r in rows if r[0] == "المخازن") == "٢"


def test_monthly_files_pass_the_smart_audit(app):
    seed_basic()
    ss.refresh_now(YEAR, MONTH)
    for tab in ss.visible_subtabs():
        assert local_audit.audit_workbook(files.tab_file(tab["key"], YEAR, MONTH), YEAR, MONTH) == []


def test_sheet_names_are_valid_short_and_unique():
    used = set()
    names = [files.sheet_title(t, used) for t in ("مخزن/الأرز:رقم*1", "م" * 40, "م" * 40, "")]
    assert all(len(n) <= 31 for n in names)
    assert len(set(names)) == len(names)
    assert not any(ch in names[0] for ch in "[]:*?/\\")


def test_open_buttons_say_clearly_when_running_in_web_preview(app, client, monkeypatch):
    import routes.stats_stores as route
    monkeypatch.setattr(route, "_open_path", lambda path: False)
    seed_basic()
    folder = client.get("/stats/stores/open-folder?sub=balances")
    assert folder.status_code == 302
    assert "err=" in folder.headers["Location"]
    workbook = client.get("/stats/stores/open-file?sub=balances&scope=year")
    assert "/stats/stores/download" in workbook.headers["Location"]
    download = client.get("/stats/stores/download?sub=balances&scope=year")
    assert download.status_code == 200 and download.data[:2] == b"PK"
    status = client.get("/stats/stores/status?sub=balances").get_json()
    assert status["ok"] and status["file"] == "أرصدة المخازن — سبتمبر ٢٠٣١.xlsx"


# ---------------------------------------------------------------- المرحلة ٤: الحصر والفحص والملف القديم
def test_live_files_panel_lists_the_stats_folders_for_month_and_year(client):
    seed_basic()
    page = client.get("/stats/?cycle=supply&tab=stores&sub=balances").get_data(as_text=True)
    assert 'id="ftZone"' in page and "section=stats_stores" in page
    assert "الملفات المحلية على جهازك — إحصائيات المخازن والثلاجات" in page
    data = client.get("/files/tree?section=stats_stores").get_json()
    root = files.root(YEAR, MONTH)
    assert data["ok"] and data["root_name"] == "المخازن والثلاجات"
    assert {c["name"] for c in data["children"]} == {f"{t['no']}. {t['name']}"
                                                    for t in ss.visible_subtabs()}
    on_disk = [p for p in root.rglob("*.xlsx") if not p.name.startswith("~$")]
    assert data["files"] == len(on_disk) == len(ss.visible_subtabs())
    yearly = client.get("/files/tree?section=stats_stores&scope=year").get_json()
    assert "00-السنوي" in yearly["root"] and yearly["files"] == len(ss.visible_subtabs())
    year_page = client.get("/stats/?cycle=supply&tab=stores&scope=year").get_data(as_text=True)
    assert "scope=year" in year_page.split('id="ftZone"')[1].split(">")[0]


def test_old_stores_stats_file_is_removed_and_not_written_again(app, client):
    from services import stores_fs
    seed_basic()
    old = stores_fs.base_dir(YEAR, MONTH) / "إحصائيات المخازن.xlsx"
    old.write_bytes(b"old")
    stores_fs.ensure_folders(YEAR, MONTH)
    assert not old.exists()
    old.write_bytes(b"old")
    stores_fs.snapshot(YEAR, MONTH)
    assert not old.exists()
    assert client.get("/stores/").status_code == 200
    assert not old.exists()
    assert "stats" not in stores_fs.TAB_FILES


def test_yearly_files_pass_the_smart_audit_with_the_year_stamp(app, client):
    seed_basic()
    ss.refresh_now(YEAR, MONTH)
    for tab in ss.visible_subtabs():
        path = files.tab_file(tab["key"], YEAR)
        assert local_audit.audit_workbook(path, YEAR, None, yearly=True) == [], tab["key"]
    report = local_audit.audit_month(YEAR, MONTH, sections=("stats_stores",))
    assert not report["problems"]
    assert set(report["sections"]) == {"stats_stores", "stats_stores_year"}
    assert report["checked"] == 2 * len(ss.visible_subtabs())
    data = client.get("/files/audit?section=stats_stores&scope=year").get_json()
    assert data["ok"] and data["checked"] == len(ss.visible_subtabs())
    page = client.get("/files/audit-page?section=stats_stores").get_data(as_text=True)
    assert "إحصائيات المخازن والثلاجات — السنوي" in page


def test_issue_and_receipt_saves_also_refresh_the_stores_section_files(app, monkeypatch):
    """قرار ٦: بعد حفظ إذون ١ و٢ مخازن تتحدث مرايا قسم المخازن نفسه (في نفس التحديث المجمّع)."""
    from services import stores_fs
    calls = []
    monkeypatch.setattr(stores_fs, "snapshot", lambda year, month: calls.append((year, month)))
    refresher.reset()
    assert ss.after_save("/warehouses/wh1/add", {}, YEAR, MONTH, background=False)
    assert ss.after_save("/calc2/save", {}, YEAR, MONTH, background=False)
    assert not ss.after_save("/tameedat/save", {}, YEAR, MONTH, background=False)
    refresher.flush()
    assert calls == [(YEAR, MONTH)]                 # حفظتان ⇒ تحديث واحد للمرايا
    refresher.reset()
    ss.after_save("/stores/save", {}, YEAR, MONTH, background=False)
    calls.clear()
    refresher.flush()
    assert calls == []                              # صفحة المخازن تحدّث مراياها بنفسها
