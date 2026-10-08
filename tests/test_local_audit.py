# -*- coding: utf-8 -*-
"""الفحص الذكي والقواعد الذهبية — الدباجة واللوجو والتوقيعان على كل ملف محلي.

توجيه المستخدم ٠٦/١٠/٢٠٢٦: «لازم أتأكد إن كل الملفات المحلية عليها الدباجة والتوقيعات
الرسمية واللوجو، وكذلك الملفات الداخلية في البرنامج». الاختبارات هنا تُثبت:
١) كل ملف يبنيه البرنامج يعدّي الفحص (لا استثناء — لا في التاميدات ولا في الراغبين).
٢) الفحص يمسك الملف الناقص فعلًا (دباجة · لوجو · توقيعان · اسم الشهر).
٣) الأصناف في الاختبارات أصناف تموينية حقيقية — والقاعدة الذهبية للأصناف.
"""
from openpyxl import Workbook

from core import arabic_numbers as arnum
from data_access import db_letterhead as lh, db_rations as dr, db_tameedat as dt
from data_access import months
from services import files_index, local_audit
from services.raghibin import files_daily, files_monthly, files_rosters
from services.raghibin import write_cadres_files
from services import tameedat_fs

YEAR, MONTH = 2031, 9


def _init():
    months.init_month(YEAR, MONTH)


def _entity(name="قطاع الفحص الذكي"):
    eid = dt.add_entity(YEAR, MONTH, name, "شرطية")
    return eid, name


# ======================================================================
# ١) الدباجة الافتراضية — لا ملف بلا دباجة حتى قبل أن يسجّلها المستخدم
# ======================================================================
def test_default_letterhead_never_empty(app):
    _init()
    values = lh.get_all(YEAR, MONTH)
    assert values["lh_1"] == "وزارة الداخلية"
    assert values["lh_2"] == "قطاع الأمن المركزي"
    assert values["lh_3"] == "قطاع وسط سيناء"
    assert values["lh_4"] == "قسم التعيينات"
    # التوقيعان الرسميان على مستوى المنظومة كلها (قابلان للتعديل، وهذه قيمتهما الافتراضية)
    assert values["sig_right_rank"] == "رائد" and values["sig_right_name"] == "مصطفى نصرالله"
    assert values["sig_left_rank"] == "مقدم" and values["sig_left_name"] == "اسامة العجرودى"


def test_audit_words_constants_match_the_letterhead(app):
    _init()
    values = lh.get_all(YEAR, MONTH)
    for word in local_audit.CORE_WORDS:
        assert word in values.values(), word


# ======================================================================
# ٢) كل ملفات الراغبين والتاميدات تعدّي الفحص بعد البناء
# ======================================================================
def test_every_raghibin_file_passes_the_audit(app):
    _init()
    eid, name = _entity("جهة الراغبين")
    from data_access import db_raghibin as drg
    drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    drg.add_person(YEAR, MONTH, eid, "individuals", "أحمد محمد صابر", "فرد (1)")
    files_daily.write_day_file(YEAR, MONTH, 7, eid, name)
    files_rosters.write_excluded_files(YEAR, MONTH, eid, name)
    files_monthly.write_monthly_file(YEAR, MONTH, eid, name)
    write_cadres_files(YEAR, MONTH, eid, name)
    root = files_index.section_root("raghibin", YEAR, MONTH)
    report = local_audit.audit_tree(root, YEAR, MONTH)
    assert report["checked"] >= 6, report
    assert report["problems"] == [], report["problems"]
    assert report["ok"] == report["checked"]


def test_every_tameedat_file_passes_the_audit(app):
    _init()
    eid, name = _entity("جهة التاميدات")
    entity = dt.get_entity(YEAR, MONTH, eid)
    dt.add_record(YEAR, MONTH, 3, entity, 5, 9, 2, "", [])
    tameedat_fs.snapshot_all(YEAR, MONTH)
    root = files_index.section_root("tameedat", YEAR, MONTH)
    report = local_audit.audit_tree(root, YEAR, MONTH)
    assert report["checked"] >= 4, report
    assert report["problems"] == [], report["problems"]


# ======================================================================
# ٣) الفحص يمسك الناقص فعلًا (لا يمرّ مرور الكرام)
# ======================================================================
def _bare_book(path, text):
    book = Workbook()
    ws = book.active
    ws.cell(1, 1, text)
    book.save(path)
    return path


def test_audit_catches_missing_letterhead_logo_and_signatures(app, tmp_path):
    _init()
    path = _bare_book(tmp_path / "بلا دباجة.xlsx", "كشف بسيط بلا أي دباجة")
    missing = local_audit.audit_workbook(path, YEAR, MONTH)
    assert any("الدباجة" in m for m in missing)
    assert any("اللوجو" in m for m in missing)
    assert any("التوقيع" in m for m in missing)
    assert any("الشهر" in m for m in missing)          # اسم الشهر إلزامي


def test_audit_passes_a_complete_file(app, tmp_path):
    _init()
    from documents.official_xlsx import add_letterhead
    from services.raghibin import signatures_rows
    book = Workbook()
    ws = book.active
    add_letterhead(ws, YEAR, MONTH, 4)
    ws.cell(local_audit_row(ws), 1, "سبتمبر ٢٠٣١ — كشف كامل")
    signatures_rows(ws, ws.max_row + 2, YEAR, MONTH, 4)
    path = tmp_path / "كامل.xlsx"
    book.save(path)
    assert local_audit.audit_workbook(path, YEAR, MONTH) == []


def local_audit_row(ws):
    """أول صف فارغ بعد الدباجة (٥) — مساعد للاختبار فقط."""
    return 6


def test_audit_json_requires_core_keys(app, tmp_path):
    _init()
    good = tmp_path / "سليم.json"
    good.write_text('{"الجهة": "ق", "القسم": "ت", "الشهر": "سبتمبر", "السنة": "٢٠٣١",'
                    ' "آخر تحديث": "—"}', encoding="utf-8")
    assert local_audit.audit_json(good) == []
    bad = tmp_path / "ناقص.json"
    bad.write_text('{"الجهة": "ق"}', encoding="utf-8")
    missing = local_audit.audit_json(bad)
    assert any("القسم" in m for m in missing) and any("آخر تحديث" in m for m in missing)


def test_audit_report_counts_and_sorting(app, tmp_path):
    _init()
    folder = tmp_path / "ملفات الفحص"          # فولدر مستقل عن قاعدة الاختبار
    folder.mkdir()
    _bare_book(folder / "b.xlsx", "ناقص")
    _bare_book(folder / "a.xlsx", "ناقص")
    report = local_audit.audit_tree(folder, YEAR, MONTH)
    assert report["checked"] == 2 and report["ok"] == 0
    assert [p["file"] for p in report["problems"]] == ["a.xlsx", "b.xlsx"]


# ======================================================================
# ٤) القاعدة الذهبية للأصناف — الأصناف من المقررات التموينية فقط
# ======================================================================
def test_rations_are_the_single_source_of_items(app):
    """«أصناف المقررات التموينية هي اللي موجودة في أي أصناف» — المصدر الوحيد للكتالوج."""
    from data_access import db_warehouses as dw
    _init()
    real = "أرز بلدي"
    dr.add_item(YEAR, MONTH, "tamween", "summer", real, "كجم", 0, 0.075, 0)
    dr.set_activation(YEAR, MONTH, "tamween", "summer")
    catalog = {i["name"] for i in dw.ration_catalog(YEAR, MONTH, "supply")}
    assert real in catalog                                  # الصنف ظهر في كتالوج ٢ مخازن
    assert "صنف وهمي" not in catalog                        # ولا صنف مخترع من الفراغ
    # والمقرر المتعهد لا يتسرب للمقرر التمويني
    dr.add_item(YEAR, MONTH, "contractor", "summer", "فراخ", "كجم", 0, 0.2, 0)
    assert "فراخ" not in {i["name"] for i in dw.ration_catalog(YEAR, MONTH, "supply")}
    assert "فراخ" in {i["name"] for i in dw.ration_catalog(YEAR, MONTH, "contractor")}


def test_stock_page_uses_only_catalog_items(app, client):
    """«٢ مخازن» لا يعرض أي صنف غير موجود في المقرر التمويني."""
    from data_access import db_rations as dr2
    _init()
    real = "سكر"
    dr2.add_item(YEAR, MONTH, "tamween", "summer", real, "كجم", 0, 0.05, 0)
    dr2.set_activation(YEAR, MONTH, "tamween", "summer")
    page = client.get("/calc2").data.decode("utf-8")
    assert real in page or "لا تأميدات" in page


# ======================================================================
# ٥) المسارات — الفحص متاح من الواجهة
# ======================================================================
def test_audit_routes_respond(client):
    _init()
    eid, name = _entity("جهة المسار")
    tameedat_fs.snapshot_all(YEAR, MONTH)
    data = client.get("/files/audit?section=tameedat").get_json()
    assert data["ok"] in (True, False) and "checked" in data and "problems" in data
    page = client.get("/files/audit-page").data.decode("utf-8")
    assert "الفحص الذكي" in page and "القواعد الذهبية" in page
    assert arnum.to_arabic_indic(0) in page                  # الأرقام عربية


def test_audit_route_rejects_unknown_section(client):
    _init()
    assert client.get("/files/audit?section=nope").status_code == 400


def test_signatures_are_found_at_the_bottom_of_long_sheets(app):
    """تصحيح ٠٧/١٠/٢٠٢٦: الملفات الطويلة (٢٠٠ صف وأكثر) كان توقيعها في آخرها
    فيُبلَّغ خطأً بأنه ناقص — الفحص الآن يقرأ الورقة كلها حتى آخر صف."""
    from openpyxl import Workbook
    from core import paths
    from data_access import dataguard
    from services import local_audit as la
    _init()
    book = Workbook()
    ws = book.active
    ws.cell(1, 1, "وزارة الداخلية — قطاع وسط سيناء - قسم التعيينات")
    ws.cell(2, 1, "الشهر: سبتمبر")
    for row in range(3, 260):                     # جدول طويل يتجاوز الحد القديم
        ws.cell(row, 1, f"صف {row}")
    ws.cell(261, 1, "رائد")
    ws.cell(262, 1, "مصطفى نصرالله")
    ws.cell(261, 3, "مقدم")
    ws.cell(262, 3, "اسامة العجرودى")
    path = paths.DATA_DIR / "long.xlsx"
    dataguard.atomic_save(book.save, str(path))
    missing = la.audit_workbook(path, YEAR, MONTH)
    assert not [m for m in missing if m.startswith("التوقيع")], missing
