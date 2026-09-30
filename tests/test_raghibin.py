# -*- coding: utf-8 -*-
"""قسم الراغبين (ج١): قوة الجهات + الحالات + نسخ القوة + ملفات الكوادر الرسمية."""
from openpyxl import load_workbook

from services import raghibin as rfs
from services.raghibin import files_monthly as mfs
from services.raghibin import files_rosters as xr

from data_access import db_raghibin as drg
from data_access import db_tameedat as dt
from services import raghibin as rfs

YEAR, MONTH = 2031, 9


def _entity(Y=YEAR, M=MONTH, name="قطاع الشهيد اشرف جاد"):
    eid = dt.add_entity(Y, M, name, "شرطية")
    return eid, name


def test_add_person_and_no_duplicates(app):
    eid, _ = _entity()
    pid, created = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    assert created
    pid2, created2 = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    assert not created2 and pid2 == pid
    assert len(drg.list_persons(YEAR, MONTH, entity_id=eid)) == 1
    persons = drg.list_persons(YEAR, MONTH, entity_id=eid, category="officers")
    assert persons[0]["rank"] == "رائد"
    # الفئات منفصلة: نفس الاسم في الأفراد ليس تكرارًا
    _, created_i = drg.add_person(YEAR, MONTH, eid, "individuals", "محمد محمود")
    assert created_i


def test_exclude_toggle_and_warning(app):
    eid, _ = _entity()
    pid, _ = drg.add_person(YEAR, MONTH, eid, "officers", "ابراهيم ناجى عطا الله", "رائد")
    drg.set_excluded(YEAR, MONTH, pid, True, "هلاكات — غير راغب")
    counts = drg.entity_force_counts(YEAR, MONTH, eid)
    assert counts[eid]["officers"] == 1
    assert counts[eid]["officers_excluded"] == 1
    assert drg.list_persons(YEAR, MONTH, entity_id=eid, excluded=False) == []
    drg.set_excluded(YEAR, MONTH, pid, False)
    assert drg.list_persons(YEAR, MONTH, entity_id=eid, excluded=False)


def test_daily_one_meal_per_day(app):
    eid, _ = _entity()
    pid, _ = drg.add_person(YEAR, MONTH, eid, "officers", "مصطفى عبدالحميد", "نقيب")
    drg.set_daily(YEAR, MONTH, pid, 22, True, source="manual")
    drg.set_daily(YEAR, MONTH, pid, 22, True, source="tamida")   # upsert لا يكرر
    drg.set_daily(YEAR, MONTH, pid, 23, True)
    assert drg.month_meals(YEAR, MONTH)[pid] == 2                # وجبة لكل يوم
    assert drg.day_state(YEAR, MONTH, 22, eid) == {pid: True}


def test_copy_force_across_months(app):
    eid, name = _entity()
    pid, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    drg.set_excluded(YEAR, MONTH, pid, True, "هلاكات")
    drg.set_daily(YEAR, MONTH, pid, 22, True)
    copied_entities, copied_names = drg.copy_force(YEAR, MONTH, 2032, 5)
    assert (copied_entities, copied_names) == (1, 1)
    targets = drg.list_persons(2032, 5)
    assert len(targets) == 1
    assert targets[0]["excluded"] == 1 and targets[0]["exclude_note"] == "هلاكات"
    assert drg.month_meals(2032, 5) == {}          # التسجيلات اليومية لا تُنسخ
    ents = dt.list_entities(2032, 5)
    assert any(e["name"] == name for e in ents)    # الجهة أنشئت بقاموس الشهر الهدف


def test_cadres_files_official_and_states(app):
    eid, name = _entity()
    p1, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    p2, _ = drg.add_person(YEAR, MONTH, eid, "officers", "ابراهيم ناجى عطا الله", "رائد")
    drg.set_excluded(YEAR, MONTH, p2, True, "هلاكات")
    drg.add_person(YEAR, MONTH, eid, "individuals", "أحمد محمد صابر", "فرد (1)")
    folder = rfs.write_cadres_files(YEAR, MONTH, eid, name)
    off_path = folder / f"{name} (ضباط).xlsx"
    ind_path = folder / f"{name} (أفراد).xlsx"
    assert off_path.exists() and ind_path.exists()
    ws = load_workbook(off_path).active
    assert ws.title == "الضباط"
    assert ws.cell(6, 1).value.startswith("الجهات والكوادر المعتمدة")
    names = [ws.cell(r, 3).value for r in range(8, 10)]
    assert "محمد محمود" in names and "ابراهيم ناجى عطا الله" in names
    states = [ws.cell(r, 4).value for r in range(8, 10)]
    assert "✓ راغب" in states and "⊘ غير راغب" in states
    # الدباجة الرسمية (صفوف ١-٤ مدموجة تمتد لأول ٤ خلايا) + صف الإجمالي بعد البيانات
    merged_starts = {str(rng).split(":")[0] for rng in ws.merged_cells.ranges}
    assert "A1" in merged_starts and "A4" in merged_starts
    total_row = 8 + 2   # صفان بيانات + صف الإجمالي في الصف العاشر
    assert "الإجمالي" in str(ws.cell(total_row, 1).value)
    ws_i = load_workbook(ind_path).active
    assert ws_i.title == "الأفراد والصف"
    assert ws_i.cell(8, 3).value == "أحمد محمد صابر"


def test_daily_save_flow_clear_and_set(app, client):
    """حفظ المربعات = كشف اليوم الكامل: المرسل راغب، والغايب من المربعات يُلغى."""
    import json as _json
    from services.raghibin.files_daily import day_dir
    eid, name = _entity()
    p1, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    p2, _ = drg.add_person(YEAR, MONTH, eid, "officers", "ابراهيم ناجى عطا الله", "رائد")
    drg.set_excluded(YEAR, MONTH, p2, True, "هلاكات")
    p3, _ = drg.add_person(YEAR, MONTH, eid, "individuals", "أحمد محمد صابر", "فرد (1)")
    r = client.post("/raghibin/daily/save", data={
        "e": eid, "d": 22, "cat": "officers", "names": ["محمد محمود"]})
    assert r.status_code == 302
    assert drg.day_state(YEAR, MONTH, 22, eid) == {p1: True}
    path = day_dir(YEAR, MONTH, 22) / f"{name}.xlsx"
    assert path.exists()
    ws = load_workbook(path).active
    assert ws.title == "يوم ٢٢"
    assert "الراغبين في وجبة الطعام" in str(ws.cell(6, 1).value)
    marks = {ws.cell(r, 3).value: (ws.cell(r, 5).value, ws.cell(r, 4).value)
             for r in range(8, 11)}
    assert marks["محمد محمود"][0] == "✓ راغب بالوجبة"
    assert marks["ابراهيم ناجى عطا الله"][0] == "⊘ غير راغب"
    assert marks["أحمد محمد صابر"][0] == "✗ لا"
    assert "الإجمالي: ١ راغبون من أصل ٣" in str(ws.cell(11, 1).value)
    # حفظ جديد بدون محمد محمود → تُلغى رغبته (clear-and-set) والمستثنى يتجاهل دايمًا
    r2 = client.post("/raghibin/daily/save", data={
        "e": eid, "d": 22, "cat": "officers",
        "names": ["ابراهيم ناجى عطا الله", "مصطفى عبدالحميد"], "new_names": "[]"})
    assert not drg.day_state(YEAR, MONTH, 22, eid).get(p1)   # محذوف من المربعات = أُلغيت رغبته
    assert "warn=" in r2.headers["Location"]      # المستثنى تحذير بدون تسجيل
    # فرد مستقل عن ضباط الفئة الأخرى
    assert drg.day_state(YEAR, MONTH, 22, eid).get(p3) is None


def test_daily_save_adds_confirmed_new_person(app, client):
    """الاسم غير الموجود + تأكيد «إضافة كقوة دائمة» → يدخل قوة الكوادر ويسجل راغبًا."""
    eid, name = _entity()
    r = client.post("/raghibin/daily/save", data={
        "e": eid, "d": 22, "cat": "officers", "names": ["شريف جمال حاتم جديد"],
        "new_names": '["شريف جمال حاتم جديد"]'})
    assert r.status_code == 302
    person = drg.find_person(YEAR, MONTH, eid, "officers", "شريف جمال حاتم جديد")
    assert person is not None                    # دخل القوة الدائمة
    assert drg.day_state(YEAR, MONTH, 22, eid) == {person["id"]: True}
    cadres_file = rfs.cadres_dir(YEAR, MONTH) / f"{name} (ضباط).xlsx"
    ws = load_workbook(cadres_file).active
    assert "شريف جمال حاتم جديد" in [ws.cell(r, 3).value for r in range(8, 20)]
    # بدون تأكيد → تجاهل بتحذير وبدون إضافة
    r2 = client.post("/raghibin/daily/save", data={
        "e": eid, "d": 23, "cat": "officers", "names": ["اسم لم يؤكد"], "new_names": "[]"})
    assert "warn=" in r2.headers["Location"]
    assert drg.find_person(YEAR, MONTH, eid, "officers", "اسم لم يؤكد") is None


def test_daily_save_back_to_tameedat(app, client):
    """حفظ من مودال التأميدات يرجع لصفحة التأميدات نفسها (ربط بالاتجاهين)."""
    eid, name = _entity()
    drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    r = client.post("/raghibin/daily/save", data={
        "e": eid, "d": 22, "cat": "officers", "names": ["محمد محمود"],
        "back": "tameedat", "source": "tamida"})
    assert r.status_code == 302
    loc = r.headers["Location"]
    assert "tameedat" in loc and "tab=day" in loc
    assert "day=" in loc and "ok=" in loc


def test_panel_fragment(app, client):
    """جسم المودال المرجعي: الجهة المختارة + قوة بتشيكات + مربعات بعدد التأميدة."""
    eid, name = _entity()
    pid, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    drg.set_daily(YEAR, MONTH, pid, 22, True)
    dt.add_record(YEAR, MONTH, 22, {"id": eid, "name": name, "entity_type": "شرطية"},
                  10, 10, 10, notes="تأميدة اعتيادية")
    r = client.get(f"/raghibin/panel?en={name}&d=22&c=officers&sid=tok")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert body.count('name="names"') == 10       # ١٠ مربعات حسب التأميدة
    assert 'value="tameedat"' in body and 'value="tamida"' in body
    assert "الجهة المختارة الحالية للتعبين" in body
    assert "جهة شرطية معتمدة" in body
    assert "الضباط المسجلون" in body and "الأفراد والصفة" in body
    assert 'class="rg-force-toggle"' in body       # قائمة القوة بتشيكات
    assert "قائمة قوة الضباط المعتمدة" in body
    assert "قسمة الأسماء في كشوفات التجهيز اليومية" in body
    assert "حفظ وأعتماد التجهيزات" in body
    assert "ضبط اليوم" in body
    assert 'checked>' in body                      # محمد محمود مشيك (راغب اليوم)
    # بدون تأميدة: لا مربعات مع رسالة إرشادية
    r2 = client.get(f"/raghibin/panel?en={name}&d=5&c=officers")
    body2 = r2.get_data(as_text=True)
    assert 'name="names"' not in body2
    assert "سجّل تأميدة" in body2


def test_panel_quick_add_and_delete(app, client):
    """إضافة عضو سريع للقوة من المودال + حذف 🗑 — الجزء يتحدث فورًا."""
    eid, name = _entity()
    r = client.post("/raghibin/panel/quick_add", data={
        "e": eid, "qa_cat": "officers", "qa_rank": "عقيد",
        "qa_name": "طارق منير عبد اللطيف", "qa_note": "ملاحظة تجربة", "d": "22"})
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "طارق منير عبد اللطيف" in body          # ظهر في قائمة القوة فورًا
    person = drg.find_person(YEAR, MONTH, eid, "officers", "طارق منير عبد اللطيف")
    assert person is not None and person["rank"] == "عقيد"
    assert person["exclude_note"] == "ملاحظة تجربة"
    assert drg.get_person(YEAR, MONTH, person["id"]) is not None
    # حذف من المودال
    r2 = client.post(f"/raghibin/panel/delete/{person['id']}", data={"e": eid, "d": "22"})
    assert r2.status_code == 200
    assert "طارق منير عبد اللطيف" not in r2.get_data(as_text=True)
    assert drg.get_person(YEAR, MONTH, person["id"]) is None


def test_exclude_clears_daily_marks(app):
    eid, _ = _entity()
    pid, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    drg.set_daily(YEAR, MONTH, pid, 22, True)
    drg.set_excluded(YEAR, MONTH, pid, True, "هلاكات")
    assert drg.month_meals(YEAR, MONTH) == {}     # الاستثناء يسقط تسجيلات اليومية
    drg.set_excluded(YEAR, MONTH, pid, False)
    assert drg.day_state(YEAR, MONTH, 22, eid) == {}   # والإرجاع لا يستعيدها


def test_daily_page_calendar_folders_and_mismatch(app, client):
    from services import raghibin as rfs
    eid, name = _entity()
    p1, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    p2, _ = drg.add_person(YEAR, MONTH, eid, "officers", "مصطفى عبدالحميد", "نقيب")
    drg.set_daily(YEAR, MONTH, p1, 22, True)
    drg.set_daily(YEAR, MONTH, p2, 22, True)
    dt.add_record(YEAR, MONTH, 22, {"id": eid, "name": name, "entity_type": "شرطية"},
                  10, 10, 10, notes="تأميدة اعتيادية")
    r = client.get(f"/raghibin?tab=daily&e={eid}&d=22")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "عدم تطابق يوم ٢٢" in body                     # ١٠ تأميدة مقابل ٢ راغبين
    assert "الضباط ٢/١٠" in body and "الأفراد ٠/١٠" in body   # عداد المربعات
    root = rfs.tab_dir(YEAR, MONTH, "daily")
    assert (root / "يوم ١").exists() and (root / "يوم ٣٠").exists()   # فولدرات الشهر كاملة


def test_rank_seniority_ordering(app):
    eid, _ = _entity()
    for rank, name in (("ملازم", "ع"),
                       ("رائد", "ب"), ("لواء", "أ"), ("نقيب", "ج")):
        drg.add_person(YEAR, MONTH, eid, "officers", name, rank)
    drg.add_person(YEAR, MONTH, eid, "individuals", "م1", "عريف")
    drg.add_person(YEAR, MONTH, eid, "individuals", "م2", "مساعد أول")
    drg.add_person(YEAR, MONTH, eid, "individuals", "م3", "رتبة غير معروفة")
    persons = drg.list_persons(YEAR, MONTH, entity_id=eid)
    officers = [p["full_name"] for p in persons if p["category"] == "officers"]
    individuals = [p["full_name"] for p in persons if p["category"] == "individuals"]
    # الضباط: اللواء ← رائد ← نقيب ← ملازم (الأقدمية)
    assert officers == ["أ", "ب", "ج", "ع"]
    # الأفراد: مساعد أول ← عريف ← وغير المعروفة آخر الفئة
    assert individuals == ["م2", "م1", "م3"]
    assert drg.rank_weight("officers", "لواء") < drg.rank_weight("officers", "رائد")
    assert drg.rank_weight("officers", "ملازم") == 9      # آخر رتب معروفة
    assert drg.rank_weight("officers", "غريب") == 11      # غير المعروفة آخر الفئة


def test_excluded_tab_and_files(app, client):
    eid, name = _entity()
    p1, _ = drg.add_person(YEAR, MONTH, eid, "officers", "ابراهيم ناجى عطا الله", "رائد")
    p2, _ = drg.add_person(YEAR, MONTH, eid, "individuals", "كريم فتحي عوض", "فرد (1)")
    drg.set_excluded(YEAR, MONTH, p2, True, "مأمورية خارج القطاع")
    folder = xr.write_excluded_files(YEAR, MONTH, eid, name)
    off = folder / f"{name} (ضباط غير راغبين).xlsx"
    ind = folder / f"{name} (أفراد غير راغبين).xlsx"
    assert off.exists() and ind.exists()
    ws = load_workbook(ind).active
    assert ws.cell(8, 3).value == "كريم فتحي عوض"
    assert ws.cell(8, 4).value == "مأمورية خارج القطاع"
    assert "الإجمالي: ١" in str(ws.cell(9, 1).value)
    r = client.get(f"/raghibin?tab=excluded&e={eid}")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "سجل عدم الراغبين" in body and "كريم فتحي عوض" in body
    assert "إرجاعه للقوة الراغبة" in body


def test_monthly_tab_and_file(app, client):
    eid, name = _entity()
    p1, _ = drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    p2, _ = drg.add_person(YEAR, MONTH, eid, "officers", "مصطفى عبدالحميد", "نقيب")
    p3, _ = drg.add_person(YEAR, MONTH, eid, "officers", "ابراهيم ناجى عطا الله", "رائد")
    drg.set_excluded(YEAR, MONTH, p3, True, "هلاكات")
    drg.set_daily(YEAR, MONTH, p1, 22, True)
    drg.set_daily(YEAR, MONTH, p1, 23, True)
    drg.set_daily(YEAR, MONTH, p2, 22, True)
    path = mfs.write_monthly_file(YEAR, MONTH, eid, name)
    book = load_workbook(path)
    assert book.sheetnames == ["الضباط", "الأفراد والصف"]
    ws = book["الضباط"]
    assert ws.cell(8, 2).value == "رائد / محمد محمود" and ws.cell(8, 3).value == "٢"
    assert ws.cell(9, 3).value == "١"                       # نقيب / مصطفى — وجبة
    names = [ws.cell(r, 2).value for r in range(8, 11)]     # المستثنى خارج الكشف
    assert all("ابراهيم" not in str(n) for n in names)
    assert "إجمالي ٢ فرد | إجمالي الوجبات: ٣ وجبة" in str(ws.cell(10, 1).value)
    r = client.get(f"/raghibin?tab=monthly&e={eid}&c=officers")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    assert "رائد / محمد محمود" in body
    assert "إجمالي الوجبات: ٣ وجبة" in body
    assert "الضباط" in body and "الأفراد" in body


def test_monthly_rebuilds_on_save(app, client):
    eid, name = _entity()
    drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    client.post("/raghibin/daily/save", data={
        "e": eid, "d": 5, "cat": "officers", "names": ["محمد محمود"]})
    ws = load_workbook(mfs.tab_dir(YEAR, MONTH, "monthly") / f"{name}.xlsx")["الضباط"]
    assert ws.cell(8, 3).value == "١"
    assert "إجمالي الوجبات: ١ وجبة" in str(ws.cell(9, 1).value)


def test_rag_check_range_warning_excess_only(app, client):
    """شيك «عدم مطابقة الأسماء وتجاهل النقص»: زيادة الراغبين تنبّه والنقص لا."""
    from core import egtime
    eid, name = _entity()
    drg.add_person(YEAR, MONTH, eid, "officers", "محمد محمود", "رائد")
    drg.add_person(YEAR, MONTH, eid, "officers", "مصطفى عبدالحميد", "نقيب")
    client.post("/raghibin/daily/save", data={
        "e": eid, "d": 22, "cat": "officers",
        "names": ["محمد محمود", "مصطفى عبدالحميد"]})
    # تأميدة ضابط واحد فقط والراغبين اتنين → زيادة → تنبيه عدم مطابقة
    r = client.post("/tameedat/records/add", data={
        "save_token": "tok-excess", "entity_name": name, "day_from": "22",
        "day_to": "22", "officers": "1", "individuals": "0", "recruits": "0",
        "notes": "", "rag_check": "1", "custom_rations_json": ""})
    assert "warn=" in r.headers["Location"]
    assert "%D8%B9%D8%AF%D9%85%20%D9%85%D8%B7%D8%A7%D8%A8%D9%82%D8%A9" in r.headers["Location"]
    # تأميدة ١٠ ضباط والراغبين ٢ → نقص → لا تنبيه (تجاهل النقص)
    r2 = client.post("/tameedat/records/add", data={
        "save_token": "tok-short", "entity_name": name, "day_from": "25",
        "day_to": "25", "officers": "10", "individuals": "10", "recruits": "0",
        "notes": "", "rag_check": "1", "custom_rations_json": ""})
    loc2 = r2.headers["Location"]
    assert "warn=" not in loc2


def test_page_and_section_redirect(client):
    r = client.get("/raghibin?tab=cadres")
    assert r.status_code == 200
    body = r.get_data(as_text=True)
    for name in ("الراغبين (يومي)", "عدم الراغبين", "تجميع الكشف العام",
                 "الجهات والكوادر المعتمدة"):
        assert name in body
    r2 = client.get("/sections/raghebeen")
    assert r2.status_code == 302 and "/raghibin" in r2.headers["Location"]
