import io
import zipfile
from pathlib import Path
import pytest
from openpyxl import load_workbook
from PIL import Image
from data_access import db_letterhead as lh, storage, db_entities, db_recruits
from documents import xlsx_rations, letterhead_docx, docx_recruits
from documents.official_xlsx import EXPORT_VERSION
from services.images import save_logo


def image_bytes(fmt="PNG"):
    stream = io.BytesIO()
    Image.new("RGB", (320, 180), (20, 40, 60)).save(stream, format=fmt)
    stream.seek(0)
    return stream


def official_setup():
    name = save_logo(image_bytes(), storage.letterhead_dir(2031, 9))
    lh.save(2031, 9, {"lh_1": "جهة اختبار معزولة", "lh_2": "السطر الثاني", "lh_3": "السطر الثالث",
                      "lh_4": "السطر الرابع", "logo_file": name,
                      "sig_right_rank": "رائد", "sig_right_name": "توقيع اليمين",
                      "sig_left_rank": "عميد", "sig_left_name": "توقيع الشمال"})


@pytest.mark.parametrize("section", ["tamween", "contractor"])
def test_embedded_logo_and_letterhead_in_every_sheet(app, section):
    official_setup()
    if section == "contractor":
        db_entities.add_entity(2031, 9, "contractor", "جهة اختبار")
    path = xlsx_rations.rebuild(2031, 9, section)
    wb = load_workbook(path)
    try:
        assert len(wb.worksheets) == (4 if section == "contractor" else 3)
        for ws in wb:
            assert ws["A1"].value == "جهة اختبار معزولة"
            assert ws["A4"].value == "السطر الرابع"
            assert len(ws._images) == 1
            assert ws._images[0].width <= 116
            assert ws.sheet_view.rightToLeft
            assert ws.freeze_panes == "A8"
            assert ws.print_title_rows == "$1:$7"
            assert ws.page_setup.fitToWidth == 1
            values = {c.value for row in ws for c in row if c.value}
            assert {"توقيع اليمين", "توقيع الشمال"} <= values
        with zipfile.ZipFile(path) as archive:
            assert not archive.testzip()
            assert len([p for p in archive.namelist() if p.startswith("xl/media/")]) == len(wb.worksheets)
    finally:
        wb.close()


def test_old_workbooks_and_letterhead_changes_rebuild_on_download(app):
    path = xlsx_rations.rebuild(2031, 9, "tamween")
    wb = load_workbook(path)
    wb.properties.version = "1"
    wb.save(path)  # Isolated test fixture only, simulating a pre-upgrade file.
    wb.close()
    official_setup()
    result = load_workbook(xlsx_rations.ensure(2031, 9, "tamween"))
    assert result.properties.version == EXPORT_VERSION
    assert result.worksheets[0]["A1"].value == "جهة اختبار معزولة"
    result.close()
    lh.save(2031, 9, {"lh_1": "تعديل لاحق"})
    result = load_workbook(xlsx_rations.ensure(2031, 9, "tamween"))
    assert result.worksheets[0]["A1"].value == "تعديل لاحق"
    result.close()


def test_logo_optional_empty_cells_not_none(app):
    wb = load_workbook(xlsx_rations.rebuild(2031, 10, "tamween"))
    assert not wb.worksheets[0]._images
    assert wb.worksheets[0]["A1"].value is None
    assert "None" not in str([c.value for row in wb.worksheets[0] for c in row if c.value])
    wb.close()


def test_docx_use_same_month_and_logo(app):
    official_setup()
    db_recruits.add_recruit(2031, 9, {"name": "اختبار مؤقت", "mil_no": "1"})
    paths = [letterhead_docx.rebuild(2031, 9), *docx_recruits.rebuild_month(2031, 9)]
    for path in paths:
        assert "09-" in str(path)
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode()
            assert "جهة اختبار معزولة" in xml and "توقيع اليمين" in xml
            assert any(name.startswith("word/media/") for name in archive.namelist())
    path = letterhead_docx.rebuild(2031, 10)
    with zipfile.ZipFile(path) as archive:
        assert "جهة اختبار معزولة" not in archive.read("word/document.xml").decode()


def test_invalid_logo_does_not_replace_saved_settings(client):
    official_setup()
    before = lh.get_all(2031, 9)
    response = client.post("/letterhead/save?year=2031&month=9", data={
        "lh_1": "must not save", "logo": (io.BytesIO(b"not an image"), "fake.png")},
        content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    assert lh.get_all(2031, 9) == before


def test_upload_webp_and_remove_logo_updates_all_xlsx(client):
    response = client.post("/letterhead/save?year=2031&month=9", data={
        "lh_1": "شهر سبتمبر", "logo": (image_bytes("WEBP"), "logo.webp")},
        content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    assert lh.logo_path(2031, 9).suffix == ".png"
    client.post("/letterhead/logo/delete?year=2031&month=9")
    for section in ("tamween", "contractor"):
        wb = load_workbook(xlsx_rations.ensure(2031, 9, section))
        assert all(not ws._images for ws in wb)
        wb.close()
