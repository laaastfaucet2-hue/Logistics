# -*- coding: utf-8 -*-
"""الألوان الثابتة للأصناف والجهات والأشخاص (توجيه ٠٦/١٠/٢٠٢٦ — أمر عسكري).

يُحصَّن: لكل صنف لون مختلف · اللون ثابت لا يتغير · المدى الغذائي للأصناف لا يلمس
ألوان الجهات ولا ألوان الأشخاص · التسجيل الجديد ياخد أول لون غير مستخدم ·
اللون يظهر في الشاشات وفي ملفات الإكسل.
"""
from core import colors
from data_access import db_rations as dr, db_tameedat as dt, months
from openpyxl import load_workbook

YEAR, MONTH = 2031, 9


def test_every_item_gets_a_different_color(app):
    names = ["أرز", "زيت", "سكر", "مكرونة", "صلصة", "شاي", "ملح"]
    assigned = [colors.color_for("item", n) for n in names]
    assert len(set(assigned)) == len(names)                  # لا صنف يشبه صنفًا
    assert all(value.startswith("#") for value in assigned)


def test_color_is_stable_across_calls(app):
    first = colors.color_for("entity", "قطاع وسط سيناء")
    assert colors.color_for("entity", "قطاع وسط سيناء") == first
    assert colors.color_for("entity", "  قطاع وسط سيناء  ") == first   # تجاهل المسافات


def test_kinds_do_not_share_palettes(app):
    item = colors.color_for("item", "أرز")
    entity = colors.color_for("entity", "قطاع وسط سيناء")
    person = colors.color_for("person", "مصطفى السيد")
    assert item not in colors.ENTITY_COLORS
    assert item not in colors.PERSON_COLORS
    assert entity not in colors.ITEM_COLORS
    assert person not in colors.ITEM_COLORS and person not in colors.ENTITY_COLORS


def test_ring_style_for_officer_name(app):
    style = colors.ring_style("person", "مصطفى السيد")
    assert "--ring:#" in style and "border-color:#" in style


def test_new_registration_takes_unused_color(app):
    first = colors.color_for("item", "صنف أول")
    second = colors.color_for("item", "صنف ثانٍ")
    assert first != second
    store = colors._load()
    assert store["item"]["صنف أول"] == first and store["item"]["صنف ثانٍ"] == second


def test_colors_appear_in_rations_page(client):
    months.init_month(YEAR, MONTH)
    dr.add_item(YEAR, MONTH, "tamween", "summer", "عدس", "كجم", 0, 0.1, 0)
    page = client.get("/rations/tamween?year=2031&month=9").data.decode("utf-8")
    assert colors.color_for("item", "عدس") in page             # لون الصنف ظاهر في الشاشة


def test_entity_color_in_tameedat_sheet(client):
    months.init_month(YEAR, MONTH)
    eid = dt.add_entity(YEAR, MONTH, "جهة الألوان", "شرطية")
    dt.add_record(YEAR, MONTH, 1, dt.get_entity(YEAR, MONTH, eid), 3, 5, 0, "", [])
    from services import tameedat_fs
    tameedat_fs.snapshot_all(YEAR, MONTH)
    path = (tameedat_fs.tab_dir(YEAR, MONTH, "dict") / tameedat_fs.TAB_XLSX["dict"])
    sheet = load_workbook(path)["قاموس الجهات"]
    header_row = next(r for r in range(1, 12)
                      if sheet.cell(r, 2).value == "الجهة")
    cell = next(sheet.cell(r, 2) for r in range(header_row + 1, sheet.max_row + 1)
                if sheet.cell(r, 2).value == "جهة الألوان")
    assert cell.font.color.rgb.lower().endswith(
        colors.color_for("entity", "جهة الألوان").lstrip("#").lower())
