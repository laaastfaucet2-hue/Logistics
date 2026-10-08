# -*- coding: utf-8 -*-
"""تاب «٢ مخازن — التغليف» (توجيه ٠٨/١٠): النافذة المتراصة وملف الإذن المنفصل.

التاب المنفصل «٢ مخازن تفاريد» اتلغى — بيانات التغليف بقت في سب-تاب «التغليف»
جوه تاب ٢ مخازن: سطر لكل إذن بإجمالي منصرف قابل للضغط يفتح النافذة المتراصة
(الأصناف فوق بعض، كل صنف ٣ سطور: الصنف/بالوحدة/بالتغليف + سطر باتش صغير:
تغليف الدفعة·مخزن·انتهاء)، وكل إذن في إكسل منفصل:
٢ مخازن\أذونات الصرف\يوم N\إذن M.xlsx + ٢ مخازن\تفاريد\يوم N\إذن M.xlsx."""
from data_access import db_permits as dp, db_rations as dr
from data_access import db_warehouses as dw, months
from services import warehouses_fs as wf

YEAR, MONTH = 2031, 9


def _init():
    months.init_month(YEAR, MONTH)


def _seed_stock_and_permit(number=3, qty=60):
    dr.add_item(YEAR, MONTH, "tamween", "summer", "أرز", "كجم", 0, 0.1, 0)
    dr.set_activation(YEAR, MONTH, "tamween", "summer")
    dw.add_opener(YEAR, MONTH, "supply", "أرز", 200, 1, handle_unit_hint="كجم",
                  exp_iso="2027-06-30", pack_kind="شكارة", pack_count=4, pack_capacity=50)
    dp.save_permit(YEAR, MONTH, {
        "number": number, "fiscal_year": 2026, "date_from": 1, "date_to": 1,
        "issue_days": 1, "mode": "box", "entity_label": "جهة التفريدة",
        "officers": 5, "individuals": 10, "recruits": 0,
        "actuals": {"tamween_أرز": qty}})


def _header_row(ws):
    """سطر العناوين تحت الدباجة (٦ صفوف) — أول سطر فيه «الصنف»."""
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v not in (None, "")]
        if "الصنف" in vals:
            return vals
    return []


def test_tabs_registered_with_new_folders(app):
    _init()
    from routes.warehouses import TABS
    keys = [t[0] for t in TABS]
    assert "tafreeda" not in keys                      # التاب المنفصل اتلغى
    assert ("wh2", "٢ مخازن", "📤") in TABS
    assert ("wh3", "٣ مخازن", "📒") in TABS
    assert wf.SUB_FOLDERS["wh2"] == "٢ مخازن"
    assert wf.SUB_FOLDERS["wh3"] == "٣ مخازن"
    assert wf.WH2_ISSUES == "أذونات الصرف" and wf.WH2_TAFARID == "تفاريد"
    assert wf.file_path(YEAR, MONTH, "supply", "wh2") is None   # كل إذن في ملفه
    assert wf.TAB_XLSX_2["wh3"] == "٣ مخازن تفاريد.xlsx"


def test_packaging_subtab_opens_stack_window(client):
    _init()
    _seed_stock_and_permit()
    page = client.get("/warehouses?cycle=supply&sub=wh2").data.decode("utf-8")
    # تابان فرعيان: «أذونات الصرف» (الافتراضي) + «التغليف»
    assert 'data-wh2sub="sijlat"' in page and 'data-wh2sub="tafared"' in page
    assert "أذونات الصرف" in page and "التغليف" in page
    # إجمالي المنصرف قابل للضغط يفتح النافذة المتراصة
    assert "wh-total-btn" in page and 'data-taf-open="packDialog-3"' in page
    # النافذة: كل صنف ٣ سطور + سطر باتش صغير + إجماليات
    import re as _re
    _dlg = _re.search(r'id="packDialog-3".*?</dialog>', page, _re.S).group(0)
    assert "المنصرف بالوحدة" in _dlg and "المنصرف بالتغليف" in _dlg
    assert "تغليف الدفعة" in _dlg and "المخزن" in _dlg and "الانتهاء" in _dlg
    assert "إجمالي المنصرف" in _dlg
    assert "جهة التفريدة" in page
    assert "٦٠" in _dlg                                # المنصرف بالوحدة (أرقام عربية)


def test_packaging_subtab_empty_state_is_explicit(client):
    _init()
    page = client.get("/warehouses?cycle=contractor&sub=wh2").data.decode("utf-8")
    assert 'data-wh2panel="tafared"' in page
    assert "لا منصرف مسجّل" in page                    # حالة فارغة صريحة
    assert "sub=tafreeda" not in page                  # مفيش وداد للتبويب الملغي


def test_permit_files_exist_with_same_columns_as_the_view(client):
    """قاعدة «الورق = الإكسل»: ملف «تفاريد\يوم N\إذن M.xlsx» بنفس بيانات الشاشة."""
    from openpyxl import load_workbook
    _init()
    _seed_stock_and_permit(number=9, qty=30)
    wf.snapshot_cycle(YEAR, MONTH, "supply")
    taf_path = wf.wh2_permit_dir(YEAR, MONTH, "supply", taf=True) / "يوم ١" / "إذن ٩.xlsx"
    issues_path = wf.wh2_permit_dir(YEAR, MONTH, "supply", taf=False) / "يوم ١" / "إذن ٩.xlsx"
    assert taf_path.exists()
    assert issues_path.exists()
    wb = load_workbook(taf_path)
    assert "التفريدة" in wb.sheetnames
    headers = _header_row(wb["التفريدة"])
    for name in ("الصنف", "الكمية بالوحدة", "المنصرف بالتغليف",
                 "تغليف الدفعة", "المخزن", "تاريخ الانتهاء", "الصلاحية المتبقية"):
        assert name in headers, name
    # هيدر الإذن: كل حاجة بتعريفه
    head = wb.worksheets[0]
    head_text = " ".join(str(c) for row in head.iter_rows(values_only=True)
                         for c in row if c is not None)
    for token in ("الجهة المستلمة", "المسؤول عن الصرف", "من يوم", "إلى يوم", "أيام الصرف"):
        assert token in head_text, token


def test_tab_bar_counts_and_no_legacy_tab(client):
    _init()
    _seed_stock_and_permit(number=11, qty=25)
    page = client.get("/warehouses?cycle=supply&sub=wh2").data.decode("utf-8")
    assert "sub=tafreeda" not in page                  # مفيش تاب منفصل في الشريط
    assert "wh2" in page and "٢ مخازن" in page
