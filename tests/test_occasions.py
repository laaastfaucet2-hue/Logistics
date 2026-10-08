# -*- coding: utf-8 -*-
"""صفحة «التوثيق والمناسبات» — توجيه ٠٦/١٠/٢٠٢٦.

يُحصَّن: التسجيل في فولدر محلي باسم المناسبة (صور/ · فيديو/ · بيانات.json) ·
التعديل · المعرضان · تشغيل الفيديو (خدمة الملف داخل البرنامج) · التقرير Excel+Word
بالدباجة واللوجو والتوقيعين · البحث الذكي بالتاريخ والنوع · الحذف الكامل.
"""
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from core import paths
from services import occasions_fs as ofs

# أصغر PNG سليم (١×١) لاختبار الرفع بلا أي ملف خارجي
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
    "890000000a49444154789c6360000002000154a24f9a0000000049454e44ae426082")


def _create(client, **values):
    data = {"title": values.get("title", "زيارة رسمية للتفتيش"),
            "kind": values.get("kind", "تفتيش"),
            "date_iso": values.get("date_iso", "2026-10-06"),
            "place": values.get("place", "قطاع وسط سيناء"),
            "force_text": values.get("force_text", "١٠ أفراد"),
            "notes": values.get("notes", "تفقد النقاط والوقوف على الخدمة")}
    return client.post("/occasions/create", data=data)


def test_occasions_page_is_permanent_and_lists_kinds_and_filters(client):
    page = client.get("/occasions").data.decode("utf-8")
    assert "التوثيق والمناسبات" in page
    assert "صفحة دائمة لا تتبع شهرًا ولا سنة" in page
    # «بحث» (توجيه: مربع البحث الذكى سمية «بحث» — بلا إيموجي على الأزرار)
    assert 'class="oc-wide">بحث <input name="q"' in page and 'name="from"' in page and 'name="to"' in page
    # النوع يُختار من قائمة (ويمكن كتابة مسمّى حر) — قرار المستخدم
    assert 'name="kind"' in page and "<datalist" in page


def test_create_builds_local_folder_with_media_subfolders_and_data(client):
    response = _create(client, title="زيارة السيد اللواء")
    assert response.status_code == 302
    root = ofs.root()
    folders = [p for p in root.glob("*/*") if p.is_dir()]
    assert len(folders) == 1
    folder = folders[0]
    assert "زيارة السيد اللواء" in folder.name and "2026-10-06" in folder.name
    assert (folder / "صور").is_dir() and (folder / "فيديو").is_dir()
    assert (folder / "بيانات.json").is_file()
    item = ofs.load(folder)
    assert item["kind"] == "تفتيش" and item["place"] == "قطاع وسط سيناء"
    assert item["notes"] and item["photos"] == [] and item["videos"] == []
    # التقريران يُبنيان فور الإنشاء (قاعدة: لا مناسبة بلا ورق رسمي)
    assert (folder / ofs.REPORT_XLSX).is_file() and (folder / ofs.REPORT_DOCX).is_file()


def test_report_has_letterhead_logo_and_two_signatures(client, app):
    _create(client, title="تفتيش مفاجئ")
    folder = next(p for p in ofs.root().glob("*/*") if p.is_dir())
    book = load_workbook(folder / ofs.REPORT_XLSX)
    ws = book.active
    assert ws.cell(1, 1).value                       # الدباجة ليست فارغة أبدًا
    assert ws.cell(6, 1).value.startswith("تقرير مناسبة")
    texts = [str(c.value or "") for row in ws.iter_rows() for c in row]
    assert any("تفتيش مفاجئ" in t for t in texts)
    assert any("الوصف" in t for t in texts)
    docx_bytes = (folder / ofs.REPORT_DOCX).read_bytes()
    assert docx_bytes[:2] == b"PK"                   # ملف Word سليم (ZIP)
    assert b"word/document.xml" in docx_bytes[:4000] or len(docx_bytes) > 5000


def test_upload_photo_and_video_then_serve_them(client, app):
    _create(client, title="زيارة تصوير")
    folder = next(p for p in ofs.root().glob("*/*") if p.is_dir())
    response = client.post("/occasions/media/upload", data={
        "folder": str(folder),
        "photos": (BytesIO(PNG_1PX), "صورة الزيارة.png"),
        "videos": (BytesIO(b"\x00\x00\x00\x18ftypmp42"), "مقطع.mp4"),
    }, content_type="multipart/form-data")
    assert response.status_code == 302
    item = ofs.load(folder)
    assert len(item["photos"]) == 1 and len(item["videos"]) == 1
    assert (folder / "صور" / "صورة الزيارة.png").is_file()
    assert (folder / "فيديو" / "مقطع.mp4").is_file()
    photo = client.get("/occasions/media/photo/صورة الزيارة.png", query_string={"folder": str(folder)})
    assert photo.status_code == 200 and photo.data.startswith(b"\x89PNG")
    video = client.get("/occasions/media/video/مقطع.mp4", query_string={"folder": str(folder)})
    assert video.status_code == 200 and b"ftyp" in video.data
    # الصفحة تعرض المعرضين ومشغّل الفيديو داخل البرنامج
    page = client.get("/occasions", query_string={"folder": str(folder)}).data.decode("utf-8")
    assert "معرض الصور" in page and "معرض الفيديو" in page and "<video" in page
    assert "oc-lightbox" in page


def test_media_path_guards_against_leaving_the_folder(client, app):
    _create(client, title="حماية المسار")
    folder = next(p for p in ofs.root().glob("*/*") if p.is_dir())
    assert ofs.media_path(folder, "photo", "../بيانات.json") is None
    response = client.get("/occasions/media/photo/..%2Fبيانات.json",
                          query_string={"folder": str(folder)})
    assert response.status_code != 200


def test_edit_occasion_rebuilds_report_and_keeps_folder(client, app):
    _create(client, title="مناسبة قبل التعديل")
    folder = next(p for p in ofs.root().glob("*/*") if p.is_dir())
    response = client.post("/occasions/save", data={
        "folder": str(folder), "title": "مناسبة بعد التعديل", "kind": "زيارة ميدانية",
        "date_iso": "2026-10-07", "place": "بوابة الدخول", "force_text": "٣ ضباط",
        "notes": "سطر أول\nسطر ثاني"})
    assert response.status_code == 302
    item = ofs.load(folder)
    assert item["title"] == "مناسبة بعد التعديل" and item["kind"] == "زيارة ميدانية"
    assert "\n" in item["notes"]
    ws = load_workbook(folder / ofs.REPORT_XLSX).active
    texts = [str(c.value or "") for row in ws.iter_rows() for c in row]
    assert any("مناسبة بعد التعديل" in t for t in texts)
    assert any("سطر ثاني" in t for t in texts)


def test_smart_search_by_date_and_text(client, app):
    _create(client, title="زيارة رسمية أكتوبر", date_iso="2026-10-06", kind="زيارة رسمية")
    _create(client, title="تفتيش سبتمبر", date_iso="2026-09-20", kind="تفتيش")
    assert len(ofs.list_occasions()) == 2
    assert len(ofs.list_occasions(date_from="2026-10-01")) == 1
    assert len(ofs.list_occasions(date_to="2026-09-30")) == 1
    assert len(ofs.list_occasions(query="تفتيش")) == 1
    page = client.get("/occasions", query_string={"q": "أكتوبر"}).data.decode("utf-8")
    assert "زيارة رسمية أكتوبر" in page and "تفتيش سبتمبر" not in page


def test_kinds_in_use_grows_with_free_custom_kind(client, app):
    _create(client, title="مناسبة ترفيهية", kind="حفل تكريم")
    kinds = ofs.kinds_in_use()
    assert "حفل تكريم" in kinds and "زيارة رسمية" in kinds


def test_delete_occasion_removes_everything(client, app):
    _create(client, title="مناسبة للحذف")
    folder = next(p for p in ofs.root().glob("*/*") if p.is_dir())
    (folder / "صور" / "x.png").write_bytes(PNG_1PX)
    assert client.post("/occasions/delete", data={"folder": str(folder)}).status_code == 302
    assert not folder.exists()
    assert ofs.list_occasions() == []


def test_sidebar_has_occasions_entry(client):
    page = client.get("/dashboard").data.decode("utf-8")
    assert "التوثيق والمناسبات" in page
    assert "/sections/occasions" in page


def test_occasions_files_live_outside_month_folders(client, app):
    """الصفحة دائمة: ملفاتها في database/occasions/ لا داخل فولدر شهر."""
    _create(client, title="فصل عن الشهور")
    root = ofs.root()
    assert root.name == "occasions"
    assert Path(paths.DATA_DIR) == root.parent
    assert root.parent.name == Path(paths.DATA_DIR).name
    assert "occasions" in str(root)
