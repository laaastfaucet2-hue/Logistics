"""Date entry/display/report regression tests, always under the isolated app fixture."""
import io
import zipfile
from datetime import date, datetime, timezone
from html.parser import HTMLParser
import pytest
from docx import Document
from core import dates, egtime
from data_access import db_recruits as recruits, db_attendance as attendance, storage, dataguard
from documents import docx_recruits
from routes.backups import _time_h


def row():
    return {'name':'سجل تاريخ معزول','mil_no':'90031', 'service_start':'2024-03-04',
            'service_end':'2033-12-31','has_cert':1,'cert_date':'2031-02-01','cert_expiry':'2032-04-03'}


def test_legacy_iso_records_are_displayed_day_first_without_rewriting(client):
    rid = recruits.add_recruit(2031,9,row())
    original = recruits.get_recruit(2031,9,rid)
    page = client.get(f'/recruits/?year=2031&month=9&edit={rid}')
    assert page.status_code == 200
    for text in ('٠٤/٠٣/٢٠٢٤','٣١/١٢/٢٠٣٣','٠١/٠٢/٢٠٣١','٠٣/٠٤/٢٠٣٢'):
        assert text in page.text
    assert 'value="٠٤/٠٣/٢٠٢٤"' in page.text
    assert 'type="date"' not in page.text
    assert recruits.get_recruit(2031,9,rid) == original
    for url in ['/recruits/print/certs','/recruits/print/recruit?rid='+str(rid)]:
        report = client.get(url)
        assert report.status_code == 200
        assert '٢٠٣٣-١٢-٣١' not in report.text
        assert '٠٤/٠٣/٢٠٢٤' in report.text or '٠١/٠٢/٢٠٣١' in report.text


def test_new_arabic_day_first_entry_stores_iso(client):
    data = {**row(), 'service_start':'03/04/2025', 'service_end':'31122033',
            'cert_date':'٠١/٠٢/٢٠٣١','cert_expiry':'۰۳/۰۴/۲۰۳۲'}
    response = client.post('/recruits/save?year=2031&month=9',data=data)
    assert response.status_code == 302
    saved = recruits.list_recruits(2031,9)[0]
    assert saved['service_start'] == '2025-04-03'  # April third, not March fourth.
    assert saved['service_end'] == '2033-12-31'
    assert saved['cert_date'] == '2031-02-01'
    assert saved['cert_expiry'] == '2032-04-03'
    assert recruits.days_to_discharge(saved,date(2033,12,30)) == 1
    assert recruits.cert_state(saved,date(2032,4,2)) == 'soon'


@pytest.mark.parametrize('values', [
    {'service_start':'31/02/2031'}, {'cert_expiry':'12/31/2031'},
    {'service_start':'05/09/2031','service_end':'04/09/2031'},
    {'cert_date':'05/09/2031','cert_expiry':'04/09/2031'},
])
def test_invalid_dates_do_not_write_records_or_uploads(client, values):
    response=client.post('/recruits/save?year=2031&month=9',data={
        **row(), **values, 'photo':(io.BytesIO(b'not written'), 'photo.png')}, content_type='multipart/form-data')
    assert response.status_code == 400
    assert 'role="alert"' in response.text and 'سجل تاريخ معزول' in response.text
    assert recruits.list_recruits(2031,9) == []
    assert not list(storage.recruits_dir(2031,9).rglob('photo_*'))


def test_rejected_edit_preserves_existing_dates_and_draft(client):
    rid=recruits.add_recruit(2031,9,row())
    original=recruits.get_recruit(2031,9,rid)
    response=client.post('/recruits/save?year=2031&month=9',data={
        **row(),'id':rid,'name':'تعديل لم يحفظ','service_end':'31/02/2034'})
    assert response.status_code == 400
    assert 'تعديل لم يحفظ' in response.text and 'value="31/02/2034"' in response.text
    assert recruits.get_recruit(2031,9,rid) == original


def test_unchecked_certificate_ignores_stale_invalid_dates(client):
    data={**row(), 'has_cert':'', 'cert_date':'invalid', 'cert_expiry':'invalid'}
    response=client.post('/recruits/save?year=2031&month=9',data=data)
    assert response.status_code == 302
    saved=recruits.list_recruits(2031,9)[0]
    assert saved['has_cert']==0 and saved['cert_date']==saved['cert_expiry']==''


def test_daily_and_leave_reports_use_complete_dates(client, monkeypatch):
    monkeypatch.setattr(egtime,'today',lambda:date(2031,9,22))
    rid=recruits.add_recruit(2031,9,row())
    attendance.set_range(2031,9,rid,'إجازة',3,5,'')
    permit=client.get(f'/recruits/print/permit?year=2031&month=9&rid={rid}&from=3&to=5')
    assert permit.status_code==200
    for text in ('٠٣/٠٩/٢٠٣١','٠٥/٠٩/٢٠٣١','٢٢/٠٩/٢٠٣١'):
        assert text in permit.text
    assert '٠٣/٠٩/٢٠٣١' in client.get('/recruits/?year=2031&month=9&tab=leaves').text
    assert '٠٣/٠٩/٢٠٣١' in client.get('/recruits/print/day?year=2031&month=9&day=3').text
    assert client.get(f'/recruits/print/permit?year=2031&month=9&rid={rid}&from=3&to=31').status_code==404
    _, leaves=docx_recruits.rebuild_month(2031,9)
    path=docx_recruits.build_leave_permit(row(),3,5,2031,9)
    for document in (leaves,path):
        with zipfile.ZipFile(document) as archive:
            xml=archive.read('word/document.xml').decode('utf-8')
            assert '٠٣/٠٩/٢٠٣١' in xml and '٠٥/٠٩/٢٠٣١' in xml
        assert Document(document).core_properties.version == docx_recruits.DATE_EXPORT_VERSION


def test_stale_monthly_word_report_is_refreshed_on_download(client):
    rid=recruits.add_recruit(2031,9,row())
    attendance.set_range(2031,9,rid,'إجازة',3,5,'')
    _, leaves=docx_recruits.rebuild_month(2031,9)
    stale=Document(leaves)
    stale.core_properties.version='previous-date-format'
    stale.save(leaves)  # Only the isolated test fixture, simulating an old cached report.
    response=client.get(f'/recruits/docx/{leaves.name}?year=2031&month=9')
    assert response.status_code==200
    assert Document(io.BytesIO(response.data)).core_properties.version==docx_recruits.DATE_EXPORT_VERSION
    assert recruits.get_recruit(2031,9,rid)['service_start']=='2024-03-04'


def test_backup_ui_uses_cairo_day_first_but_archive_ids_stay_unchanged(client):
    epoch=datetime(2026,9,21,22,15,tzinfo=timezone.utc).timestamp()
    assert _time_h(epoch)=='٢٢/٠٩/٢٠٢٦ — ٠١:١٥:٠٠'
    result=dataguard.create_backup('manual')
    name=result['name']
    response=client.get('/backups/')
    assert response.status_code==200 and 'التاريخ والوقت' in response.text
    assert name in response.text  # Kept in the download/restore URL, not renamed on disk.
    assert dataguard.valid_name(name) and result['path'].exists()
    assert 'class="date-display" dir="ltr"' in response.text


def test_default_month_uses_cairo_date_not_server_localtime(app, monkeypatch):
    from data_access import database
    from core.auth_core import current_context
    conn = database.get_conn()
    conn.execute("DELETE FROM user_context WHERE user_id=1")
    conn.commit(); conn.close()
    monkeypatch.setattr(egtime, 'today', lambda: date(2031, 1, 1))
    with app.test_request_context('/dashboard'):
        year, month = current_context(1)
    assert year == 2031 and month == 1
