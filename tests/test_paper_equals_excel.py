# -*- coding: utf-8 -*-
"""«الورق = الإكسل» — كل كشف مطبوع بنفس أعمدة ملفه المحلي بالحرف (توجيه ٠٦/١٠/٢٠٢٦).

قاعدة المستخدم: «الورق اللي في البرنامج هو بالضبط اللي موجود في الإكسل».
الاختبار يقارن **أعمدة ملف الإكسل الفعلي** (من openpyxl) بـ**أعمدة الجدول في الشاشة/الطباعة**
(من HTML الحقيقي للصفحة) ويؤكد التطابق اسمًا وترتيبًا — فلا يقع فرق صامت مرة أخرى.
"""
import re
from openpyxl import load_workbook

from data_access import db_tameedat as dt
from services import sheet_columns as sc, tameedat_fs

YEAR, MONTH = 2031, 9


def _header_row(path, first_header):
    """صف الأعمدة في ملف الإكسل: أول صف فيه اسم العمود الأول (م/اليوم)."""
    ws = load_workbook(path).active
    for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
        if row and str(row[0] or "").strip() == first_header:
            return [str(v or "").strip() for v in row if str(v or "").strip()]
    raise AssertionError(f"لم أجد صف الأعمدة في {path.name}")


def _table_headers(html, table_id=None):
    """أعمدة أول جدول في الصفحة (أو الجدول ذي المعرّف المحدد) — بلا أعمدة الأزرار."""
    scope = html
    if table_id:
        scope = html[html.index(f'id="{table_id}"'):]
        scope = scope[:scope.index("</table>")]
    else:
        scope = html[html.index("<table"):]
        scope = scope[:scope.index("</table>")]
    head = scope[:scope.index("</thead>")] if "</thead>" in scope else scope
    cols = [re.sub(r"<[^>]+>", "", th).strip()
            for th in re.findall(r"<th[^>]*>(.*?)</th>", head, re.S)]
    return [c for c in cols if c and c != "إجراءات"]


def _seed_month():
    eid = dt.add_entity(YEAR, MONTH, "قطاع أعمدة", "شرطية", 12, 40)
    entity = dt.get_entity(YEAR, MONTH, eid)
    dt.add_record(YEAR, MONTH, 3, entity, 10, 25, 8, "", [])
    tameedat_fs.snapshot_all(YEAR, MONTH)
    return entity


# ======================================================================
# ١) قاموس الجهات: الشاشة = الملف
# ======================================================================
def test_dict_tab_columns_match_excel(client):
    entity = _seed_month()
    page = client.get("/tameedat?tab=dict").get_data(as_text=True)
    path = tameedat_fs.tab_dir(YEAR, MONTH, "dict") / tameedat_fs.TAB_XLSX["dict"]
    assert _table_headers(page, "tmDictTable") == _header_row(path, "م")
    assert _table_headers(page, "tmDictTable") == sc.TAMEEDAT_DICT
    assert entity["name"] in page


# ======================================================================
# ٢) الجهات المومدة: الشاشة = الملف = التقرير المطبوع
# ======================================================================
def test_momoda_tab_columns_match_excel(client):
    _seed_month()
    page = client.get("/tameedat?tab=momoda").get_data(as_text=True)
    path = tameedat_fs.tab_dir(YEAR, MONTH, "momoda") / tameedat_fs.TAB_XLSX["momoda"]
    book = load_workbook(path)
    sheet = book["ملخص الجهات المومدة"]
    header = None
    for row in sheet.iter_rows(min_row=1, max_row=20, values_only=True):
        cells = [str(c or "").strip() for c in row if str(c or "").strip()]
        if cells[:1] == ["م"]:
            header = cells
            break
    assert header == sc.TAMEEDAT_MOMODA
    assert _table_headers(page) == sc.TAMEEDAT_MOMODA


# ======================================================================
# ٣) التقرير الشامل المطبوع = ملف التقرير
# ======================================================================
def test_printed_report_columns_match_excel(client):
    _seed_month()
    html = client.get("/tameedat/report/print",
                      follow_redirects=True).get_data(as_text=True)
    heads = re.findall(r"<thead>(.*?)</thead>", html, re.S)
    printed = [[re.sub(r"<[^>]+>", "", th).strip()
                for th in re.findall(r"<th[^>]*>(.*?)</th>", block, re.S)] for block in heads]
    printed = [cols for cols in printed if cols]
    assert printed[0] == sc.TAMEEDAT_DAY          # التفصيل اليومي
    assert printed[1] == sc.TAMEEDAT_MOMODA       # الملخص الشهري
    from documents import xlsx_tameedat
    report = xlsx_tameedat.rebuild(YEAR, MONTH)

    def _cols(sheet, first):
        for row in sheet.iter_rows(min_row=1, max_row=20, values_only=True):
            cells = [str(c or "").strip() for c in row if str(c or "").strip()]
            if cells[:1] == [first]:
                return cells
        raise AssertionError(f"لا صف أعمدة في {sheet.title}")

    book = load_workbook(report)
    assert _cols(book["التفصيل اليومي"], "اليوم") == sc.TAMEEDAT_DAY
    assert _cols(book["الملخص الشهري"], "م") == sc.TAMEEDAT_MOMODA


# ======================================================================
# ٤) الراغبين: ملف اليوم = نفس أعمدة الشاشة المطبوعة
# ======================================================================
def test_raghibin_day_file_columns_are_shared_constant():
    from services.raghibin import files_daily
    assert files_daily.HEADERS[:3] == sc.RAGHIBIN_DAY[:3]
    from services.raghibin import files_rosters
    assert files_rosters.HEADERS == sc.RAGHIBIN_ROSTER


# ======================================================================
# ٥) ٢ مخازن تفاريد: التبويب = الملف (حرس رجوع)
# ======================================================================
def test_tafreeda_tab_columns_match_excel(client):
    """جولة ٢٣: تاب التفاريد جوه صفحة ٢ مخازن — وكل إذن في إكسل منفصل،
    والنافذة المتراصة (بالوحدة/بالتغليف) = شيتي ملف الإذن حرفيًا."""
    from services import warehouses_fs
    from data_access import db_stores, db_tameedat as dt
    from data_access import db_rations as dr, months
    months.init_month(YEAR, MONTH)
    dr.add_item(YEAR, MONTH, "tamween", "summer", "أرز بلدي", "طن", 0, 0, 0)
    db_stores.add_store("مخزن التفريدة")
    store = db_stores.list_stores()[0]
    assert client.post("/warehouses/wh1/add?cycle=supply", data={
        "cycle": "supply", "item_name": "أرز بلدي", "qty": "٥٠٠", "day": "٣",
        "pack_kind": "شكارة", "pack_count": "١٠", "pack_capacity": "٥٠",
        "producer": "مطاحن الاختبار",
        "store_id": [str(store["id"])], "store_qty": ["٥٠٠"]}).status_code == 302
    ent_id = dt.add_entity(YEAR, MONTH, "جهة التفاريد", "شرطية")
    rec_id = dt.add_record(YEAR, MONTH, 5, dt.get_entity(YEAR, MONTH, ent_id), 1, 5, 0, "")
    assert client.post("/calc2/save", data={
        "date_from": "5", "date_to": "5", "issue_days": "1", "number": "1",
        "selected_json": '["main:%d"]' % rec_id, "entity_label": "جهة التفاريد",
        "meal_lunch": "1", "actual_tamween_أرز بلدي": "60"}).status_code == 302
    # الرابط القديم بيريدريكت لصفحة ٢ مخازن جوا تاب التفاريد
    page = client.get("/warehouses?cycle=supply&sub=tafreeda",
                      follow_redirects=True).get_data(as_text=True)
    assert 'data-wh2sub="tafared"' in page
    assert page.count("<th") > 0
    assert "المنصرف بالوحدة" in page and "المنصرف بالتغليف" in page
    warehouses_fs.snapshot_cycle(YEAR, MONTH, "supply")
    path = warehouses_fs.wh2_permit_file_path(YEAR, MONTH, "supply", 1, 5, taf=True)
    assert path.is_file(), "ملف التفاريد لم يُبنَ"

    def _cols(sheet_title, first):
        ws = load_workbook(path)[sheet_title]
        for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
            vals = [str(v or "").strip() for v in row]
            if vals and vals[0] == first:
                return [v for v in vals if v]
        raise AssertionError(f"لا صف أعمدة في {sheet_title}")

    # الورق = الإكسل: نفس الأعمدة الحرفية جوه شيتي الإذن
    assert _cols("الأصناف", "م") == ["م", "الصنف", "المنصرف بالوحدة", "الوحدة",
                                     "المنصرف بالتغليف"]
    assert _cols("التفريدة", "م") == ["م", "الصنف", "الكمية بالوحدة", "المنصرف بالتغليف",
                                      "تغليف الدفعة", "المخزن", "تاريخ الانتهاء",
                                      "الصلاحية المتبقية", "ملاحظات"]
    # نفس الكمية حرفيًا: قيمة شيت «الأصناف» = سطر «المنصرف بالوحدة» في النافذة المتراصة
    window = re.search(r'id="packDialog-1".*?</dialog>', page, re.S).group(0)
    item_row = next(r for r in load_workbook(path)["الأصناف"].iter_rows(min_row=2, values_only=True)
                    if str(r[1] or "").strip() == "أرز بلدي")
    qty = str(item_row[2] or "").strip()
    assert qty and qty in window
