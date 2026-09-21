# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""مولّد ملف Word للدباجة والتوقيعات الرسمية — يُعاد بناؤه بعد كل حفظ.

يحفظ في: database/الدباجة/الدباجة والتوقيعات.docx
الشكل: أسطر الدباجة يمين أعلى الصفحة + اللوجو شمال أعلى الصفحة
       + توقيع (رتبة/اسم) يمين أسفل الصفحة وآخر شمال أسفلها.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import database as db
import storage

NAVY = RGBColor(0x0A, 0x12, 0x30)
GOLD = RGBColor(0xB8, 0x86, 0x0B)


def docx_path():
    return storage.letterhead_dir() / "الدباجة والتوقيعات.docx"


def _line(doc, text, size, bold=True, color=NAVY):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Cairo"
    return p


def _sig_cell(cell, rank, name):
    """كتلة التوقيع: قيمة الرتبة سطر، وقيمة الاسم سطر تحته — بدون تسميات."""
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for value in (rank, name):
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value or "........................")
        run.font.bold = True
        run.font.size = Pt(13)
        run.font.name = "Cairo"


def rebuild():
    """يعيد بناء ملف Word من الإعدادات الحالية — يُستدعى بعد كل حفظ."""
    doc = Document()
    # ترويسة: اللوجو شمال (فقرة يسار) + أسطر الدباجة يمين
    logo = db.get_setting("logo_file")
    logo_path = storage.letterhead_dir() / logo if logo else None
    if logo_path and logo_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(logo_path), width=Cm(3))
    for i, size in ((1, 16), (2, 14), (3, 13), (4, 13)):
        text = db.get_setting(f"lh_{i}")
        if text:
            _line(doc, text, size, color=(GOLD if i == 1 else NAVY))
    doc.add_paragraph("\n\n\n")

    # توقيعان: يمين وشمال أسفل الصفحة
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _sig_cell(table.cell(0, 0), db.get_setting("sig_left_rank"),
              db.get_setting("sig_left_name"))
    _sig_cell(table.cell(0, 1), db.get_setting("sig_right_rank"),
              db.get_setting("sig_right_name"))

    path = docx_path()
    doc.save(str(path))
    return path


def ensure():
    """يتأكد أن الملف موجود (يبنيه لو اختفى) ويرجع مساره."""
    path = docx_path()
    if not path.exists():
        rebuild()
    return path
