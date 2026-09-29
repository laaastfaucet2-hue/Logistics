# -*- coding: utf-8 -*-
"""معاينة ملفات Excel المحلية — تقلب على database/2026/09-سبتمبر وتعرض كل الشيتات HTML.

التشغيل:  .venv/bin/python scripts/excel_gallery.py
الناتج:   /home/user/معاينة-ملفات-الاكسل.html
"""
import html
import os
import sys

BASE = "/home/user/Logistics/database/2026/09-سبتمبر"
OUT = "/home/user/معاينة-ملفات-الاكسل.html"

from openpyxl import load_workbook


def render_sheet(ws):
    rows = []
    for r in ws.iter_rows(values_only=True):
        if all(v is None for v in r):
            continue
        rows.append(r)
    if not rows:
        return "<p class='empty'>شيت فاضي</p>"
    ncols = max(len(r) for r in rows)
    out = ["<table>"]
    for i, r in enumerate(rows):
        cells = []
        for c in range(ncols):
            v = r[c] if c < len(r) else None
            cells.append("<td>%s</td>" % html.escape(str(v)) if v is not None else "<td></td>")
        cls = " class='head'" if i == 0 else ""
        out.append("<tr%s>%s</tr>" % (cls, "".join(cells)))
    out.append("</table>")
    return "\n".join(out)


def collect():
    """يرجع [(مسار نسبي, المسار الكامل)] مرتب."""
    files = []
    for root, dirs, names in os.walk(BASE):
        dirs.sort()
        for n in sorted(names):
            if n.endswith(".xlsx") and not n.startswith("~$"):
                p = os.path.join(root, n)
                files.append((os.path.relpath(p, BASE), p))
    return files


def main():
    files = collect()
    blocks = []
    for rel, path in files:
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as exc:
            blocks.append("<details><summary>📄 %s <b class='bad'>(تالف: %s)</b></summary></details>"
                          % (html.escape(rel), html.escape(str(exc))))
            continue
        sheets = []
        for ws in wb.worksheets:
            sheets.append("<h4>ورقة: %s</h4>%s" % (html.escape(ws.title), render_sheet(ws)))
        blocks.append(
            "<details open><summary>📄 %s <span class='meta'>(%d ورقة)</span></summary>%s</details>"
            % (html.escape(rel), len(wb.worksheets), "\n".join(sheets)))
    head = """<!doctype html><html dir="rtl" lang="ar"><head><meta charset="utf-8">
<title>معاينة ملفات Excel — ٩/٢٠٢٦</title><style>
body{font-family:'Segoe UI',Tahoma,sans-serif;background:#eef3fa;color:#1a2a3a;margin:20px}
h1{color:#123c73;text-align:center}
details{background:#fff;border:1px solid #c5d5ea;border-radius:8px;margin:10px 0;padding:8px 14px}
summary{cursor:pointer;font-weight:bold;color:#123c73;font-size:15px}
table{border-collapse:collapse;margin:8px 0;min-width:60pc}
td{border:1px solid #b9c9dd;padding:4px 10px;font-size:13px}
tr.head td{background:#2f6db5;color:#fff;font-weight:bold}
.empty{color:#888}.meta{color:#777;font-weight:normal;font-size:12px}
.bad{color:#c0392b}
h4{color:#2f6db5;margin:12px 0 2px}
</style></head><body>
<h1>📊 معاينة ملفات Excel المحلية — شهر ٩ / ٢٠٢٦</h1>
<p style="text-align:center;color:#555">كل الملفات المولّدة في 06-مستودعات وسجلات · 07-التاميدات · 08-المجندين · 10-المخازن والثلاجات · 11-الترفية · 13-الصحة والدباجة</p>"""
    doc = head + "\n" + "\n".join(blocks) + "\n</body></html>"
    with open(OUT, "w", encoding="utf-8") as f:
        f.write(doc)
    print("written: %s الحجم: %d بايت — %d ملف" % (OUT, os.path.getsize(OUT), len(files)))


if __name__ == "__main__":
    sys.exit(main())
