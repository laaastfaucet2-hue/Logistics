"""Shared official XLSX letterhead: real cells + embedded logo, never a screen-only preview."""
import io
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from data_access import db_letterhead as lh
from services.images import png_bytes

TITLE_ROW = 6
TABLE_ROW = 7
DATA_ROW = 8
EXPORT_VERSION = "2.0-monthly-letterhead"


def add_letterhead(ws, year, month, ncols):
    values = lh.get_all(year, month)
    for row in range(1, 5):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols - 2)
        cell = ws.cell(row, 1, values[f"lh_{row}"] or "")
        cell.font = Font(name="Arial", size=14 if row == 1 else 12, bold=True,
                         color="996C27" if row == 1 else "132638")
        cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
        ws.row_dimensions[row].height = 25
    ws.row_dimensions[5].height = 12
    path = lh.logo_path(year, month)
    if path:
        image = Image(io.BytesIO(png_bytes(path, max_size=(116, 116))))
        ws.add_image(image, f"{get_column_letter(ncols - 1)}1")
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if ncols > 7 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35)
    ws.print_title_rows = f"1:{TABLE_ROW}"
    ws.freeze_panes = f"A{DATA_ROW}"
