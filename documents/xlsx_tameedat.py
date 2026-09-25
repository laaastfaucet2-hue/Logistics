# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""التقرير الشامل للتأميدات (Excel) — بالدباجة واللوجو والتوقيعات الرسمية.

يُحفظ محليًا في: database/<السنة>/<الشهر>/07-التاميدات/طباعة التقرير الشامل/
وشيتان: «التفصيل اليومي» (كل تأميدة بملحقاتها وإجماليها) و«الملخص الشهري»
(الجهات المومدة مستقلة + إجمالي القوة + أيام التميد) — بدون «متوسط تغذية».
الدباجة واللوجو يدخلان الملف نفسه عبر documents.official_xlsx.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import arabic_numbers as arnum, dates, egtime
from core.config import MONTH_NAMES
from data_access import dataguard
from data_access import db_letterhead as lhdb
from data_access import db_tameedat as dt
from documents.official_xlsx import add_letterhead, TITLE_ROW, TABLE_ROW, DATA_ROW, EXPORT_VERSION
from services import tameedat_fs

NAVY, GOLD, GOLD_L, WHITE = "0A1230", "B8860B", "F6D47C", "FFFFFF"
SUBTOTAL = "E9EDF7"
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
_BORDER = Border(*[Side(style="thin", color="9AA6C7")] * 4)

EXPORT_TAG = "tameedat-report-1.0"
FILE_BASE = "التقرير الشامل - تأميدات"


def report_name(year, month):
    return f"{FILE_BASE} {MONTH_NAMES[month - 1]} {year}"


def xlsx_path(year, month):
    return tameedat_fs.report_dir(year, month) / f"{report_name(year, month)}.xlsx"


def _title(ws, text, ncols):
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=ncols)
    cell = ws.cell(TITLE_ROW, 1, text)
    cell.font = Font(bold=True, size=15, color=GOLD_L, name="Cairo")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    ws.row_dimensions[TITLE_ROW].height = 34


def _header(ws, headers):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(TABLE_ROW, col, text)
        cell.font = Font(bold=True, size=12, color=WHITE, name="Cairo")
        cell.fill = PatternFill("solid", fgColor=GOLD)
        cell.alignment = CENTER
        cell.border = _BORDER
    ws.row_dimensions[TABLE_ROW].height = 26


def _cell(ws, row, col, value, bold=False, fill=None):
    cell = ws.cell(row, col, value)
    cell.alignment = CENTER
    cell.border = _BORDER
    cell.font = Font(size=11, name="Cairo", bold=bold)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    return cell


def _signatures(ws, row, ncols, year, month):
    """توقيعان أسفل الجدول: قيمة الرتبة سطر وقيمة الاسم سطر تحته — بدون تسميات."""
    blocks = [(1, lhdb.get_setting(year, month, "sig_right_rank"),
               lhdb.get_setting(year, month, "sig_right_name")),
              (ncols - 1, lhdb.get_setting(year, month, "sig_left_rank"),
               lhdb.get_setting(year, month, "sig_left_name"))]
    for col, rank, name in blocks:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        top = ws.cell(row, col, rank or "........................")
        top.font = Font(bold=True, size=12, name="Cairo")
        top.alignment = CENTER
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        bottom = ws.cell(row + 1, col, name or "........................")
        bottom.font = Font(bold=True, size=12, name="Cairo")
        bottom.alignment = CENTER


def _daily_sheet(wb, year, month, records):
    ws = wb.create_sheet("التفصيل اليومي")
    headers = ["اليوم", "التاريخ", "الجهة", "النوع", "ضباط", "أفراد", "مجندين",
               "الإجمالي", "ملاحظات"]
    add_letterhead(ws, year, month, len(headers))
    _title(ws, f"التقرير الشامل للتأميدات — التفصيل اليومي — "
               f"{MONTH_NAMES[month - 1]} {arnum.to_arabic_indic(year)}", len(headers))
    _header(ws, headers)
    row = DATA_ROW
    for rec in records:
        stamp = f"{day_count_cell(rec, month, year)}"
        if rec["range_days"] > 1:  # تأميدة واحدة سارية من يوم إلى يوم — بلا تكرار
            end_iso = "{:04d}-{:02d}-{:02d}".format(year, month, rec["day_to"])
            stamp = f"تأميدة واحدة من {stamp} إلى {dates.format_date(end_iso)}"
        values = [arnum.to_arabic_indic(rec["day"]), stamp, rec["entity_name"],
                  rec["entity_type"], arnum.to_arabic_indic(rec["officers"]),
                  arnum.to_arabic_indic(rec["individuals"]),
                  arnum.to_arabic_indic(rec["recruits"]),
                  arnum.to_arabic_indic(rec["total"]), rec["notes"] or ""]
        for col, value in enumerate(values, start=1):
            _cell(ws, row, col, value, bold=(col == 3),
                  fill=("F4F6FB" if row % 2 == 0 else None))
        row += 1
        for att in rec["attachments"]:
            att_values = ["", "", f"ملحقة: {att['name']}", att["entity_type"] or "—",
                          arnum.to_arabic_indic(att["officers"]),
                          arnum.to_arabic_indic(att["individuals"]),
                          arnum.to_arabic_indic(att["recruits"]),
                          arnum.to_arabic_indic(att["officers"] + att["individuals"]
                                                + att["recruits"]), ""]
            for col, value in enumerate(att_values, start=1):
                _cell(ws, row, col, value)
            row += 1
        if rec["attachments"]:
            merged = ["", "", f"إجمالي تأميدة «{rec['entity_name']}» مع الملحقات", "",
                      arnum.to_arabic_indic(rec["officers"] + sum(a["officers"] for a in rec["attachments"])),
                      arnum.to_arabic_indic(rec["individuals"] + sum(a["individuals"] for a in rec["attachments"])),
                      arnum.to_arabic_indic(rec["recruits"] + sum(a["recruits"] for a in rec["attachments"])),
                      arnum.to_arabic_indic(rec["grand_total"]), ""]
            for col, value in enumerate(merged, start=1):
                _cell(ws, row, col, value, bold=True, fill=SUBTOTAL)
            row += 1
    if row == DATA_ROW:
        _cell(ws, row, 1, "لا توجد تأميدات مسجلة في هذا الشهر بعد")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        row += 1
    next_row = row + 2
    _signatures(ws, next_row, len(headers), year, month)
    ws.print_area = f"A1:{get_column_letter(len(headers))}{next_row + 1}"
    for col, width in enumerate([7, 13, 30, 10, 8, 8, 9, 9, 26], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def day_count_cell(rec, month, year):
    """تاريخ اليوم كاملًا dd/mm/yyyy عربي داخل خلية الإكسل."""
    return dates.format_date(f"{year:04d}-{month:02d}-{rec['day']:02d}")


def _summary_sheet(wb, year, month, summary, totals):
    ws = wb.create_sheet("الملخص الشهري")
    headers = ["م", "الجهة المومدة", "النوع", "أيام التميد", "عدد التأميدات",
               "ضباط", "أفراد", "مجندين", "الإجمالي",
               "متوسط ضباط", "متوسط أفراد", "متوسط مجندين"]
    add_letterhead(ws, year, month, len(headers))
    _title(ws, f"الجهات المومدة بالشهر الحالي — {MONTH_NAMES[month - 1]} "
               f"{arnum.to_arabic_indic(year)}", len(headers))
    _header(ws, headers)
    row = DATA_ROW
    for index, item in enumerate(summary, start=1):
        name = item["name"] + (" (ملحقة)" if item["kind"] == "attachment" else "")
        values = [arnum.to_arabic_indic(index), name, item["entity_type"],
                  arnum.to_arabic_indic(item["active_days"]),
                  arnum.to_arabic_indic(item["records"]),
                  arnum.to_arabic_indic(item["total_officers"]),
                  arnum.to_arabic_indic(item["total_individuals"]),
                  arnum.to_arabic_indic(item["total_recruits"]),
                  arnum.to_arabic_indic(item["grand_total"]),
                  arnum.fmt_qty(round(item["avg_officers"], 3)),
                  arnum.fmt_qty(round(item["avg_individuals"], 3)),
                  arnum.fmt_qty(round(item["avg_recruits"], 3))]
        for col, value in enumerate(values, start=1):
            _cell(ws, row, col, value, bold=(col == 2),
                  fill=("E9EDF7" if item["kind"] == "attachment"
                        else "F4F6FB" if row % 2 == 0 else None))
        row += 1
    total_values = ["", "إجمالي القوة", "", arnum.to_arabic_indic(totals["active_days"]),
                    arnum.to_arabic_indic(totals["records"]),
                    arnum.to_arabic_indic(totals["officers"]),
                    arnum.to_arabic_indic(totals["individuals"]),
                    arnum.to_arabic_indic(totals["recruits"]),
                    arnum.to_arabic_indic(totals["grand"]), "", "", ""]
    for col, value in enumerate(total_values, start=1):
        _cell(ws, row, col, value, bold=True, fill=SUBTOTAL)
    row += 2
    _signatures(ws, row, len(headers), year, month)
    ws.print_area = f"A1:{get_column_letter(len(headers))}{row + 1}"
    for col, width in enumerate([6, 28, 10, 11, 12, 9, 9, 10, 10, 12, 12, 12], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def rebuild(year, month):
    """يبني ملف الإكسل كاملًا بالدباجة واللوجو ويحفظه محليًا — كتابة ذرّية."""
    tameedat_fs.ensure_folders(year, month)
    records = dt.month_records(year, month)
    summary, totals = dt.month_summary(year, month)
    wb = Workbook()
    wb.properties.version = EXPORT_VERSION
    wb.properties.keywords = EXPORT_TAG
    wb.remove(wb.active)
    _daily_sheet(wb, year, month, records)
    _summary_sheet(wb, year, month, summary, totals)
    path = xlsx_path(year, month)
    dataguard.atomic_save(wb.save, path)
    return path
