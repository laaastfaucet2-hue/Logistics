# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""التقرير الشامل للتأميدات (Word) — بالدباجة واللوجو والتوقيعات الرسمية.

يُحفظ محليًا بجوار ملف الإكسل في فولدر «طباعة التقرير الشامل» داخل التاميدات.
المحتوى: ترويسة رسمية + جدول التفصيل اليومي + جدول الملخص الشهري + توقيعان.
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from core import arabic_numbers as arnum, dates
from core.config import MONTH_NAMES
from data_access import dataguard
from data_access import db_letterhead as lhdb
from data_access import db_tameedat as dt
from data_access import storage
from documents.xlsx_tameedat import report_name
from services import tameedat_fs

NAVY = RGBColor(0x0A, 0x12, 0x30)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
FONT = "Cairo"


def docx_path(year, month):
    return tameedat_fs.report_dir(year, month) / f"{report_name(year, month)}.docx"


def _para(doc, text, size=12, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return p


def _letterhead(doc, year, month):
    logo = lhdb.logo_path(year, month)
    if logo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(logo), width=Cm(3))
    for index, size in ((1, 16), (2, 14), (3, 13), (4, 13)):
        text = lhdb.get_setting(year, month, f"lh_{index}")
        if text:
            _para(doc, text, size, color=(GOLD if index == 1 else NAVY))


def _head_cell(cell, text):
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].add_run(text)
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.name = FONT
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shading = cell._tc.get_or_add_tcPr().makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd",
        {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill": "0A1230"})
    cell._tc.get_or_add_tcPr().append(shading)


def _row(table, values, bold=False):
    row = table.add_row()
    for index, value in enumerate(values):
        cell = row.cells[index]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(str(value))
        run.font.size = Pt(10)
        run.font.bold = bold
        run.font.name = FONT


def _daily_table(doc, year, month, records):
    _para(doc, "أولًا: التفصيل اليومي للتأميدات", 13, color=NAVY,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, text in enumerate(["اليوم", "التاريخ", "الجهة", "النوع", "ضباط",
                                  "أفراد", "مجندين", "الإجمالي"]):
        _head_cell(table.rows[0].cells[index], text)
    if not records:
        _row(table, ["—"] * 3 + ["لا توجد تأميدات مسجلة هذا الشهر"] + ["—"] * 4)
        return
    for rec in records:
        _row(table, [arnum.to_arabic_indic(rec["day"]),
                     dates.document_date(f"{year:04d}-{month:02d}-{rec['day']:02d}"),
                     rec["entity_name"], rec["entity_type"],
                     arnum.to_arabic_indic(rec["officers"]),
                     arnum.to_arabic_indic(rec["individuals"]),
                     arnum.to_arabic_indic(rec["recruits"]),
                     arnum.to_arabic_indic(rec["total"])], bold=True)
        for att in rec["attachments"]:
            _row(table, ["", "", f"ملحقة: {att['name']}", att["entity_type"] or "—",
                         arnum.to_arabic_indic(att["officers"]),
                         arnum.to_arabic_indic(att["individuals"]),
                         arnum.to_arabic_indic(att["recruits"]),
                         arnum.to_arabic_indic(att["officers"] + att["individuals"]
                                               + att["recruits"])])
        if rec["attachments"]:
            _row(table, ["", "", f"إجمالي التأميدة مع الملحقات", "",
                         arnum.to_arabic_indic(
                             rec["officers"] + sum(a["officers"] for a in rec["attachments"])),
                         arnum.to_arabic_indic(
                             rec["individuals"] + sum(a["individuals"] for a in rec["attachments"])),
                         arnum.to_arabic_indic(
                             rec["recruits"] + sum(a["recruits"] for a in rec["attachments"])),
                         arnum.to_arabic_indic(rec["grand_total"])], bold=True)


def _summary_table(doc, summary, totals):
    doc.add_paragraph()
    _para(doc, "ثانيًا: الملخص الشهري — الجهات المومدة", 13, color=NAVY,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, text in enumerate(["م", "الجهة", "النوع", "أيام التميد", "ضباط",
                                  "أفراد", "مجندين", "الإجمالي"]):
        _head_cell(table.rows[0].cells[index], text)
    for index, item in enumerate(summary, start=1):
        _row(table, [arnum.to_arabic_indic(index), item["name"], item["entity_type"],
                     arnum.to_arabic_indic(item["active_days"]),
                     arnum.to_arabic_indic(item["total_officers"]),
                     arnum.to_arabic_indic(item["total_individuals"]),
                     arnum.to_arabic_indic(item["total_recruits"]),
                     arnum.to_arabic_indic(item["grand_total"])])
    _row(table, ["", "إجمالي القوة", "", arnum.to_arabic_indic(totals["active_days"]),
                 arnum.to_arabic_indic(totals["officers"]),
                 arnum.to_arabic_indic(totals["individuals"]),
                 arnum.to_arabic_indic(totals["recruits"]),
                 arnum.to_arabic_indic(totals["grand"])], bold=True)


def _signatures(doc, year, month):
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, side in ((0, "left"), (1, "right")):
        cell = table.cell(0, column)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for key in ("rank", "name"):
            value = lhdb.get_setting(year, month, f"sig_{side}_{key}")
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value or "........................")
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.name = FONT


def rebuild(year, month):
    """يبني ملف الوورد كاملًا ويحفظه محليًا — كتابة ذرّية."""
    tameedat_fs.ensure_folders(year, month)
    records = dt.month_records(year, month)
    summary, totals = dt.month_summary(year, month)
    doc = Document()
    _letterhead(doc, year, month)
    _para(doc, report_name(year, month), 15, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
    _daily_table(doc, year, month, records)
    _summary_table(doc, summary, totals)
    _signatures(doc, year, month)
    path = docx_path(year, month)
    dataguard.atomic_save(doc.save, path)
    return path
