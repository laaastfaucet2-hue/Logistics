# -*- coding: utf-8 -*-
# ⚠️ قاعدة إلزامية: لا يزيد أي ملف عن 1000 سطر — الترتيب المعماري موثّق في CONTRIBUTING.md
"""صفحة الدباجة والتوقيعات الرسمية — تُعمَّم على كل ملفات الإكسل والوورد الرسمية.

المحتوى: مربعات أسطر الدباجة (يمين أعلى المستند) + الشعار (شمال أعلى المستند)
+ توقيعان (رتبة/اسم) يمين وشمال أسفل المستند + معاينة حية لشكل الوورد والإكسل.
"""
from data_access import dataguard
from flask import g
from core.downloads import attachment
import logging
import os
import sys
import subprocess
from urllib.parse import quote

from flask import (Blueprint, render_template, request, redirect,
                   url_for, abort, send_file)

from core.auth_core import login_required, current_session, current_context
from data_access import db_letterhead as db
from data_access import months as months_mod
from services import legacy_months
from services.images import save_logo
from services.document_refresh import refresh_month
from core import arabic_numbers as arnum
from core.config import MONTH_NAMES, month_folder
from data_access import storage
from documents import xlsx_rations
from documents import letterhead_docx

letterhead_bp = Blueprint("letterhead", __name__, url_prefix="/letterhead")

ALLOWED_LOGO = {"png", "jpg", "jpeg", "webp"}
SETTING_KEYS = ["lh_1", "lh_2", "lh_3", "lh_4",
                "sig_right_rank", "sig_right_name",
                "sig_left_rank", "sig_left_name"]
PREVIEW_EXTS = {".xlsx", ".docx"}


def _rb(ok=None, err=None):
    _, token = current_session()
    year, month = _ctx()
    return redirect(url_for("letterhead.page", year=year, month=month, sid=token, ok=ok, err=err))


def _ctx():
    return current_context(g.user["id"])


def _page_vars():
    year, month = _ctx()
    v = db.template_vars(year, month)
    v.update(year=year, month=month, month_name=MONTH_NAMES[month-1],
             legacy_letterhead=legacy_months.letterhead_available() and not legacy_months.imported("letterhead", year, month),
             folder_path=f"database/{year}/{month_folder(month)}/الدباجة/")
    v["can_open"] = os.name == "nt"
    return v


# ======================================================================
# الصفحة + المعاينة الحية (شكل عام + معاينة أي ملف محلي — مولّدة، من غير قراءة الملف)
# ======================================================================
def _month_root(year, month):
    return months_mod.month_db_path(year, month).parent


def _local_files(year, month):
    """كل الملفات المحلية للشهر (xlsx/docx) لقائمة المعاينة — أسماء فقط."""
    root = _month_root(year, month)
    files = []
    if not root.is_dir():
        return files
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.suffix.lower() in PREVIEW_EXTS \
                and not p.name.startswith("~$"):
            rel = p.relative_to(root).as_posix()
            section = rel.split("/", 1)[0] if "/" in rel else "ملفات عامة"
            files.append({"path": rel, "name": p.name, "section": section})
    return files


def _group_files(files):
    """{القسم: [ملفات...]} — الأقسام برقمها الطبيعي أولًا."""
    groups = {}
    for f in files:
        groups.setdefault(f["section"], []).append(f)
    def _key(sec):
        import re
        m = re.match(r"(\d+)", sec or "")
        return (int(m.group(1)) if m else 99, sec)
    return [(s, groups[s]) for s in sorted(groups, key=_key)]


def _safe_local_path(year, month, rel):
    """مسار آمن داخل مجلد الشهر وبامتداد مدعوم — وإلا None."""
    if not rel:
        return None
    root = _month_root(year, month).resolve()
    p = (root / rel).resolve()
    if p.suffix.lower() not in PREVIEW_EXTS:
        return None
    if p.is_file() and str(p).startswith(str(root) + os.sep):
        return p
    return None


def _sheet_names(path):
    """أسماء أوراق عمل ملف إكسل — قراءة أسماء فقط (بلا بيانات الخلايا)."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _gen_rows(seed, n=7):
    """صفوف معاينة مولّدة بأرقام عشوائية ثابتة لنفس الملف (stable per file)."""
    import random
    rnd = random.Random(seed)
    names = ["أرز بلدي", "سكر", "دقيق", "شاي", "زيت طعام", "ملح",
             "لحمة مجمدة", "تونة", "سمنة", "فاصوليا"]
    units = ["كجم", "جم", "لتر", "عبوة"]
    rows = []
    for i in range(n):
        rate = round(rnd.uniform(0.1, 9.5), 3)
        days = rnd.randint(1, 31)
        rows.append([arnum.to_arabic_indic(str(i + 1)), rnd.choice(names),
                     rnd.choice(units), arnum.fmt_qty(rate),
                     arnum.fmt_qty(days), arnum.fmt_qty(round(rate * days, 3)), ""])
    return rows


def _preview_xlsx(path, year, month, sheet_i, sheets):
    seed = abs(hash((path.name, sheets[sheet_i])))
    stamp = f"{MONTH_NAMES[int(month) - 1]} {arnum.to_arabic_indic(str(year))}"
    return {"kind": "xlsx", "rel": path.relative_to(_month_root(year, month)).as_posix(),
            "title": path.name, "sheet": sheets[sheet_i],
            "sheets": sheets, "sheet_i": sheet_i,
            "stamp": stamp, "rows": _gen_rows(seed)}


def _preview_docx(path, year, month):
    stamp = f"{MONTH_NAMES[int(month) - 1]} {arnum.to_arabic_indic(str(year))}"
    return {"kind": "docx", "rel": path.relative_to(_month_root(year, month)).as_posix(),
            "title": path.name, "sheets": [], "sheet_i": 0,
            "stamp": stamp, "rows": []}


def _preview_vars(year, month, v):
    """اختيار الملف/ورقة العمل من باراميترات الصفحة + معاينة من التوليد التجريبي الحقيقي."""
    v["local_files"] = _group_files(_local_files(year, month))
    sel = (request.args.get("file") or "").strip()
    v["preview_file"] = None
    p = _safe_local_path(year, month, sel)
    if not p:
        return v
    try:
        if p.suffix.lower() == ".xlsx":
            # أسماء الأوراق عمل من ناتج التوليد التجريبي المقابل (مش من ملف المستخدم)
            from services import letterhead_preview as lhp
            src = lhp.find_source(p.relative_to(_month_root(year, month)).as_posix())
            sheets = _sheet_names(src) if src else [p.stem]
            if not sheets:
                return v
            try:
                sheet_i = int(request.args.get("sheet") or 0)
            except ValueError:
                sheet_i = 0
            sheet_i = max(0, min(sheet_i, len(sheets) - 1))
            pv = lhp.build_preview(p.relative_to(_month_root(year, month)).as_posix(),
                                   year, month, sheet_i)
            if pv:
                v["preview_file"] = pv
            else:
                v["preview_file"] = {"kind": "xlsx", "rel": p.relative_to(_month_root(year, month)).as_posix(),
                                     "title": p.name, "sheet": sheets[sheet_i],
                                     "sheets": sheets, "sheet_i": sheet_i,
                                     "stamp": f"{MONTH_NAMES[int(month) - 1]} {arnum.to_arabic_indic(str(year))}",
                                     "rows": _gen_rows(abs(hash((p.name, sheets[sheet_i]))))}
        else:
            from services import letterhead_preview as lhp
            pv = lhp.build_preview(p.relative_to(_month_root(year, month)).as_posix(),
                                   year, month)
            if pv:
                v["preview_file"] = pv
            else:
                stamp = f"{MONTH_NAMES[int(month) - 1]} {arnum.to_arabic_indic(str(year))}"
                v["preview_file"] = {"kind": "docx", "rel": p.relative_to(_month_root(year, month)).as_posix(),
                                     "title": p.name, "sheets": [], "sheet_i": 0,
                                     "stamp": stamp, "rows": []}
    except Exception:
        logging.exception("letterhead preview render failed")
        v["preview_file"] = None
    return v


@letterhead_bp.route("")
@letterhead_bp.route("/")
@login_required
def page():
    v = _page_vars()
    year, month = _ctx()
    return render_template("letterhead.html", **_preview_vars(year, month, v))


# ======================================================================
# حفظ الإعدادات + رفع الشعار (يعيد بناء الوورد ويحدّث توقيعات كل الإكسل)
# ======================================================================
@letterhead_bp.route("/save", methods=["POST"])
@login_required
def save():
    year, month = _ctx()
    values = {key: (request.form.get(key) or "").strip() for key in SETTING_KEYS}
    file = request.files.get("logo")
    if file and file.filename:
        try:
            values["logo_file"] = save_logo(file.stream, storage.letterhead_dir(year, month))
        except ValueError as exc:
            return _rb(err=str(exc))
    db.save(year, month, values)
    try:
        refresh_month(year, month)
    except (OSError, ValueError):
        return _rb(err="حُفظت الدباجة، لكن تعذّر تحديث مستند؛ أغلق Excel أو Word ثم أعد الحفظ")
    return _rb(ok="حُفظت دباجة هذا الشهر مع الشعار وتحدّثت ملفات Excel وWord الخاصة به — والتوقيعان حُفظا على مستوى المنظومة كلها فسريا على جميع الشهور والمستندات")


# ======================================================================
# حذف الشعار
# ======================================================================
@letterhead_bp.route("/logo/delete", methods=["POST"])
@login_required
def logo_delete():
    year, month = _ctx()
    db.save(year, month, {"logo_file": ""})
    refresh_month(year, month)
    return _rb(ok="حُذف الشعار من هذا الشهر وتحدّثت مستنداته")


# ======================================================================
# عرض الشعار داخل الصفحة
# ======================================================================
@letterhead_bp.route("/logo")
@login_required
def logo_view():
    logo = db.get_setting(*_ctx(), "logo_file")
    path = storage.letterhead_dir(*_ctx()) / logo if logo else None
    if not path or not path.exists():
        abort(404)
    return send_file(str(path))


# ======================================================================
# ملف الوورد: تنزيل/فتح + فتح المجلد
# ======================================================================
def _open_path(path):
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(path))  # type: ignore[attr-defined]  # noqa
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)],
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return True
    except Exception:
        return False


@letterhead_bp.route("/download-docx")
@login_required
def download_docx():
    path = letterhead_docx.ensure(*_ctx())
    return attachment(path, "letterhead.docx")


@letterhead_bp.route("/download-zip")
@login_required
def download_zip():
    """تنزيل مجلد الدباجة كاملًا (الوورد + الشعار) أرشيف ZIP — يعمل من أي متصفح."""
    letterhead_docx.ensure(*_ctx())
    zip_path = dataguard.zip_folder_tmp(storage.letterhead_dir(*_ctx()), prefix="letterhead")
    resp = send_file(str(zip_path), as_attachment=True, download_name="letterhead.zip", conditional=False, max_age=0)
    resp.call_on_close(lambda: os.path.exists(zip_path) and os.remove(str(zip_path)))
    return resp


@letterhead_bp.route("/open-docx")
@login_required
def open_docx():
    path = letterhead_docx.ensure(*_ctx())
    if _open_path(path):
        return _rb(ok="تم فتح ملف الوورد 📄")
    return redirect(url_for("letterhead.download_docx",
                            sid=request.args.get("sid", "")))


@letterhead_bp.route("/open-folder")
@login_required
def open_folder():
    if _open_path(storage.letterhead_dir(*_ctx())):
        return _rb(ok="تم فتح مجلد الدباجة 📂")
    return _rb(err="فتح المجلد متاح عند تشغيل البرنامج على جهازك — استخدم زر التنزيل هنا")
